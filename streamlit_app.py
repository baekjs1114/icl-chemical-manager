# ICL V8.6 standalone website. Keep existing requirements and Streamlit Secrets.
import sys as _icl_sys
import types as _icl_types
_icl_guard = _icl_types.ModuleType("icl_ocr_ui_guard")
_icl_guard.__file__ = __file__ + ":icl_ocr_ui_guard"
exec(compile('"""Apply after all existing UI transformations; compatible with the supplied ICL app.\n\nThis module uses the application\'s existing requests and Streamlit objects. No ML\npackages or API key belong in the Streamlit repository.\n"""\nfrom __future__ import annotations\n\nimport ast\nimport hashlib\nimport json\n\n\ndef is_v8(result):\n    return isinstance(result, dict) and result.get("schema_version") == "icl.ocr.v8"\n\n\ndef candidates(result):\n    return list(result.get("review_candidates", [])) if is_v8(result) else []\n\n\ndef candidate_label(candidate):\n    sources = {"reviewed_reference": "Reviewed manufacturer reference", "official_product_page": "Manufacturer product page",\n               "pubchem_exact_cas_synonym": "PubChem CAS match", "inventory_unreviewed": "Existing inventory candidate"}\n    parts = [str(candidate.get("name", ""))]\n    if candidate.get("cas"):\n        parts.append("CAS " + str(candidate["cas"]))\n    maker = candidate.get("manufacturer") or candidate.get("observed_manufacturer") or ""\n    catalog = candidate.get("catalog") or candidate.get("observed_catalog") or ""\n    if catalog:\n        suffix = "" if candidate.get("catalog") else " (label)"\n        parts.append((str(maker) + " " + str(catalog)).strip() + suffix)\n    parts.append(sources.get(candidate.get("source"), "Source needs review"))\n    return " | ".join(parts)\n\n\ndef metadata(result):\n    defaults = {"label_name": "", "verified_label_name": "", "verified_label_name_source": "",\n                "label_name_confidence": 0, "manufacturer": "", "catalog": "", "cas": "",\n                "catalog_candidates": [], "catalog_ambiguous": False, "catalog_confidence": 0,\n                "raw_cas_candidates": [], "cas_candidates": [], "cas_verified": False, "cas_ambiguous": False,\n                "lot": "", "purity": "", "amount": None, "unit": ""}\n    if is_v8(result):\n        defaults.update(result.get("metadata", {}))\n        defaults["label_name_confidence"] = defaults.get("label_name_confidence") or 0\n    return defaults\n\n\ndef review_key(result, candidate):\n    value = str(result.get("review_token", "")) + ":" + str(candidate.get("candidate_id", ""))\n    return "icl_v8_review_" + hashlib.sha256(value.encode()).hexdigest()[:20]\n\n\ndef choice_key(result):\n    value = str(result.get("scan_id")) + ":" + str(result.get("review_token", ""))\n    return "icl_v8_choice_" + hashlib.sha256(value.encode()).hexdigest()[:20]\n\n\ndef render_review(env, result):\n    st = env["st"]\n    if not is_v8(result):\n        st.warning("이전 OCR 결과입니다. 새 OCR API로 다시 판독하세요.")\n        return\n    state = result.get("verification_status")\n    if state == "PRODUCT_VERIFIED":\n        st.success("제품 정보 확인 — 저장할 이름은 확인된 제품명을 사용합니다.")\n    elif state == "IDENTITY_VERIFIED":\n        st.info("물질 확인 — 제조사/제품번호는 별도 확인이 필요할 수 있습니다.")\n    elif state == "CONFLICT":\n        st.error("식별 정보 충돌 — 사진 그룹 또는 라벨을 확인하세요. 자동 적용할 수 없습니다.")\n    else:\n        st.warning("미확정 — 아래 후보와 원본을 비교한 뒤 선택하세요.")\n    for item in result.get("conflicts", []):\n        st.error(item)\n    with st.expander("OCR 원문과 제품 조회 근거"):\n        st.code("\\n".join(o["text"] for o in result.get("observations", []) if o.get("pass_id") == "full"))\n        st.json({"fields": result.get("fields", {}), "lookups": result.get("lookup_events", []),\n                 "warnings": result.get("warnings", []), "runtime": result.get("runtime", {})})\n    packages = result.get("fields", {}).get("package", [])\n    if packages:\n        st.caption("판독한 포장 표기: " + " / ".join(dict.fromkeys(x["value"] for x in packages)) + " · Remaining amount은 Manual entry하세요.")\n\n\ndef render_candidate(env, result, candidate):\n    if not is_v8(result):\n        return\n    st = env["st"]\n    st.write("적용할 이름: **" + str(candidate.get("name", "")) + "**")\n    maker = candidate.get("manufacturer") or candidate.get("observed_manufacturer") or ""\n    catalog = candidate.get("catalog") or candidate.get("observed_catalog") or ""\n    product_note = "" if candidate.get("product_verified") else " · label observation"\n    st.write("CAS: " + str(candidate.get("cas", "")) + " · Manufacturer: " + str(maker or "Not confirmed") + " · Catalog: " + str(catalog or "Not confirmed") + product_note)\n    if candidate.get("source_url"):\n        st.markdown("[제품/물질 출처 확인](" + candidate["source_url"] + ")")\n    if not candidate.get("product_verified"):\n        st.checkbox("원본 라벨과 후보의 이름·CAS·제품번호를 대조했습니다.", key=review_key(result, candidate))\n\n\ndef _post(env, path, payload):\n    response = env["requests"].post(env["get_ocr_api_url"]().rstrip("/") + path,\n        headers={"X-API-Key": env["get_ocr_api_key"]()}, json=payload, timeout=60)\n    if response.status_code != 200:\n        raise ValueError("OCR API " + str(response.status_code) + ": " + response.text[:250])\n    return response.json()\n\n\ndef apply_result(env, result, original_bytes, filename, selected_candidate=None, source_scan_id=None, evidence_photos=None):\n    st = env["st"]\n    if not is_v8(result):\n        st.error("This OCR result is from an older version. Analyze again or use manual entry.")\n        st.stop()\n    selected = selected_candidate or result.get("verified_identity")\n    if not selected:\n        st.warning("Select a product or chemical candidate first, or use manual entry.")\n        st.stop()\n    try:\n        draft = _post(env, "/draft", {"review_token": result.get("review_token"),\n            "candidate_id": selected.get("candidate_id"),\n            "reviewed": bool(st.session_state.get(review_key(result, selected), False))})\n    except Exception as exc:\n        st.error(str(exc))\n        st.stop()  # Prevent the old caller from displaying its success banner/rerun.\n    fields = draft["fields"]\n    codes = {str(x.get("value", "")).strip().upper() for x in result.get("fields", {}).get("catalog", [])}\n    if fields.get("catalog") and codes and str(fields["catalog"]).upper() not in codes and not draft.get("catalog_correction"):\n        st.error("Catalog mismatch. The selected product differs from the label; please rescan.")\n        st.stop()\n    for key, field in (("add_name", "name"), ("add_cas", "cas"), ("add_formula", "formula"),\n                       ("add_mw", "mw"), ("add_manufacturer", "manufacturer"), ("add_catalog", "catalog")):\n        value = fields.get(field)\n        st.session_state[key] = "" if value is None else str(value)\n    for key in ("add_lot", "add_purity", "add_description"):\n        st.session_state[key] = ""\n    # Clear a previous bottle\'s amount; no package quantity is inferred as remaining stock.\n    st.session_state["add_initial_amount"] = 0.0\n    st.session_state["add_remaining_amount"] = 0.0\n    st.session_state["add_unit"] = "g"\n    quantity = draft.get("quantity") or {}\n    if quantity.get("unit") in ("g", "mg", "kg", "mL", "L") and isinstance(quantity.get("amount"), (int, float)) and quantity["amount"] > 0:\n        st.session_state["add_initial_amount"] = float(quantity["amount"])\n        st.session_state["add_unit"] = quantity["unit"]\n    st.session_state["add_manufacturer_widget_needs_sync"] = True\n    st.session_state["add_ocr_result"] = result\n    st.session_state["add_ocr_original"] = original_bytes\n    st.session_state["add_ocr_filename"] = filename\n    st.session_state["icl_v8_registration_draft"] = draft\n    if evidence_photos is not None:\n        st.session_state["add_ocr_group_photos"] = evidence_photos\n    if source_scan_id:\n        st.session_state["add_source_scan_id"] = source_scan_id\n    st.session_state["_pending_navigation"] = "Add Chemical"\n    return draft\n\n\ndef merge_group(env, records, group_id):\n    successful = [r for r in records if not r.get("error") and r.get("result")]\n    base = {"group_id": group_id, "records": records, "result": {}, "error": "", "conflicts": []}\n    if not successful or not all(is_v8(r["result"]) for r in successful):\n        return {**base, "error": "Analyze this photo group again with the current OCR API."}\n    unique = {r["result"]["scan_id"]: r for r in successful}\n    if len(unique) == 1:\n        return {**base, "result": next(iter(unique.values()))["result"], "conflicts": successful[0]["result"].get("conflicts", [])}\n    tokens = sorted(r["result"]["evidence_token"] for r in unique.values())\n    signature = hashlib.sha256(json.dumps(tokens).encode()).hexdigest()\n    cache = env["st"].session_state.setdefault("icl_v8_group_cache", {})\n    try:\n        if signature not in cache:\n            # Errors are deliberately not cached. A subsequent rerun can retry 503s.\n            result = _post(env, "/resolve-group", {"evidence_tokens": tokens})\n            if not result.get("ok"):\n                raise ValueError("Group resolution failed")\n            cache[signature] = result\n        result = dict(cache[signature])\n        result["label_crop_base64"] = successful[0]["result"].get("label_crop_base64", "")\n        return {**base, "result": result, "conflicts": result.get("conflicts", [])}\n    except Exception as exc:\n        return {**base, "error": str(exc)}\n\n\ndef transform_source(source):\n    """Targeted source-span patches preserve unrelated UI and database code.\n\n    Run on the final UI-transformed source, including all fallback compile paths.\n    Unsupported source is an error, never a silently partial safety patch.\n    """\n    if "# ICL_V8_GUARD_APPLIED" in source:\n        return source\n    tree = ast.parse(source)\n    lines = source.splitlines(keepends=True)\n    offsets = [0]\n    for line in lines:\n        offsets.append(offsets[-1] + len(line))\n    def offset(lineno, byte_column):\n        # AST columns are UTF-8 byte offsets; Python string indices are characters.\n        return offsets[lineno - 1] + len(lines[lineno - 1].encode()[:byte_column].decode())\n    edits, found = [], set()\n    def replace(node, text):\n        edits.append((offset(node.lineno, node.col_offset), offset(node.end_lineno, node.end_col_offset), text))\n    helper = \'__import__("icl_ocr_ui_guard")\'\n    bodies = {\n        "ocr_candidate_label": f"return {helper}.candidate_label(candidate)",\n        "resolve_current_bottle_metadata": f"return {helper}.metadata(result)",\n        "merge_add_ocr_group": f"return {helper}.merge_group(globals(), records, group_id)",\n        "load_add_form_from_ocr": f"return {helper}.apply_result(globals(), result, original_bytes, filename, selected_candidate, source_scan_id, evidence_photos)",\n        "load_add_form_from_existing_bottle": f"return {helper}.apply_result(globals(), result, original_bytes, filename, source_scan_id=source_scan_id)",\n    }\n    replaced_regions = []\n    for node in tree.body:\n        if isinstance(node, ast.FunctionDef) and node.name in bodies:\n            found.add(node.name)\n            replace(node, "def " + node.name + "(" + ast.unparse(node.args) + "):\\n    " + bodies[node.name])\n            replaced_regions.append((node.lineno, node.end_lineno))\n    missing = set(bodies) - found\n    if missing:\n        raise ValueError("Unsupported Streamlit source; missing functions: " + ", ".join(sorted(missing)))\n    for node in ast.walk(tree):\n        if not hasattr(node, "lineno") or any(a <= node.lineno <= b for a, b in replaced_regions):\n            continue\n        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "add_candidates" for t in node.targets):\n            indent = " " * node.col_offset\n            replace(node, f"add_candidates = {helper}.candidates(add_ocr_result)\\n{indent}{helper}.render_review(globals(), add_ocr_result)")\n        if isinstance(node, ast.If) and isinstance(node.test, ast.Name) and node.test.id == "chosen_candidate":\n            first = node.body[0]\n            start = offset(first.lineno, first.col_offset)\n            edits.append((start, start, f"{helper}.render_candidate(globals(), add_ocr_result, chosen_candidate)\\n" + " " * first.col_offset))\n        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and isinstance(node.func.value, ast.Name) and node.func.value.id == "st":\n            if node.func.attr == "radio" and any(k.arg == "key" and isinstance(k.value, ast.Constant) and k.value.value == "add_chemical_candidate_choice" for k in node.keywords):\n                for k in node.keywords:\n                    if k.arg == "index":\n                        replace(k.value, "None")\n                    elif k.arg == "key":\n                        replace(k.value, f"{helper}.choice_key(add_ocr_result)")\n            if node.func.attr == "success" and node.args and isinstance(node.args[0], ast.Constant) and str(node.args[0].value) in ("Label fields verified", "Core label fields verified"):\n                replace(node, \'st.caption("필드 표시 완료 · 제품/물질 검증 상태는 아래에서 확인하세요.")\')\n    for start, end, text in sorted(edits, reverse=True):\n        source = source[:start] + text + source[end:]\n    source += "\\n# ICL_V8_GUARD_APPLIED\\n"\n    compile(source, "<ICL v8 guarded source>", "exec")\n    return source\n', _icl_guard.__file__, "exec"), _icl_guard.__dict__)
_icl_sys.modules["icl_ocr_ui_guard"] = _icl_guard

import json
import base64
import hashlib
import re
import os
import zipfile
import queue
import threading
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from difflib import SequenceMatcher

import psycopg
from psycopg.rows import dict_row
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps


# Per-click timing shown at the bottom of the sidebar: total script time and
# how much of it was spent waiting on the database.
_ICL_PERF = {"start": time.perf_counter(), "db_calls": 0, "db_seconds": 0.0}


def _icl_db_timed(fn, *args, **kwargs):
    started = time.perf_counter()
    try:
        return fn(*args, **kwargs)
    finally:
        _ICL_PERF["db_calls"] += 1
        _ICL_PERF["db_seconds"] += time.perf_counter() - started


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="ICL Chemical Manager",
    page_icon="🧪",
    layout="wide",
)

def _icl_exception_message(exc):
    message = str(exc).strip()
    if not message:
        message = "No error details were returned. Share this screen with the error type."
    return type(exc).__name__ + ": " + message


st.caption("ICL V8.6 · vendor-aware label decisions")
st.markdown("""
<style>


/* =========================================================
   GLOBAL
   ========================================================= */

.stApp {
    background-color: #F7F9FC;
}

/* Main content shell — shared by ALL categories */
.block-container {
    width: 100%;
    max-width: 1180px !important;
    margin-left: auto !important;
    margin-right: auto !important;
    padding-top: 1.55rem !important;
    padding-bottom: 2.6rem !important;
    padding-left: 1.35rem !important;
    padding-right: 1.35rem !important;
}

/* Keep the whole application compact and centered on large monitors. */
@media (min-width: 1500px) {
    .block-container {
        max-width: 1180px !important;
    }
}

/* Tablet / small desktop */
@media (max-width: 1100px) {
    .block-container {
        max-width: 100% !important;
        padding-left: 1rem !important;
        padding-right: 1rem !important;
    }
}

/* Mobile */
@media (max-width: 700px) {
    .block-container {
        padding-top: 1rem !important;
        padding-left: 0.75rem !important;
        padding-right: 0.75rem !important;
    }
}



/* =========================================================
   TYPOGRAPHY — original UI + 2px
   ========================================================= */

.stApp,
.stApp p,
.stApp li,
.stApp label,
.stApp input,
.stApp textarea,
.stApp button,
.stApp [data-baseweb="select"],
[data-testid="stSidebar"] {
    font-size: 18px !important;
}

.stApp [data-testid="stCaptionContainer"] {
    font-size: 16px !important;
}

h1 {
    font-size: 36px !important;
}

h2 {
    font-size: 30px !important;
}

h3 {
    font-size: 25px !important;
}

h4 {
    font-size: 21px !important;
}


/* =========================================================
   SIDEBAR
   ========================================================= */

[data-testid="stSidebar"] {
    background-color: #FFFFFF;
    border-right: 1px solid #E5E7EB;
}

[data-testid="stSidebar"] > div:first-child {
    padding-top: 1.5rem;
}


/* =========================================================
   TITLES
   ========================================================= */

h1 {
    color: #172033;
    font-weight: 700;
    letter-spacing: -0.5px;
}

h2, h3 {
    color: #25324B;
}


/* =========================================================
   INPUTS
   ========================================================= */

.stTextInput input {
    border-radius: 10px;
    border: 1px solid #D8DEE9;
    background-color: #FFFFFF;
}

.stTextInput input:focus {
    border-color: #3B82F6;
}


/* =========================================================
   SELECTBOX
   ========================================================= */

[data-baseweb="select"] > div {
    border-radius: 10px;
}


/* =========================================================
   BUTTONS
   ========================================================= */

.stButton > button {
    border-radius: 8px;
    padding: 0.45rem 1rem;
    font-weight: 500;
    border: 1px solid #D8DEE9;
    white-space: nowrap;
}

.stButton > button:hover {
    border-color: #3B82F6;
}


/* Primary button */
button[kind="primary"] {
    background-color: #2563EB;
    color: white;
    border: none;
}

button[kind="primary"]:hover {
    background-color: #1D4ED8;
}


/* =========================================================
   DATAFRAME
   ========================================================= */

[data-testid="stDataFrame"] {
    border-radius: 12px;
    overflow: hidden;
    border: 1px solid #E5E7EB;
}


/* =========================================================
   METRIC / CARDS
   ========================================================= */

[data-testid="stMetric"] {
    background: #FFFFFF;
    padding: 18px;
    border-radius: 12px;
    border: 1px solid #E5E7EB;
}


/* =========================================================
   FORMS
   ========================================================= */

[data-testid="stForm"] {
    background-color: #FFFFFF;
    padding: 22px;
    border-radius: 14px;
    border: 1px solid #E5E7EB;
}


/* =========================================================
   TABS
   ========================================================= */

button[data-baseweb="tab"] {
    font-size: 17px;
    font-weight: 500;
}



/* =========================================================
   SHARED UX / UI — applied to ALL categories
   ========================================================= */

/* More deliberate page rhythm */
main [data-testid="stVerticalBlock"] {
    gap: 0.9rem;
}

main [data-testid="stHorizontalBlock"] {
    gap: 0.9rem;
}

/* Page headings: less oversized, easier to scan */
h1 {
    margin-top: 0.1rem !important;
    margin-bottom: 0.45rem !important;
    line-height: 1.18 !important;
}

h2 {
    margin-top: 0.55rem !important;
    margin-bottom: 0.35rem !important;
    line-height: 1.2 !important;
}

h3, h4 {
    margin-top: 0.45rem !important;
    margin-bottom: 0.25rem !important;
}

/* Labels become consistent across Search / Inventory / Add / Storage / SDS */
.stApp label {
    font-weight: 600 !important;
    color: #30394A !important;
}

/* Uniform control height and density */
.stTextInput input,
.stNumberInput input,
.stDateInput input,
.stTextArea textarea,
[data-baseweb="select"] > div {
    min-height: 44px;
}

.stTextArea textarea {
    min-height: 92px;
}

/* Expander = reusable compact card */
[data-testid="stExpander"] {
    background: #FFFFFF;
    border: 1px solid #E3E8F0 !important;
    border-radius: 12px !important;
    overflow: hidden;
}

[data-testid="stExpander"] details summary {
    padding-top: 0.55rem !important;
    padding-bottom: 0.55rem !important;
}

/* File uploader: smaller and less dominant */
[data-testid="stFileUploader"] {
    margin-top: 0.1rem;
    margin-bottom: 0.2rem;
}

[data-testid="stFileUploaderDropzone"] {
    background: #FFFFFF !important;
    border: 1px dashed #C7D0DF !important;
    border-radius: 12px !important;
    padding: 0.65rem 0.8rem !important;
    min-height: 76px !important;
}

[data-testid="stFileUploaderDropzoneInstructions"] {
    padding-top: 0 !important;
    padding-bottom: 0 !important;
}

/* Alerts should read like compact information panels */
[data-testid="stAlert"] {
    border-radius: 11px !important;
    padding: 0.7rem 0.9rem !important;
    margin-top: 0.15rem !important;
    margin-bottom: 0.35rem !important;
}

/* Forms across Add Chemical / Storage / Admin / SDS */
[data-testid="stForm"] {
    padding: 1rem 1.1rem !important;
    border-radius: 12px !important;
    box-shadow: 0 1px 2px rgba(24, 36, 56, 0.03);
}

/* Metrics: smaller cards so dashboards do not sprawl */
[data-testid="stMetric"] {
    padding: 0.8rem 0.9rem !important;
    min-height: 92px;
}

/* Tables stay inside the shared content shell and scroll when necessary */
[data-testid="stDataFrame"] {
    width: 100% !important;
    max-width: 100% !important;
}

/* Buttons: unified height */
.stButton > button,
.stDownloadButton > button,
[data-testid="stFormSubmitButton"] > button {
    min-height: 42px;
}

/* Divider density */
hr {
    margin-top: 0.7rem !important;
    margin-bottom: 0.7rem !important;
}

/* Captions are readable but visually secondary */
[data-testid="stCaptionContainer"] {
    color: #6B7280 !important;
    line-height: 1.45 !important;
}

/* Sidebar remains readable without growing with the main content. */
[data-testid="stSidebar"] {
    min-width: 245px;
}

@media (max-width: 900px) {
    main [data-testid="stHorizontalBlock"] {
        gap: 0.65rem;
    }

    [data-testid="stMetric"] {
        min-height: 82px;
    }
}


/* =========================================================
   HIDE STREAMLIT DEFAULT UI
   ========================================================= */

#MainMenu {
    visibility: hidden;
}

footer {
    visibility: hidden;
}

</style>
""", unsafe_allow_html=True)

st.markdown(
    '<style>:root {\n  --icl-navy:#01194C;\n  --icl-navy-2:#0A2A64;\n  --icl-blue:#2F5FB7;\n  --icl-bg:#F7F8FB;\n  --icl-border:#E1E6EF;\n  --icl-text:#172033;\n  --icl-muted:#6D7480;\n}\n\n.stApp {\n  background:linear-gradient(180deg,#fff 0,#fff 116px,var(--icl-bg) 116px)!important;\n  color:var(--icl-text);\n}\n\n.block-container {\n  max-width:1220px!important;\n  padding-top:1.05rem!important;\n  padding-bottom:2.2rem!important;\n}\n\n[data-testid="stHeader"] { background:transparent!important; }\n\n/* Original lab website character: white left rail + typography-first navigation. */\n[data-testid="stSidebar"] {\n  min-width:270px!important;\n  background:#fff!important;\n  border-right:1px solid #E7EAF0!important;\n}\n[data-testid="stSidebar"] > div:first-child {\n  padding-top:.4rem!important;\n}\n[data-testid="stSidebar"] img {\n  display:block!important;\n  width:auto!important;\n  max-width:238px!important;\n  height:auto!important;\n  max-height:84px!important;\n  object-fit:contain!important;\n  margin:.05rem auto .08rem auto!important;\n}\n.icl-sidebar-product {\n  color:#8A92A0;\n  font-size:13px;\n  font-weight:800;\n  letter-spacing:.15em;\n  margin:.16rem .1rem .5rem .1rem;\n}\n\n/* Legacy radio navigation, retained for any untouched page fragments. */\n[data-testid="stSidebar"] [role="radiogroup"] { gap:.16rem!important; }\n[data-testid="stSidebar"] [role="radiogroup"] label {\n  border-radius:0!important;\n  padding:.48rem .62rem!important;\n  font-weight:600!important;\n  letter-spacing:.025em!important;\n  color:#222833!important;\n  border-left:3px solid transparent!important;\n  transition:background .12s ease,border-color .12s ease,color .12s ease;\n}\n[data-testid="stSidebar"] [role="radiogroup"] label:hover {\n  background:#F4F6FA!important;\n  color:var(--icl-navy)!important;\n}\n[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) {\n  background:#F2F5FA!important;\n  border-left-color:var(--icl-navy)!important;\n  color:var(--icl-navy)!important;\n}\n\n/* UX v3 hierarchical navigation. */\n[data-testid="stSidebar"] .stButton > button {\n  justify-content:flex-start!important;\n  text-align:left!important;\n  width:100%!important;\n  min-height:39px!important;\n  padding:.44rem .62rem!important;\n  margin:0!important;\n  border:0!important;\n  border-left:3px solid transparent!important;\n  border-radius:0!important;\n  background:transparent!important;\n  color:#222833!important;\n  box-shadow:none!important;\n  font-size:17px!important;\n  font-weight:600!important;\n  letter-spacing:.018em!important;\n}\n[data-testid="stSidebar"] .stButton > button:hover {\n  background:#F4F6FA!important;\n  color:var(--icl-navy)!important;\n}\n[data-testid="stSidebar"] .stButton > button[kind="primary"] {\n  background:#F2F5FA!important;\n  color:var(--icl-navy)!important;\n  border-left-color:var(--icl-navy)!important;\n}\n\n/* Compact ICL workspace header: content-sized instead of full-width billboard. */\n.icl-shell-header {\n  position:relative;\n  display:inline-flex;\n  width:auto;\n  max-width:100%;\n  align-items:center;\n  justify-content:flex-start;\n  gap:.72rem;\n  overflow:hidden;\n  background:#fff;\n  border:1px solid var(--icl-border);\n  border-radius:4px;\n  padding:.58rem .72rem .58rem 1rem;\n  margin-bottom:.72rem;\n}\n.icl-shell-header::before {\n  content:"";\n  position:absolute;\n  left:0;top:0;bottom:0;width:4px;\n  background:var(--icl-navy);\n}\n.icl-shell-copy { min-width:0; }\n.icl-shell-eyebrow {\n  font-size:13px;\n  font-weight:800;\n  letter-spacing:.14em;\n  color:#7D8795;\n  margin-bottom:.12rem;\n  white-space:nowrap;\n  overflow:hidden;\n  text-overflow:ellipsis;\n}\n.icl-shell-title {\n  font-size:29.5px;\n  font-weight:780;\n  letter-spacing:-.035em;\n  color:var(--icl-navy);\n  line-height:1.02;\n}\n.icl-shell-pills {\n  display:flex;\n  flex-wrap:wrap;\n  justify-content:flex-start;\n  gap:.25rem;\n}\n.icl-shell-pills span {\n  display:inline-flex;\n  align-items:center;\n  min-height:24px;\n  padding:.18rem .44rem;\n  border:1px solid #DCE3EE;\n  background:#F7F9FC;\n  border-radius:999px;\n  color:#526070;\n  font-size:14px;\n  font-weight:700;\n  letter-spacing:.02em;\n}\n\nh1,h2,h3,h4 { color:var(--icl-navy)!important;letter-spacing:-.025em!important; }\nh1 { font-size:34px!important; }\nh2 { font-size:27px!important; }\nh3 { font-size:21px!important; }\n\n.stApp,.stApp p,.stApp li,.stApp label,.stApp input,.stApp textarea,\n.stApp button,.stApp [data-baseweb="select"] { font-size:17px!important; }\n\n.stTextInput input,.stNumberInput input,.stDateInput input,.stTextArea textarea,\n[data-baseweb="select"] > div {\n  border-color:#D7DDE8!important;\n  border-radius:6px!important;\n  background:#fff!important;\n}\n\n/* Main-canvas actions: readable at a glance when reviewing many bottles. */\nmain .stButton > button,\nmain .stDownloadButton > button,\nmain [data-testid="stFormSubmitButton"] > button,\nmain [data-testid="stLinkButton"] a {\n  border-radius:5px!important;\n  font-weight:680!important;\n  min-height:48px!important;\n  padding:.64rem 1.05rem!important;\n  font-size:18px!important;\n  line-height:1.2!important;\n}\n\n/* Primary button text must stay readable on ICL navy. */\nmain button[kind="primary"],\nmain button[kind="primary"] *,\nmain [data-testid="stFormSubmitButton"] button[kind="primary"],\nmain [data-testid="stFormSubmitButton"] button[kind="primary"] * {\n  background:var(--icl-navy)!important;\n  border-color:var(--icl-navy)!important;\n  color:#fff!important;\n  -webkit-text-fill-color:#fff!important;\n}\nmain button[kind="primary"]:hover,\nmain button[kind="primary"]:hover * {\n  background:var(--icl-navy-2)!important;\n  border-color:var(--icl-navy-2)!important;\n  color:#fff!important;\n  -webkit-text-fill-color:#fff!important;\n}\n\n/* Disabled actions should look disabled, not like dark unreadable active buttons. */\nmain button:disabled,\nmain button:disabled * {\n  opacity:1!important;\n  background:#E7ECF4!important;\n  border-color:#D7DFEA!important;\n  color:#7D8797!important;\n  -webkit-text-fill-color:#7D8797!important;\n}\n\n[data-testid="stForm"],[data-testid="stExpander"],[data-testid="stMetric"],[data-testid="stDataFrame"] {\n  border-color:var(--icl-border)!important;\n  border-radius:7px!important;\n  box-shadow:none!important;\n}\n[data-testid="stAlert"] { border-radius:6px!important; }\n\n/* Chemical/Search + Storage details open as modal dialogs. */\n[data-testid="stDialog"] [role="dialog"] {\n  border-radius:8px!important;\n  border:1px solid #DDE3ED!important;\n  box-shadow:0 24px 80px rgba(1,25,76,.18)!important;\n}\n\n/* Add Chemical is the intentionally redesigned workflow. */\n.icl-add-hero {\n  display:flex;\n  align-items:center;\n  justify-content:space-between;\n  gap:1rem;\n  background:var(--icl-navy);\n  color:#fff;\n  border-radius:9px;\n  padding:1.22rem 1.38rem;\n  margin:.1rem 0 1rem 0;\n}\n.icl-add-kicker { font-size:13px;font-weight:800;letter-spacing:.16em;opacity:.72; }\n.icl-add-title { font-size:29px;font-weight:780;letter-spacing:-.025em;margin-top:.12rem; }\n.icl-add-steps {\n  display:flex;\n  align-items:center;\n  flex-wrap:wrap;\n  justify-content:flex-end;\n  gap:.42rem;\n}\n.icl-add-steps span {\n  display:inline-flex;\n  align-items:center;\n  gap:.3rem;\n  padding:.38rem .58rem;\n  border:1px solid rgba(255,255,255,.24);\n  border-radius:999px;\n  background:rgba(255,255,255,.08);\n  font-size:14.5px;\n  font-weight:650;\n  white-space:nowrap;\n}\n.icl-add-steps b { color:#BFD0FF;font-size:12.7px;letter-spacing:.06em; }\n.icl-result-line {\n  border-left:4px solid var(--icl-navy);background:#fff;\n  border-top:1px solid var(--icl-border);border-right:1px solid var(--icl-border);\n  border-bottom:1px solid var(--icl-border);padding:.72rem .82rem;\n  border-radius:4px;margin:.25rem 0 .55rem 0;\n}\n\n/* Inventory Scan: compact workflow card instead of developer prose. */\n.icl-scan-hero {\n  background:#fff;\n  border:1px solid var(--icl-border);\n  border-left:6px solid var(--icl-navy);\n  border-radius:7px;\n  padding:1.05rem 1.15rem;\n  margin:.1rem 0 .75rem 0;\n}\n.icl-scan-kicker {\n  font-size:13px;\n  font-weight:800;\n  letter-spacing:.15em;\n  color:var(--icl-blue);\n}\n.icl-scan-title {\n  color:var(--icl-navy);\n  font-size:27px;\n  font-weight:760;\n  letter-spacing:-.025em;\n  margin-top:.15rem;\n}\n.icl-scan-flow {\n  display:flex;\n  align-items:center;\n  flex-wrap:wrap;\n  gap:.4rem;\n  margin:.42rem 0 .48rem 0;\n}\n.icl-scan-flow span {\n  display:inline-flex;\n  align-items:center;\n  gap:.35rem;\n  border:1px solid #DCE3EE;\n  border-radius:999px;\n  padding:.38rem .62rem;\n  background:#fff;\n  color:#334155;\n  font-size:15px;\n  font-weight:650;\n}\n.icl-scan-flow span b {\n  display:inline-grid;\n  place-items:center;\n  width:1.25rem;\n  height:1.25rem;\n  border-radius:50%;\n  background:var(--icl-navy);\n  color:#fff;\n  font-size:13.5px;\n}\n.icl-scan-flow i {\n  color:#94A3B8;\n  font-style:normal;\n}\n.icl-scan-note-row {\n  display:flex;\n  flex-wrap:wrap;\n  gap:.38rem;\n}\n.icl-scan-note-row span {\n  display:inline-flex;\n  align-items:center;\n  min-height:25px;\n  padding:.22rem .5rem;\n  border-radius:4px;\n  background:#F2F5FA;\n  color:#5B6574;\n  font-size:13.7px;\n  font-weight:650;\n}\n\n/* Bottle-photo uploader: taller hit target + larger, legible Upload button. */\n[data-testid="stFileUploaderDropzone"] {\n  min-height:132px!important;\n  padding:1.45rem 1.25rem!important;\n  display:flex!important;\n  align-items:center!important;\n  background:#fff!important;\n  border:1.5px dashed #92A3BD!important;\n  border-radius:9px!important;\n}\n[data-testid="stFileUploaderDropzone"] > div { width:100%!important; }\n[data-testid="stFileUploaderDropzone"] p,\n[data-testid="stFileUploaderDropzone"] span,\n[data-testid="stFileUploaderDropzone"] small {\n  font-size:16px!important;\n  line-height:1.45!important;\n}\n[data-testid="stFileUploaderDropzone"] button {\n  min-height:48px!important;\n  padding:.64rem 1.05rem!important;\n  font-size:18px!important;\n  font-weight:700!important;\n}\n\n/* Mobile camera capture should feel like the primary field workflow, not a debug widget. */\n[data-testid="stCameraInput"] {\n  border:1px solid var(--icl-border)!important;\n  border-radius:9px!important;\n  background:#fff!important;\n  padding:.45rem!important;\n}\n[data-testid="stCameraInput"] button {\n  min-height:52px!important;\n  font-size:18px!important;\n  font-weight:750!important;\n}\n[data-baseweb="tab-list"] {\n  gap:.3rem!important;\n}\n[data-baseweb="tab"] {\n  min-height:44px!important;\n  font-weight:700!important;\n}\n\n/* Responsive inventory table: scroll horizontally instead of collapsing. */\n[data-testid="stDataFrame"] {\n  width:100%!important;\n  max-width:100%!important;\n  overflow-x:auto!important;\n}\n\n@media (max-width:800px) {\n  [data-testid="stSidebar"] { min-width:unset!important; }\n  .block-container {\n    padding-left:.65rem!important;\n    padding-right:.65rem!important;\n    padding-top:.65rem!important;\n  }\n  .icl-shell-header,.icl-add-hero {\n    align-items:flex-start;\n    flex-direction:column;\n  }\n  .icl-shell-header {\n    padding:.62rem .68rem .62rem .92rem;\n    margin-bottom:.62rem;\n  }\n  .icl-shell-eyebrow { font-size:12px;letter-spacing:.12em; }\n  .icl-shell-title { font-size:24.5px; }\n  .icl-shell-pills,.icl-add-steps { justify-content:flex-start; }\n  h1 { font-size:27.5px!important; }\n  h2 { font-size:23.5px!important; }\n  h3 { font-size:19.5px!important; }\n  [data-testid="stFileUploaderDropzone"] {\n    min-height:116px!important;\n    padding:1rem .85rem!important;\n  }\n  [data-testid="stCameraInput"] {\n    padding:.25rem!important;\n  }\n  [data-testid="stCameraInput"] button {\n    min-height:52px!important;\n    width:100%!important;\n    font-size:18px!important;\n  }\n  .icl-scan-hero { padding:.9rem .85rem;margin-bottom:.6rem; }\n  .icl-scan-title { font-size:23.5px; }\n  .icl-scan-flow { gap:.32rem;margin-bottom:.4rem; }\n  .icl-scan-flow span { padding:.34rem .52rem;font-size:14.2px; }\n  .icl-scan-flow i { display:none; }\n  .icl-scan-note-row span { font-size:13px; }\n  main .stButton > button,\n  main .stDownloadButton > button,\n  main [data-testid="stFormSubmitButton"] > button {\n    min-height:46px!important;\n    font-size:17px!important;\n  }\n  [data-testid="stDataFrame"] { min-height:280px; }\n  [data-testid="stDialog"] [role="dialog"] {\n    width:calc(100vw - 1rem)!important;\n    max-width:calc(100vw - 1rem)!important;\n    margin:.5rem!important;\n  }\n}\n\n/* Product UI polish — intentional density, quiet surfaces, visible focus. */\n:root {\n  --icl-space-1:4px;\n  --icl-space-2:8px;\n  --icl-space-3:12px;\n  --icl-space-4:16px;\n  --icl-space-5:24px;\n  --icl-input:#F3F5F8;\n  --icl-border-soft:#E7EAF0;\n  --icl-focus:#2F5FB7;\n}\n\n.icl-shell-header {\n  padding:12px 16px !important;\n  gap:12px !important;\n  border-color:var(--icl-border-soft) !important;\n}\n.icl-shell-title { font-size:27px !important; }\n.icl-shell-pills { gap:4px !important; }\n.icl-add-hero {\n  padding:20px 24px !important;\n  margin:8px 0 24px 0 !important;\n  border-radius:8px !important;\n  box-shadow:none !important;\n}\n.icl-add-steps { gap:8px !important; }\n\n.stTextInput input,\n.stNumberInput input,\n.stDateInput input,\n.stTextArea textarea,\n[data-baseweb="select"] > div {\n  background:var(--icl-input) !important;\n  border-color:transparent !important;\n  box-shadow:inset 0 0 0 1px #DDE3EC !important;\n}\n.stTextInput input:focus,\n.stNumberInput input:focus,\n.stDateInput input:focus,\n.stTextArea textarea:focus {\n  box-shadow:inset 0 0 0 1px var(--icl-focus), 0 0 0 2px rgba(47,95,183,.18) !important;\n}\n\nmain button:focus-visible,\n[data-testid="stSidebar"] button:focus-visible,\nmain a:focus-visible,\n[data-baseweb="select"]:focus-within {\n  outline:2px solid var(--icl-focus) !important;\n  outline-offset:2px !important;\n}\n\n[data-testid="stMetricValue"],\ninput[type="number"],\n[data-testid="stDataFrame"] {\n  font-variant-numeric:tabular-nums;\n}\n\n[data-testid="stForm"],\n[data-testid="stExpander"],\n[data-testid="stMetric"],\n[data-testid="stDataFrame"] {\n  border-color:var(--icl-border-soft) !important;\n  box-shadow:none !important;\n}\n[data-testid="stForm"] { background:#fff !important; }\n[data-testid="stAlert"] { box-shadow:none !important; }\n\n[data-testid="stSidebar"] .stButton > button { transition:none !important; }\n@media (hover:hover) {\n  [data-testid="stSidebar"] .stButton > button:hover {\n    background:#F3F5F8 !important;\n  }\n}\n\n@media (max-width:800px) {\n  .icl-shell-pills { display:none !important; }\n  .icl-add-hero { padding:16px !important; margin-bottom:16px !important; }\n}\n</style>',
    unsafe_allow_html=True,
)

st.markdown(
    "<style>" + '\n/* ICL UX v3.9 compact desktop layout */\n@media (min-width: 1101px) {\n  .block-container {\n    max-width: 820px !important;\n    margin-left: auto !important;\n    margin-right: auto !important;\n  }\n}\n\n/* The global Chemical Manager header and Add Chemical hero use the exact same\n   content width. The older shell intentionally used inline-flex/auto width,\n   which made it look shorter than the full-width Add Chemical card. */\n.icl-shell-header,\n.icl-add-hero {\n  width: 100% !important;\n  max-width: 100% !important;\n  box-sizing: border-box !important;\n}\n.icl-shell-header {\n  display: flex !important;\n  justify-content: space-between !important;\n}\n.icl-shell-pills {\n  margin-left: auto !important;\n}\n\n/* Defensive cap for Streamlit image DOM variants. Explicit st.image widths below\n   are the primary sizing mechanism; this is only a browser-side safeguard. */\nmain img {\n  max-width: 280px;\n  height: auto;\n}\n\n@media (max-width: 800px) {\n  .block-container { max-width: 100% !important; }\n  main img { max-width: min(280px, 100%); }\n  .icl-shell-header,\n  .icl-add-hero { width: 100% !important; }\n}\n' + "</style>",
    unsafe_allow_html=True,
)

# ============================================================
# DATABASE — SUPABASE POSTGRESQL
# ============================================================


def _database_url():
    """Use Codespaces env var locally or Streamlit Secrets in deployment."""
    env_url = os.getenv("SUPABASE_DB_URL", "").strip()
    if env_url:
        return env_url

    # Streamlit Community Cloud exposes top-level secrets as environment
    # variables in many deployments, but also support direct st.secrets access.
    try:
        direct_url = str(st.secrets.get("SUPABASE_DB_URL", "")).strip()
        if direct_url:
            return direct_url
    except Exception:
        pass

    try:
        nested_url = str(st.secrets["database"]["url"]).strip()
        if nested_url:
            return nested_url
    except Exception as exc:
        raise RuntimeError(
            "Supabase database URL is not configured. "
            "Set SUPABASE_DB_URL locally or add it to Streamlit Secrets."
        ) from exc

    raise RuntimeError(
        "Supabase database URL is not configured. "
        "Set SUPABASE_DB_URL locally or add it to Streamlit Secrets."
    )


def _translate_qmarks(sql):
    """Translate SQLite-style ? placeholders to psycopg %s safely."""
    out = []
    in_single = False
    in_double = False
    i = 0

    while i < len(sql):
        ch = sql[i]

        if ch == "'" and not in_double:
            # SQL escaped quote: ''
            if in_single and i + 1 < len(sql) and sql[i + 1] == "'":
                out.append("''")
                i += 2
                continue
            in_single = not in_single
            out.append(ch)

        elif ch == '"' and not in_single:
            in_double = not in_double
            out.append(ch)

        elif ch == "?" and not in_single and not in_double:
            out.append("%s")

        else:
            out.append(ch)

        i += 1

    return "".join(out)


class CompatRow(dict):
    """Row compatible with sqlite3.Row: supports row['name'] and row[0]."""

    def __getitem__(self, key):
        if isinstance(key, int):
            return list(self.values())[key]
        return super().__getitem__(key)


class CompatCursor:
    def __init__(self, cursor):
        self._cursor = cursor
        self.lastrowid = None

    def execute(self, sql, params=None):
        translated = _translate_qmarks(sql)
        stripped = translated.lstrip()
        is_insert = stripped.upper().startswith("INSERT INTO ")
        has_returning = bool(re.search(r"\bRETURNING\b", translated, flags=re.I))

        if is_insert and not has_returning:
            translated = translated.rstrip().rstrip(";") + " RETURNING id"

        _icl_db_timed(self._cursor.execute, translated, params or ())

        if is_insert:
            try:
                returned = self._cursor.fetchone()
                if returned:
                    if isinstance(returned, dict):
                        self.lastrowid = returned.get("id")
                    else:
                        self.lastrowid = returned[0]
            except psycopg.ProgrammingError:
                self.lastrowid = None

        return self

    def fetchone(self):
        row = self._cursor.fetchone()
        return CompatRow(row) if row is not None else None

    def fetchall(self):
        return [CompatRow(row) for row in self._cursor.fetchall()]

    @property
    def rowcount(self):
        return self._cursor.rowcount

    def close(self):
        self._cursor.close()


class _RawConnectionPool:
    """Small thread-safe PostgreSQL connection pool for Streamlit reruns.

    Streamlit reruns the script for most widget interactions. The old app opened
    a brand-new TLS/PostgreSQL connection every time get_connection() was called,
    often several times during one rerun. Reusing a handful of live connections
    removes that handshake from normal clicks.
    """

    def __init__(self, max_size=6):
        self.max_size = int(max(1, max_size))
        self._available = queue.LifoQueue(maxsize=self.max_size)
        self._lock = threading.Lock()
        self._created = 0

    def _create(self):
        return psycopg.connect(
            _database_url(),
            row_factory=dict_row,
            connect_timeout=10,
            keepalives=1,
            keepalives_idle=30,
            keepalives_interval=10,
            keepalives_count=3,
        )

    def acquire(self, timeout=10):
        while True:
            try:
                raw = self._available.get_nowait()
            except queue.Empty:
                create_new = False
                with self._lock:
                    if self._created < self.max_size:
                        self._created += 1
                        create_new = True

                if create_new:
                    try:
                        return self._create()
                    except Exception:
                        with self._lock:
                            self._created = max(0, self._created - 1)
                        raise

                raw = self._available.get(timeout=timeout)

            if raw is None:
                continue

            if getattr(raw, "closed", False):
                with self._lock:
                    self._created = max(0, self._created - 1)
                continue

            return raw

    def release(self, raw):
        if raw is None:
            return

        if getattr(raw, "closed", False):
            with self._lock:
                self._created = max(0, self._created - 1)
            return

        # SELECT statements also open a transaction with psycopg's default
        # transaction mode. Reset before another session/request reuses it.
        try:
            raw.rollback()
        except Exception:
            try:
                raw.close()
            except Exception:
                pass
            with self._lock:
                self._created = max(0, self._created - 1)
            return

        try:
            self._available.put_nowait(raw)
        except queue.Full:
            try:
                raw.close()
            finally:
                with self._lock:
                    self._created = max(0, self._created - 1)


@st.cache_resource(show_spinner=False)
def _get_db_connection_pool():
    # Six connections is enough for the Streamlit UI plus a few concurrent lab
    # users without creating a large persistent Supabase connection footprint.
    return _RawConnectionPool(max_size=6)


class CompatConnection:
    def __init__(self, raw_conn, pool=None):
        self._conn = raw_conn
        self._pool = pool
        self._returned = False

    def execute(self, sql, params=None):
        cur = CompatCursor(self._conn.cursor(row_factory=dict_row))
        return cur.execute(sql, params)

    def commit(self):
        self._conn.commit()

    def rollback(self):
        self._conn.rollback()

    def close(self):
        if self._returned:
            return
        self._returned = True

        if self._pool is None:
            self._conn.close()
        else:
            _icl_db_timed(self._pool.release, self._conn)

    def __del__(self):
        # Many callers close only on the happy path. An exception, st.rerun()
        # or st.stop() between get_connection() and close() used to keep the
        # connection checked out forever; after six of those every user got
        # queue.Empty at login. Return it when the wrapper is collected.
        try:
            self.close()
        except Exception:
            pass


def get_connection():
    pool = _get_db_connection_pool()
    try:
        raw = _icl_db_timed(pool.acquire, timeout=1)
    except queue.Empty:
        # Every pooled connection is busy: use a one-off connection rather
        # than failing the page. It is closed (not pooled) on close().
        return CompatConnection(pool._create(), pool=None)
    return CompatConnection(raw, pool=pool)


@st.cache_resource(show_spinner=False)
def init_db():
    """Initialize the Supabase schema once per app process.

    A PostgreSQL transaction-level advisory lock serializes schema initialization
    so simultaneous Streamlit sessions cannot deadlock while creating indexes.
    """
    conn = get_connection()

    # Prevent concurrent app sessions from running DDL/index creation together.
    conn.execute("SELECT pg_advisory_xact_lock(742013)")

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS storage_units (
            id BIGSERIAL PRIMARY KEY,
            storage_name TEXT NOT NULL,
            notes TEXT,
            created_at TEXT,
            updated_at TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS ux_storage_units_name_lower
        ON storage_units (LOWER(storage_name))
        """
    )
    # Storage units such as a disposal/throw list can be kept out of the
    # default Chemical Search list.
    conn.execute(
        "ALTER TABLE storage_units ADD COLUMN IF NOT EXISTS hidden_from_lists BOOLEAN NOT NULL DEFAULT FALSE"
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS storage_shelves (
            id BIGSERIAL PRIMARY KEY,
            storage_unit_id BIGINT NOT NULL REFERENCES storage_units(id),
            shelf_number INTEGER NOT NULL,
            UNIQUE(storage_unit_id, shelf_number)
        )
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS chemicals (
            id BIGSERIAL PRIMARY KEY,
            chemical_name TEXT NOT NULL,
            cas_number TEXT,
            manufacturer TEXT,
            catalog_number TEXT,
            amount DOUBLE PRECISION,
            unit TEXT,
            location TEXT,
            status TEXT,
            created_at TEXT,
            bottle_id TEXT,
            legacy_id TEXT,
            company_management_id TEXT,
            category_classification TEXT,
            category_name TEXT,
            description TEXT,
            lot_number TEXT,
            initial_amount DOUBLE PRECISION,
            remaining_amount DOUBLE PRECISION,
            safety_stock TEXT,
            concentration_density TEXT,
            molecular_weight TEXT,
            mol TEXT,
            purity TEXT,
            solubility TEXT,
            molecular_formula TEXT,
            storage_temperature TEXT,
            flash_point TEXT,
            boiling_point TEXT,
            melting_point TEXT,
            owner TEXT,
            registered_by TEXT,
            purchaser TEXT,
            purchase_date TEXT,
            purchase_price TEXT,
            source_registered_date TEXT,
            source_last_modified TEXT,
            opened_date TEXT,
            expiration_date TEXT,
            comments TEXT,
            tags TEXT,
            legacy_location_path TEXT,
            legacy_storage_name TEXT,
            storage_unit_id BIGINT REFERENCES storage_units(id),
            shelf_number INTEGER,
            updated_at TEXT,
            disposed_at TEXT,
            disposal_reason TEXT,
            disposal_note TEXT,
            disposal_location TEXT,
            pubchem_cid TEXT,
            safety_signal_word TEXT,
            safety_h_codes TEXT,
            safety_hazard_statements TEXT,
            safety_pictograms TEXT,
            safety_source_url TEXT,
            safety_checked_at TEXT,
            safety_match_term TEXT
        )
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS sds_documents (
            id BIGSERIAL PRIMARY KEY,
            chemical_id BIGINT NOT NULL UNIQUE REFERENCES chemicals(id),
            file_name TEXT,
            file_data BYTEA,
            source_name TEXT,
            source_url TEXT,
            revision_date TEXT,
            uploaded_at TEXT
        )
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS inventory_observations (
            id BIGSERIAL PRIMARY KEY,
            chemical_id BIGINT NOT NULL REFERENCES chemicals(id) ON DELETE CASCADE,
            original_file_name TEXT,
            original_image_data BYTEA,
            label_image_data BYTEA,
            ocr_result_json TEXT,
            observed_storage_unit_id BIGINT REFERENCES storage_units(id),
            observed_shelf_number INTEGER,
            captured_at TEXT NOT NULL,
            uploaded_by TEXT
        )
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_inventory_observations_chemical
        ON inventory_observations(chemical_id, id DESC)
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS app_users (
            id BIGSERIAL PRIMARY KEY,
            email TEXT NOT NULL UNIQUE,
            display_name TEXT,
            role TEXT NOT NULL DEFAULT 'member',
            active BOOLEAN NOT NULL DEFAULT TRUE,
            created_at TEXT,
            updated_at TEXT,
            created_by TEXT
        )
        """
    )

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS audit_logs (
            id BIGSERIAL PRIMARY KEY,
            user_email TEXT NOT NULL,
            user_name TEXT,
            action TEXT NOT NULL,
            target_type TEXT,
            target_id BIGINT,
            bottle_id TEXT,
            details TEXT,
            created_at TEXT NOT NULL
        )
        """
    )

    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_app_users_email ON app_users(LOWER(email))"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_audit_created_at ON audit_logs(created_at)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_audit_user_email ON audit_logs(LOWER(user_email))"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_audit_action ON audit_logs(action)"
    )

    # Keep status and amount fields normalized without changing user data.
    status_mapping = {
        "미개봉": "Unopened",
        "사용 중": "In Use",
        "소진": "Disposal Pending",
        "폐기 예정": "Disposal Pending",
        "폐기": "Disposed",
        "Empty": "Disposal Pending",
    }
    for old_value, new_value in status_mapping.items():
        conn.execute(
            "UPDATE chemicals SET status=? WHERE status=?",
            (new_value, old_value),
        )

    conn.execute(
        "UPDATE chemicals SET status='Unopened' "
        "WHERE status IS NULL OR TRIM(status)=''"
    )
    conn.execute(
        """
        UPDATE chemicals
        SET initial_amount = amount
        WHERE initial_amount IS NULL AND amount IS NOT NULL
        """
    )
    conn.execute(
        """
        UPDATE chemicals
        SET remaining_amount = amount
        WHERE remaining_amount IS NULL AND amount IS NOT NULL
        """
    )
    conn.execute(
        """
        UPDATE chemicals
        SET bottle_id = 'ICL-' || LPAD(id::text, 6, '0')
        WHERE bottle_id IS NULL OR TRIM(bottle_id)=''
        """
    )

    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chem_name ON chemicals(chemical_name)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chem_cas ON chemicals(cas_number)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chem_status ON chemicals(status)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chem_legacy ON chemicals(legacy_id)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chem_storage ON chemicals(storage_unit_id)"
    )

    conn.commit()
    conn.close()
    return True


init_db()


# ============================================================
# HELPERS
# ============================================================

LOCATION_EXPR = """
CASE
    WHEN su.storage_name IS NOT NULL AND c.shelf_number IS NOT NULL
        THEN su.storage_name || ' / Shelf ' || c.shelf_number
    WHEN su.storage_name IS NOT NULL
        THEN su.storage_name
    ELSE COALESCE(NULLIF(c.location, ''), 'Not Assigned')
END
"""


def clean_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip()
    if text.lower() == "nan" or text == "-":
        return ""
    return text


def clean_number(value):
    text = clean_text(value)
    if not text:
        return None

    try:
        return float(text.replace(",", ""))
    except Exception:
        return None


def clean_date(value):
    if value is None:
        return ""

    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")

    text = clean_text(value)
    if not text:
        return ""

    parsed = pd.to_datetime(text, errors="coerce")
    if not pd.isna(parsed):
        return parsed.strftime("%Y-%m-%d")

    return text


def normalize_unit(value):
    text = clean_text(value)
    mapping = {
        "ml": "mL",
        "milliliter": "mL",
        "milliliters": "mL",
        "l": "L",
        "liter": "L",
        "liters": "L",
        "g": "g",
        "mg": "mg",
        "kg": "kg",
        "ea": "ea",
    }
    return mapping.get(text.lower(), text)


@st.cache_data(ttl=5, show_spinner=False)
def get_storage_options():
    conn = get_connection()
    units = conn.execute(
        "SELECT id, storage_name FROM storage_units ORDER BY storage_name"
    ).fetchall()
    shelves = conn.execute(
        """
        SELECT storage_unit_id, shelf_number
        FROM storage_shelves
        ORDER BY storage_unit_id, shelf_number
        """
    ).fetchall()
    conn.close()

    shelves_by_unit = {}
    for shelf in shelves:
        shelves_by_unit.setdefault(shelf["storage_unit_id"], []).append(
            shelf["shelf_number"]
        )

    options = [(None, None, "Not Assigned")]

    for unit in units:
        # Storage unit without a specific shelf
        options.append((unit["id"], None, unit["storage_name"]))

        for shelf_number in shelves_by_unit.get(unit["id"], []):
            options.append(
                (
                    unit["id"],
                    shelf_number,
                    f"{unit['storage_name']} / Shelf {shelf_number}",
                )
            )

    return options


def get_or_create_storage(conn, storage_name):
    storage_name = clean_text(storage_name)
    if not storage_name:
        return None

    row = conn.execute(
        "SELECT id FROM storage_units WHERE LOWER(storage_name)=LOWER(?)",
        (storage_name,),
    ).fetchone()

    if row:
        return row["id"]

    cursor = conn.execute(
        """
        INSERT INTO storage_units
        (storage_name, notes, created_at)
        VALUES (?, ?, ?)
        """,
        (
            storage_name,
            "Imported from Lab Manager Excel",
            datetime.now().strftime("%Y-%m-%d %H:%M"),
        ),
    )
    return cursor.lastrowid


def current_location_index(options, storage_unit_id, shelf_number):
    for index, option in enumerate(options):
        if option[0] == storage_unit_id and option[1] == shelf_number:
            return index
    return 0


def set_notice(message):
    st.session_state["notice"] = message


# ============================================================
# AUTHENTICATION / USERS / AUDIT
# ============================================================

def _secret_text(name):
    """Read a top-level secret from env or Streamlit Secrets."""
    env_value = os.getenv(name, "").strip()
    if env_value:
        return env_value

    try:
        return str(st.secrets.get(name, "")).strip()
    except Exception:
        return ""


def write_audit_log(
    action,
    target_type=None,
    target_id=None,
    bottle_id=None,
    details=None,
    conn=None,
    user=None,
):
    """Write one user activity record. Reuse an existing transaction when supplied."""
    actor = user or globals().get("CURRENT_USER") or {}
    user_email = str(actor.get("email") or "system").strip().lower()
    user_name = str(actor.get("name") or actor.get("display_name") or user_email).strip()

    owns_connection = conn is None
    if owns_connection:
        conn = get_connection()

    try:
        conn.execute(
            """
            INSERT INTO audit_logs (
                user_email, user_name, action, target_type,
                target_id, bottle_id, details, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user_email,
                user_name,
                action,
                target_type,
                target_id,
                bottle_id,
                details,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ),
        )
        if owns_connection:
            conn.commit()
    finally:
        if owns_connection:
            conn.close()


def _get_app_user(email):
    email = str(email or "").strip().lower()
    if not email:
        return None

    conn = get_connection()
    row = conn.execute(
        """
        SELECT id, email, display_name, role, active
        FROM app_users
        WHERE LOWER(email)=LOWER(?)
        """,
        (email,),
    ).fetchone()
    conn.close()
    return row


def _bootstrap_admin_if_allowed(email, identity_name):
    """
    Create the first/declared administrator only when the signed-in email
    exactly matches BOOTSTRAP_ADMIN_EMAIL from Streamlit Secrets.
    """
    bootstrap_email = _secret_text("BOOTSTRAP_ADMIN_EMAIL").lower()
    email = str(email or "").strip().lower()

    if not bootstrap_email or email != bootstrap_email:
        return None

    existing = _get_app_user(email)
    if existing:
        return existing

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    conn.execute(
        """
        INSERT INTO app_users (
            email, display_name, role, active,
            created_at, updated_at, created_by
        )
        VALUES (?, ?, 'admin', TRUE, ?, ?, ?)
        ON CONFLICT(email) DO NOTHING
        """,
        (
            email,
            identity_name or email,
            now,
            now,
            "BOOTSTRAP_ADMIN_EMAIL",
        ),
    )
    conn.commit()
    conn.close()
    return _get_app_user(email)


def _login_screen():
    st.markdown(
        """
<div style="max-width:620px; margin:7rem auto 0 auto; background:#FFFFFF;
padding:36px; border:1px solid #E5E7EB; border-radius:18px; text-align:center;">
<div style="font-size:32px; font-weight:700; color:#172033;">🧪 ICL Chemical Manager</div>
<div style="margin-top:8px; color:#64748B;">
Laboratory Chemical Inventory Management System
</div>
</div>
        """,
        unsafe_allow_html=True,
    )
    st.write("")
    left, middle, right = st.columns([1, 1.3, 1])
    with middle:
        st.button(
            "Sign in with Google",
            type="primary",
            width="stretch",
            on_click=st.login,
        )


def require_authenticated_user():
    """Require Google OIDC login and an active app_users whitelist entry."""
    try:
        logged_in = bool(st.user.is_logged_in)
    except Exception:
        st.error(
            "Authentication is not configured yet. Add the [auth] Google OIDC "
            "settings to this app's Streamlit Secrets."
        )
        st.stop()

    if not logged_in:
        _login_screen()
        st.stop()

    identity = st.user.to_dict()
    email = str(identity.get("email") or "").strip().lower()
    identity_name = str(
        identity.get("name")
        or identity.get("given_name")
        or email
    ).strip()

    if not email:
        st.error("Your Google account did not provide an email address.")
        if st.button("Log out"):
            st.logout()
        st.stop()

    # Every click reruns the script; re-check the account at most once a minute
    # instead of on every click. Deactivation still takes effect within 60 s.
    cached = st.session_state.get("_icl_app_user_cache")
    if cached and cached[0] == email and time.time() - cached[1] < 60:
        app_user = cached[2]
    else:
        app_user = _get_app_user(email)
        if app_user is not None:
            st.session_state["_icl_app_user_cache"] = (email, time.time(), app_user)
    if app_user is None:
        app_user = _bootstrap_admin_if_allowed(email, identity_name)

    if app_user is None or not bool(app_user["active"]):
        st.error("⛔ This Google account is not authorized to use ICL Chemical Manager.")
        st.caption(f"Signed in as: {email}")
        if not _secret_text("BOOTSTRAP_ADMIN_EMAIL"):
            st.info(
                "Initial setup: add BOOTSTRAP_ADMIN_EMAIL to Streamlit Secrets "
                "using the Google email that should become the first administrator."
            )
        st.button("Log out", on_click=st.logout)
        st.stop()

    current_user = {
        "id": app_user["id"],
        "email": str(app_user["email"]).strip().lower(),
        "name": str(app_user["display_name"] or identity_name or email).strip(),
        "role": str(app_user["role"] or "member").strip().lower(),
    }

    if not st.session_state.get("_login_audit_written"):
        write_audit_log(
            "LOGIN",
            target_type="session",
            details="Google OIDC login",
            user=current_user,
        )
        st.session_state["_login_audit_written"] = True

    return current_user


CURRENT_USER = require_authenticated_user()


# ============================================================
# SAFETY / PUBCHEM HELPERS
# ============================================================

PUBCHEM_BASE = "https://pubchem.ncbi.nlm.nih.gov"
PUBCHEM_USER_AGENT = "ICL-Chemical-Manager/1.0"


def http_get_json(url, timeout=15):
    """Small JSON HTTP helper using only the Python standard library."""
    request = Request(
        url,
        headers={
            "User-Agent": PUBCHEM_USER_AGENT,
            "Accept": "application/json",
        },
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError):
        return None


def iter_pubchem_information(node):
    """Yield all PubChem PUG-View Information objects recursively."""
    if isinstance(node, dict):
        information = node.get("Information")
        if isinstance(information, list):
            for item in information:
                if isinstance(item, dict):
                    yield item

        for value in node.values():
            yield from iter_pubchem_information(value)

    elif isinstance(node, list):
        for item in node:
            yield from iter_pubchem_information(item)


def pubchem_string_values(info):
    values = []
    value = info.get("Value") or {}

    for item in value.get("StringWithMarkup", []) or []:
        if not isinstance(item, dict):
            continue
        text = str(item.get("String") or "").strip()
        if text:
            values.append(text)

    return values


def pubchem_pictogram_names(info):
    names = []

    for item in (info.get("Value") or {}).get("StringWithMarkup", []) or []:
        if not isinstance(item, dict):
            continue

        text = str(item.get("String") or "").strip()
        if text and text.lower() not in {"pictogram", "pictograms"}:
            names.append(text)

        for markup in item.get("Markup", []) or []:
            if not isinstance(markup, dict):
                continue
            extra = str(markup.get("Extra") or "").strip()
            if extra:
                names.append(extra)

    # Keep order while removing duplicates.
    return list(dict.fromkeys(names))


def find_pubchem_cid(cas_number, chemical_name):
    terms = []

    cas_number = clean_text(cas_number)
    chemical_name = clean_text(chemical_name)

    if cas_number:
        terms.append(cas_number)
    if chemical_name and chemical_name not in terms:
        terms.append(chemical_name)

    for term in terms:
        url = (
            f"{PUBCHEM_BASE}/rest/pug/compound/name/"
            f"{quote(term, safe='')}/cids/JSON"
        )
        data = http_get_json(url)
        if not data:
            continue

        cids = (data.get("IdentifierList") or {}).get("CID") or []
        if cids:
            return str(cids[0]), term

    return None, None


def fetch_pubchem_safety(cas_number, chemical_name):
    """
    Retrieve GHS screening information from PubChem.

    Important: this is NOT treated as the manufacturer's official SDS.
    """
    cid, matched_term = find_pubchem_cid(cas_number, chemical_name)

    if not cid:
        return None

    url = (
        f"{PUBCHEM_BASE}/rest/pug_view/data/compound/{cid}/JSON/"
        f"?heading={quote('GHS Classification', safe='')}"
    )
    data = http_get_json(url)

    if not data:
        return None

    signal_words = []
    hazard_statements = []
    pictograms = []

    for info in iter_pubchem_information(data):
        name = str(info.get("Name") or "").strip().lower()

        if name in {"signal", "signal word"}:
            signal_words.extend(pubchem_string_values(info))

        elif "ghs hazard statements" in name or name == "hazard statements":
            hazard_statements.extend(pubchem_string_values(info))

        elif "pictogram" in name:
            pictograms.extend(pubchem_pictogram_names(info))

    signal_words = list(dict.fromkeys(x for x in signal_words if x))
    hazard_statements = list(dict.fromkeys(x for x in hazard_statements if x))
    pictograms = list(dict.fromkeys(x for x in pictograms if x))

    h_codes = []
    for statement in hazard_statements:
        for code in re.findall(r"\b(?:H\d{3}[A-Z]?|EUH\d{3}[A-Z]?)\b", statement):
            if code not in h_codes:
                h_codes.append(code)

    # If multiple contributors exist, Danger takes precedence for the display.
    normalized_signals = [x.strip().title() for x in signal_words]
    if any(x.lower() == "danger" for x in normalized_signals):
        signal_word = "Danger"
    elif any(x.lower() == "warning" for x in normalized_signals):
        signal_word = "Warning"
    else:
        signal_word = normalized_signals[0] if normalized_signals else ""

    return {
        "cid": cid,
        "match_term": matched_term or "",
        "signal_word": signal_word,
        "h_codes": ", ".join(h_codes),
        "hazard_statements": "\n".join(hazard_statements),
        "pictograms": ", ".join(pictograms),
        "source_url": f"{PUBCHEM_BASE}/compound/{cid}",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def save_pubchem_safety(chemical_id, safety):
    conn = get_connection()
    conn.execute(
        """
        UPDATE chemicals
        SET pubchem_cid=?,
            safety_signal_word=?,
            safety_h_codes=?,
            safety_hazard_statements=?,
            safety_pictograms=?,
            safety_source_url=?,
            safety_checked_at=?,
            safety_match_term=?,
            updated_at=?
        WHERE id=?
        """,
        (
            safety.get("cid", ""),
            safety.get("signal_word", ""),
            safety.get("h_codes", ""),
            safety.get("hazard_statements", ""),
            safety.get("pictograms", ""),
            safety.get("source_url", ""),
            safety.get("checked_at", ""),
            safety.get("match_term", ""),
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            chemical_id,
        ),
    )
    write_audit_log(
        "SAFETY_SCREENING",
        target_type="chemical",
        target_id=chemical_id,
        details=f"PubChem CID: {safety.get('cid', '')}",
        conn=conn,
    )
    conn.commit()
    conn.close()


def safety_badge(signal_word, hazard_statements, checked_at):
    signal = clean_text(signal_word).lower()

    if signal == "danger":
        return "🔴 Danger"
    if signal == "warning":
        return "🟠 Warning"
    if clean_text(hazard_statements):
        return "🟡 Hazard info"
    if clean_text(checked_at):
        return "⚪ Check SDS"
    return "⚪ Not checked"


def render_safety_panel(
    signal_word="",
    h_codes="",
    hazard_statements="",
    pictograms="",
    source_url="",
    checked_at="",
    source_label="PubChem safety screening",
):
    """Render safety screening without presenting it as an official SDS."""
    signal = clean_text(signal_word)
    h_codes = clean_text(h_codes)
    pictograms = clean_text(pictograms)
    source_url = clean_text(source_url)
    checked_at = clean_text(checked_at)
    statements = [
        line.strip()
        for line in str(hazard_statements or "").splitlines()
        if line.strip()
    ]

    if not checked_at:
        st.info("Safety screening has not been checked yet. Verify the official SDS before use.")
        return

    if signal.lower() == "danger":
        st.error("⚠️ DANGER — hazardous chemical. Review the official SDS before handling.")
    elif signal.lower() == "warning":
        st.warning("⚠️ WARNING — hazardous chemical. Review the official SDS before handling.")
    elif statements or pictograms:
        st.warning("⚠️ Hazard information is available. Review the official SDS before handling.")
    else:
        st.info(
            "No GHS hazard statement was returned by this PubChem screening. "
            "This does NOT mean the product is safe; verify the official SDS."
        )

    info1, info2, info3 = st.columns(3)
    info1.markdown(f"**Signal Word**\n\n{signal or '-'}")
    info2.markdown(f"**H-Codes**\n\n{h_codes or '-'}")
    info3.markdown(f"**Pictograms / Hazards**\n\n{pictograms or '-'}")

    if statements:
        with st.expander("Hazard Statements", expanded=signal.lower() == "danger"):
            for statement in statements[:20]:
                st.write(f"• {statement}")
            if len(statements) > 20:
                st.caption(f"Showing 20 of {len(statements)} statements.")

    footer = f"Source: {source_label}"
    if checked_at:
        footer += f"  |  Checked: {checked_at}"
    st.caption(footer)

    if source_url:
        st.link_button("Open PubChem Safety Source", source_url)


# ============================================================
# DIALOGS
# ============================================================

@st.dialog("✏️ Edit Chemical", width="large")
def edit_chemical_dialog(chemical_id):
    conn = get_connection()
    chemical = conn.execute(
        "SELECT * FROM chemicals WHERE id=?",
        (chemical_id,),
    ).fetchone()
    conn.close()

    if not chemical:
        st.error("Chemical not found.")
        return

    st.markdown(f"### {chemical['chemical_name']}")
    st.caption(
        f"Bottle ID: {chemical['bottle_id'] or '-'}   |   "
        f"Legacy ID: {chemical['legacy_id'] or '-'}"
    )

    storage_options = get_storage_options()
    storage_index = current_location_index(
        storage_options,
        chemical["storage_unit_id"],
        chemical["shelf_number"],
    )

    unit_options = ["g", "mg", "kg", "mL", "L", "ea"]
    current_unit = chemical["unit"] or "g"
    if current_unit not in unit_options:
        unit_options.append(current_unit)

    status_options = ["Unopened", "In Use", "Disposal Pending"]
    current_status = chemical["status"] or "Unopened"
    if current_status not in status_options:
        status_options.append(current_status)

    with st.form(f"edit_chemical_form_{chemical_id}", enter_to_submit=False):
        edit_name = st.text_input(
            "Chemical Name *",
            value=chemical["chemical_name"] or "",
            key=f"edit_name_{chemical_id}",
        )
        edit_description = st.text_input(
            "Description / Grade",
            value=chemical["description"] or "",
            key=f"edit_description_{chemical_id}",
        )

        col1, col2 = st.columns(2)

        with col1:
            edit_cas = st.text_input(
                "CAS No.",
                value=chemical["cas_number"] or "",
                key=f"edit_cas_{chemical_id}",
            )
            edit_manufacturer = st.text_input(
                "Manufacturer",
                value=chemical["manufacturer"] or "",
                key=f"edit_manufacturer_{chemical_id}",
            )
            edit_catalog = st.text_input(
                "Catalog No.",
                value=chemical["catalog_number"] or "",
                key=f"edit_catalog_{chemical_id}",
            )
            edit_lot = st.text_input(
                "Lot No.",
                value=chemical["lot_number"] or "",
                key=f"edit_lot_{chemical_id}",
            )
            edit_initial = st.number_input(
                "Initial Amount",
                min_value=0.0,
                value=float(
                    chemical["initial_amount"]
                    if chemical["initial_amount"] is not None
                    else chemical["amount"] or 0
                ),
                key=f"edit_initial_{chemical_id}",
            )

        with col2:
            edit_remaining = st.number_input(
                "Remaining Amount",
                min_value=0.0,
                value=float(
                    chemical["remaining_amount"]
                    if chemical["remaining_amount"] is not None
                    else chemical["amount"] or 0
                ),
                key=f"edit_remaining_{chemical_id}",
            )
            edit_unit = st.selectbox(
                "Unit",
                unit_options,
                index=unit_options.index(current_unit),
                key=f"edit_unit_{chemical_id}",
            )
            edit_storage = st.selectbox(
                "Storage Location",
                storage_options,
                index=storage_index,
                format_func=lambda option: option[2],
                key=f"edit_storage_{chemical_id}",
            )
            edit_owner = st.text_input(
                "Owner / Person in Charge",
                value=chemical["owner"] or "",
                key=f"edit_owner_{chemical_id}",
            )
            edit_expiration = st.text_input(
                "Expiration Date",
                value=chemical["expiration_date"] or "",
                placeholder="YYYY-MM-DD",
                key=f"edit_expiration_{chemical_id}",
            )

        edit_status = st.selectbox(
            "Status",
            status_options,
            index=status_options.index(current_status),
            key=f"edit_status_{chemical_id}",
        )

        save_col, cancel_col = st.columns(2)
        with save_col:
            save_edit = st.form_submit_button(
                "Save Changes",
                type="primary",
                width="stretch",
            )
        with cancel_col:
            cancel_edit = st.form_submit_button(
                "Cancel",
                width="stretch",
            )

    if save_edit:
        if not edit_name.strip():
            st.error("Chemical name is required.")
        else:
            conn = get_connection()
            conn.execute(
                """
                UPDATE chemicals
                SET chemical_name=?,
                    description=?,
                    cas_number=?,
                    manufacturer=?,
                    catalog_number=?,
                    lot_number=?,
                    amount=?,
                    initial_amount=?,
                    remaining_amount=?,
                    unit=?,
                    storage_unit_id=?,
                    shelf_number=?,
                    location=?,
                    owner=?,
                    expiration_date=?,
                    status=?,
                    updated_at=?
                WHERE id=?
                """,
                (
                    edit_name.strip(),
                    edit_description.strip(),
                    edit_cas.strip(),
                    edit_manufacturer.strip(),
                    edit_catalog.strip(),
                    edit_lot.strip(),
                    edit_initial,
                    edit_initial,
                    edit_remaining,
                    edit_unit,
                    edit_storage[0],
                    edit_storage[1],
                    edit_storage[2] if edit_storage[0] is not None else "",
                    edit_owner.strip(),
                    edit_expiration.strip(),
                    edit_status,
                    datetime.now().strftime("%Y-%m-%d %H:%M"),
                    chemical_id,
                ),
            )
            write_audit_log(
                "EDIT_CHEMICAL",
                target_type="chemical",
                target_id=chemical_id,
                bottle_id=chemical["bottle_id"],
                details=json.dumps(
                    {
                        "chemical_name": edit_name.strip(),
                        "remaining_amount": edit_remaining,
                        "unit": edit_unit,
                        "location": (
                            edit_storage[2]
                            if edit_storage[0] is not None
                            else "Not Assigned"
                        ),
                        "owner": edit_owner.strip(),
                        "status": edit_status,
                    },
                    ensure_ascii=False,
                ),
                conn=conn,
            )
            conn.commit()
            conn.close()
            set_notice("✅ Chemical information has been updated.")
            st.rerun()

    if cancel_edit:
        st.rerun()

@st.dialog("🗑️ Dispose Chemical", width="medium")
def dispose_chemical_dialog(chemical_id):
    conn = get_connection()
    chemical = conn.execute(
        f"""
        SELECT
            c.id,
            c.bottle_id,
            c.legacy_id,
            c.chemical_name,
            c.cas_number,
            c.manufacturer,
            COALESCE(c.remaining_amount, c.amount) AS remaining_amount,
            c.unit,
            {LOCATION_EXPR} AS location_text
        FROM chemicals c
        LEFT JOIN storage_units su ON c.storage_unit_id = su.id
        WHERE c.id = ?
        """,
        (chemical_id,),
    ).fetchone()
    conn.close()

    if not chemical:
        st.error("Chemical not found.")
        return

    st.markdown(f"### {chemical['chemical_name']}")

    id_parts = []
    if chemical["bottle_id"]:
        id_parts.append(f"Bottle ID: {chemical['bottle_id']}")
    if chemical["legacy_id"]:
        id_parts.append(f"Legacy ID: {chemical['legacy_id']}")
    if id_parts:
        st.caption("   |   ".join(id_parts))

    info_col1, info_col2 = st.columns(2)

    with info_col1:
        st.markdown("**Manufacturer**")
        st.write(chemical["manufacturer"] or "-")

        st.markdown("**CAS No.**")
        st.write(chemical["cas_number"] or "-")

    with info_col2:
        st.markdown("**Location**")
        st.write(chemical["location_text"] or "Not Assigned")

        st.markdown("**Remaining**")
        if chemical["remaining_amount"] is not None:
            st.write(f"{chemical['remaining_amount']:g} {chemical['unit'] or ''}")
        else:
            st.write("-")

    st.divider()

    disposal_reason = st.selectbox(
        "Reason for Disposal *",
        [
            "Empty / Fully Used",
            "Expired",
            "Damaged Container",
            "Contaminated",
            "Quality Concern",
            "No Longer Needed",
            "Duplicate / Excess Stock",
            "Other",
        ],
        key=f"dialog_disposal_reason_{chemical_id}",
    )

    disposal_note = st.text_area(
        "Additional Note",
        placeholder="Optional details about the disposal",
        key=f"dialog_disposal_note_{chemical_id}",
    )

    confirm = st.checkbox(
        "I confirm that this chemical has been disposed.",
        key=f"dialog_disposal_confirm_{chemical_id}",
    )

    cancel_col, confirm_col = st.columns(2)

    with cancel_col:
        if st.button(
            "Cancel",
            key=f"dialog_cancel_disposal_{chemical_id}",
            width="stretch",
        ):
            st.rerun()

    with confirm_col:
        if st.button(
            "Confirm Disposal",
            type="primary",
            key=f"dialog_confirm_disposal_{chemical_id}",
            width="stretch",
        ):
            if not confirm:
                st.error("Please check the confirmation box.")
            else:
                now = datetime.now().strftime("%Y-%m-%d %H:%M")

                conn = get_connection()
                conn.execute(
                    """
                    UPDATE chemicals
                    SET status='Disposed',
                        disposed_at=?,
                        disposal_reason=?,
                        disposal_note=?,
                        disposal_location=?,
                        updated_at=?
                    WHERE id=?
                    """,
                    (
                        now,
                        disposal_reason,
                        disposal_note.strip(),
                        chemical["location_text"],
                        now,
                        chemical_id,
                    ),
                )
                write_audit_log(
                    "DISPOSE_CHEMICAL",
                    target_type="chemical",
                    target_id=chemical_id,
                    bottle_id=chemical["bottle_id"],
                    details=json.dumps(
                        {
                            "chemical_name": chemical["chemical_name"],
                            "reason": disposal_reason,
                            "note": disposal_note.strip(),
                            "location": chemical["location_text"],
                        },
                        ensure_ascii=False,
                    ),
                    conn=conn,
                )
                conn.commit()
                conn.close()

                set_notice(
                    f"✅ {chemical['chemical_name']} has been disposed."
                )
                st.rerun()


# ============================================================
# OCR WORKER / INVENTORY OBSERVATIONS
# ============================================================


def _ocr_secret(name):
    """
    App-wide OCR backend configuration.

    Production policy:
    - OCR_API_URL and OCR_API_KEY live in Streamlit Secrets (or env vars).
    - They are shared by every logged-in user of this Streamlit deployment.
    - End users never enter backend credentials.
    """
    env_value = os.getenv(name, "").strip()
    if env_value:
        return env_value

    try:
        return str(st.secrets.get(name, "")).strip()
    except Exception:
        return ""


def get_ocr_api_url():
    # The old OCR_API_URL may still point to v7. v8 uses this known deployment.
    return (_ocr_secret("OCR_V8_API_URL") or "https://eddiekim995--icl-reagent-ocr-v8-ocrservice-web.modal.run").rstrip("/")



def get_ocr_api_key():
    if _ocr_user_is_admin():
        temporary = str(st.session_state.get("icl_v8_temporary_api_key", "")).strip()
        if temporary:
            return temporary
    return _ocr_secret("OCR_API_KEY")



def ocr_backend_configured():
    return bool(
        get_ocr_api_url()
        and get_ocr_api_key()
    )


def _ocr_user_is_admin():
    current = globals().get("CURRENT_USER") or {}
    return str(
        current.get("role") or ""
    ).strip().lower() == "admin"


def _ocr_backend_summary(info):
    if not isinstance(info, dict):
        return ""

    parts = []

    backend = clean_text(info.get("backend"))
    if backend:
        parts.append(backend)

    device = clean_text(
        info.get("device")
        or info.get("gpu")
    )
    if device:
        parts.append(device)

    engine_patch = clean_text(
        info.get("engine_patch")
    )
    if engine_patch:
        parts.append(engine_patch)

    inventory_count = info.get(
        "inventory_count"
    )
    if inventory_count is not None:
        try:
            parts.append(
                f"{int(inventory_count):,} identities"
            )
        except Exception:
            pass

    return " · ".join(parts)


def render_ocr_backend_panel(
    key_prefix="ocr",
    expanded=False,
):
    """
    Compact Modal GPU status UI.

    /health is intentionally NOT called on every page render because a health
    request would wake a scale-to-zero Modal GPU just to draw the page.
    The normal inventory/OCR request verifies the backend when it is needed.
    Admins can explicitly run diagnostics.
    """
    panel_col, _panel_spacer = st.columns(
        [1.2, 2.0]
    )

    with panel_col:
        if not ocr_backend_configured():
            st.error(
                "☁️ Modal GPU OCR is not configured."
            )
            st.caption(
                "Admin: add OCR_API_URL and OCR_API_KEY "
                "to Streamlit Secrets."
            )
            return

        last_info = (
            st.session_state.get(
                "ocr_last_server_info"
            )
            or st.session_state.get(
                "ocr_last_health"
            )
            or {}
        )

        st.success(
            "☁️ Modal GPU OCR ready"
        )

        summary = _ocr_backend_summary(
            last_info
        )

        if summary:
            st.caption(summary)
        else:
            st.caption(
                "Serverless GPU backend · "
                "starts automatically when OCR is requested"
            )

        if _ocr_user_is_admin():
            with st.expander(
                "Admin diagnostics",
                expanded=expanded,
            ):
                endpoint = get_ocr_api_url()

                st.caption(
                    f"Endpoint: {endpoint}"
                )
                st.caption(
                    "API key: configured in Streamlit Secrets"
                )

                if st.button(
                    "Check Modal backend",
                    key=f"{key_prefix}_admin_health",
                    width="stretch",
                ):
                    try:
                        response = requests.get(
                            endpoint + "/health",
                            timeout=120,
                        )
                        response.raise_for_status()
                        info = response.json()

                    except Exception as exc:
                        st.session_state.pop(
                            "ocr_last_health",
                            None,
                        )
                        st.error(
                            f"Modal backend check failed: {exc}"
                        )

                    else:
                        st.session_state[
                            "ocr_last_health"
                        ] = info
                        st.success(
                            "Modal backend online"
                        )
                        st.json(info)


def _current_inventory_for_ocr():
    conn = get_connection()
    rows = conn.execute(
        f"""
        SELECT
            c.chemical_name AS "Chemical Name",
            c.description AS "Description",
            c.cas_number AS "CAS No.",
            c.manufacturer AS "Manufacturer",
            c.catalog_number AS "Catalog No.",
            COALESCE(c.remaining_amount, c.amount) AS "Amount",
            c.unit AS "Unit",
            c.molecular_formula AS "Molecular Formula",
            c.molecular_weight AS "Molecular Weight",
            {LOCATION_EXPR} AS "Storage"
        FROM chemicals c
        LEFT JOIN storage_units su ON c.storage_unit_id = su.id
        WHERE c.status != 'Disposed'
        ORDER BY c.id
        """
    ).fetchall()
    conn.close()
    return pd.DataFrame([dict(row) for row in rows])


def sync_inventory_to_ocr_worker(force=False):
    url = get_ocr_api_url()
    key = get_ocr_api_key()
    if not url or not key:
        raise RuntimeError("Modal OCR is not configured. Admin must set OCR_API_URL and OCR_API_KEY in Streamlit Secrets.")

    df = _current_inventory_for_ocr()
    csv_bytes = df.to_csv(index=False).encode("utf-8-sig")
    signature = hashlib.sha1(csv_bytes).hexdigest()
    cache_key = f"{url}::{signature}"

    if (
        not force
        and st.session_state.get("ocr_inventory_sync_signature") == cache_key
    ):
        return {
            "ok": True,
            "cached": True,
            "inventory_count": st.session_state.get("ocr_inventory_sync_count", len(df)),
        }

    response = requests.post(
        url + "/inventory",
        headers={"X-API-Key": key},
        files={
            "file": (
                "chemical_manager_live_inventory.csv",
                csv_bytes,
                "text/csv",
            )
        },
        timeout=120,
    )
    if response.status_code != 200:
        raise RuntimeError(f"Inventory sync {response.status_code}: {response.text}")

    payload = response.json()
    st.session_state["ocr_last_inventory_sync"] = payload

    if isinstance(payload, dict):
        previous_server_info = (
            st.session_state.get(
                "ocr_last_server_info"
            )
            or {}
        )
        st.session_state[
            "ocr_last_server_info"
        ] = {
            **previous_server_info,
            "backend": payload.get(
                "backend",
                previous_server_info.get(
                    "backend",
                    "Modal",
                ),
            ),
            "inventory_count": payload.get(
                "inventory_count",
                previous_server_info.get(
                    "inventory_count"
                ),
            ),
        }

    st.session_state["ocr_inventory_sync_signature"] = cache_key
    st.session_state["ocr_inventory_sync_count"] = payload.get(
        "inventory_count", len(df)
    )
    return payload


def _ocr_api_request_bytes(
    filename,
    data,
    mime_type,
    enable_pubchem=True,
    extract_only=False,
):
    """
    Thread-safe OCR HTTP request.

    This helper does not touch Streamlit session_state, so it can safely run
    inside the bulk ThreadPoolExecutor.
    """
    url = get_ocr_api_url()
    key = get_ocr_api_key()

    if not url or not key:
        raise RuntimeError(
            "Modal OCR is not configured. Admin must set "
            "OCR_API_URL and OCR_API_KEY in Streamlit Secrets."
        )

    response = requests.post(
        url + "/ocr",
        headers={
            "X-API-Key": key
        },
        files={
            "file": (
                filename,
                data,
                mime_type
                or "application/octet-stream",
            )
        },
        data={
            "enable_pubchem":
                str(
                    bool(
                        enable_pubchem
                    )
                ).lower(),
            "extract_only":
                str(
                    bool(
                        extract_only
                    )
                ).lower(),
        },
        timeout=60,
    )

    if response.status_code != 200:
        raise RuntimeError(
            f"OCR API {response.status_code}: "
            f"{response.text}"
        )

    return response.json()


def analyze_label_with_ocr(
    uploaded_file,
    enable_pubchem,
):
    payload = _ocr_api_request_bytes(
        uploaded_file.name,
        uploaded_file.getvalue(),
        uploaded_file.type
        or "application/octet-stream",
        enable_pubchem=enable_pubchem,
    )

    if isinstance(
        payload,
        dict,
    ):
        server_info = payload.get(
            "server"
        )

        if isinstance(
            server_info,
            dict,
        ):
            st.session_state[
                "ocr_last_server_info"
            ] = server_info

    return payload



class MemoryUploadedImage:
    """Small UploadedFile-compatible wrapper for images extracted from ZIP."""
    def __init__(self, name, data, mime_type="image/jpeg"):
        self.name = name
        self._data = data
        self.type = mime_type

    def getvalue(self):
        return self._data


def _guess_image_mime(filename):
    ext = str(filename or "").lower().rsplit(".", 1)[-1]
    return {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "webp": "image/webp",
    }.get(ext, "application/octet-stream")


def _icl_expand_uploads_cached(uploaded_files, cache_key):
    """expand_inventory_scan_uploads, run once per upload set instead of on every rerun.

    Expanding copies, hashes and (for ZIPs) decompresses every uploaded photo.
    With ~100 phone photos that is hundreds of MB of work on each click.
    """
    files = list(uploaded_files or [])
    signature = tuple((getattr(f, 'file_id', None) or f.name, getattr(f, 'size', None)) for f in files)
    cached = st.session_state.get(cache_key)
    if cached and cached[0] == signature:
        return cached[1], cached[2]
    items, warnings = expand_inventory_scan_uploads(files)
    st.session_state[cache_key] = (signature, items, warnings)
    return items, warnings


def expand_inventory_scan_uploads(uploaded_files, max_images=120):
    """
    Accept direct images and/or ZIP files.
    ZIP safety:
    - image extensions only
    - ignore __MACOSX / hidden files
    - no extraction to disk
    - per-file and total uncompressed size limits
    - SHA1 duplicate removal
    """
    image_exts = {".jpg", ".jpeg", ".png", ".webp"}
    max_single = 30 * 1024 * 1024
    max_zip_total = 500 * 1024 * 1024

    items = []
    seen_hashes = set()
    warnings = []

    skipped_over_limit = 0

    def add_item(name, data, source):
        nonlocal skipped_over_limit
        if len(items) >= max_images:
            skipped_over_limit += 1
            return
        if not data:
            return
        if len(data) > max_single:
            warnings.append(f"{name}: skipped (>30 MB)")
            return
        digest = hashlib.sha1(data).hexdigest()
        if digest in seen_hashes:
            warnings.append(f"{name}: duplicate image skipped")
            return
        seen_hashes.add(digest)
        items.append(
            {
                "scan_id": digest[:14],
                "filename": Path(name).name,
                "bytes": data,
                "mime": _guess_image_mime(name),
                "source": source,
            }
        )

    for uploaded in uploaded_files or []:
        name = uploaded.name
        suffix = Path(name).suffix.lower()
        raw = uploaded.getvalue()

        if suffix in image_exts:
            add_item(name, raw, "direct")
            continue

        if suffix != ".zip":
            warnings.append(f"{name}: unsupported file skipped")
            continue

        try:
            with zipfile.ZipFile(BytesIO(raw)) as zf:
                infos = [
                    info
                    for info in zf.infolist()
                    if not info.is_dir()
                ]

                total_uncompressed = sum(
                    max(0, int(info.file_size or 0))
                    for info in infos
                )
                if total_uncompressed > max_zip_total:
                    warnings.append(
                        f"{name}: ZIP skipped (uncompressed size >500 MB)"
                    )
                    continue

                for info in infos:
                    inner = info.filename.replace("\\", "/")
                    parts = [part for part in inner.split("/") if part]

                    if (
                        not parts
                        or "__MACOSX" in parts
                        or any(part.startswith(".") for part in parts)
                    ):
                        continue

                    if Path(inner).suffix.lower() not in image_exts:
                        continue

                    if info.file_size > max_single:
                        warnings.append(
                            f"{inner}: skipped (>30 MB)"
                        )
                        continue

                    try:
                        data = zf.read(info)
                    except Exception:
                        warnings.append(
                            f"{inner}: could not read from ZIP"
                        )
                        continue

                    add_item(inner, data, f"zip:{name}")

                    if len(items) >= max_images:
                        warnings.append(
                            f"Image limit reached ({max_images}). Remaining files skipped."
                        )
                        break
        except zipfile.BadZipFile:
            warnings.append(f"{name}: invalid ZIP")

        if len(items) >= max_images:
            break

    if skipped_over_limit and not any(w.startswith("Image limit reached") for w in warnings):
        warnings.append(f"Image limit reached ({max_images}). {skipped_over_limit} more photo(s) were skipped.")
    return items, warnings


def analyze_scan_memory_item(item, enable_pubchem=False):
    uploaded = MemoryUploadedImage(
        item["filename"],
        item["bytes"],
        item["mime"],
    )
    return analyze_label_with_ocr(
        uploaded,
        enable_pubchem=enable_pubchem,
    )



LABEL_VENDOR_ALIASES = {
    "Aladdin": (
        "ALADDIN",
        "ALADIN",
    ),
    "TCI": (
        "TCI",
        "TOKYO CHEMICAL INDUSTRY",
        "TOKYOCHEMICALINDUSTRY",
    ),
    "Sigma-Aldrich": (
        "SIGMA-ALDRICH",
        "SIGMA ALDRICH",
        "SIGMAALDRICH",
        "ALDRICH",
        "SIGMA",
    ),
    "Acros Organics": (
        "ACROS ORGANICS",
        "ACROSORGANICS",
        "ACROS",
    ),
    "Alfa Aesar": (
        "ALFA AESAR",
        "ALFAAESAR",
    ),
    "Thermo Scientific": (
        "THERMO SCIENTIFIC",
        "THERMOSCIENTIFIC",
        "THERMO",
    ),
    "Fisher Scientific": (
        "FISHER SCIENTIFIC",
        "FISHERSCIENTIFIC",
        "FISHER",
    ),
    "Merck": ("MERCK",),
    "Junsei": ("JUNSEI",),
    "DAEJUNG": ("DAEJUNG",),
    "SAMCHUN": ("SAMCHUN",),
}


ADD_CHEMICAL_MANUFACTURERS = [
    "",
    "TCI",
    "Sigma-Aldrich",
    "Acros Organics",
    "Alfa Aesar",
    "Thermo Scientific",
    "Fisher Scientific",
    "Merck",
    "Junsei",
    "DAEJUNG",
    "SAMCHUN",
    "Aladdin",
]

OTHER_MANUFACTURER_OPTION = "Other / Manual input"


def sync_add_manufacturer_widget():
    """Sync OCR-prefilled manufacturer into the selector/manual widgets."""
    raw_value = clean_text(
        st.session_state.get(
            "add_manufacturer",
            "",
        )
    )
    canonical = canonical_label_vendor(
        raw_value
    )

    if canonical in ADD_CHEMICAL_MANUFACTURERS:
        st.session_state[
            "add_manufacturer_choice"
        ] = canonical
        st.session_state[
            "add_manufacturer_manual"
        ] = ""
    elif raw_value:
        st.session_state[
            "add_manufacturer_choice"
        ] = OTHER_MANUFACTURER_OPTION
        st.session_state[
            "add_manufacturer_manual"
        ] = raw_value
    else:
        st.session_state[
            "add_manufacturer_choice"
        ] = ""
        st.session_state[
            "add_manufacturer_manual"
        ] = ""



def normalize_vendor_key(value):
    return re.sub(
        r"[^A-Z0-9]+",
        "",
        str(value or "").upper(),
    )


def canonical_label_vendor(value):
    raw = clean_text(value)
    if not raw:
        return ""

    normalized = normalize_vendor_key(raw)

    for display, aliases in LABEL_VENDOR_ALIASES.items():
        for alias in aliases:
            alias_norm = normalize_vendor_key(alias)
            if (
                normalized == alias_norm
                or alias_norm in normalized
            ):
                return display

    return raw


def detect_label_vendor_from_result(result):
    """
    Manufacturer of the CURRENT PHYSICAL BOTTLE.

    OCR label evidence wins over DB/PubChem candidate metadata.
    """
    metadata = result.get("metadata") or {}

    metadata_vendor = canonical_label_vendor(
        metadata.get("manufacturer")
    )
    if metadata_vendor:
        return metadata_vendor

    evidence_parts = [
        str(result.get("full_ocr_text") or ""),
    ]

    for query in result.get("ocr_name_queries") or []:
        evidence_parts.append(
            str(query.get("text") or "")
        )

    normalized = normalize_vendor_key(
        " ".join(evidence_parts)
    )

    # TCI is intentionally checked with the same alias table.
    for display, aliases in LABEL_VENDOR_ALIASES.items():
        for alias in aliases:
            alias_norm = normalize_vendor_key(alias)
            if alias_norm and alias_norm in normalized:
                return display

    return ""


def choose_label_catalog(result, manufacturer=""):
    """
    Catalog number printed on the CURRENT bottle.

    For TCI, prefer one-letter + four-digit codes.
    """
    metadata = result.get("metadata") or {}

    if bool(
        metadata.get(
            "catalog_ambiguous"
        )
    ):
        # v18: safer to leave the registration field blank than silently save
        # 185030 when independent OCR passes disagree with 185930.
        return ""

    catalogs = [
        clean_text(value)
        for value in (
            metadata.get("catalog_candidates") or []
        )
        if clean_text(value)
    ]

    if not catalogs:
        return ""

    maker = canonical_label_vendor(manufacturer)

    if maker == "TCI":
        # First: already-correct TCI pattern.
        for value in catalogs:
            compact = re.sub(
                r"[^A-Za-z0-9]",
                "",
                value,
            ).upper()

            if re.fullmatch(
                r"[A-Z]\d{4}",
                compact,
            ):
                return compact

        # Defensive repair for common OCR confusion in the first character:
        # B->8, S->5, O->0, I->1.
        first_char_repairs = {
            "8": "B",
            "5": "S",
            "0": "O",
            "1": "I",
        }

        for value in catalogs:
            compact = re.sub(
                r"[^A-Za-z0-9]",
                "",
                value,
            ).upper()

            if re.fullmatch(
                r"\d{5}",
                compact,
            ):
                first = compact[0]

                if first in first_char_repairs:
                    return (
                        first_char_repairs[first]
                        + compact[1:]
                    )

    return catalogs[0]


def resolve_current_bottle_metadata(result):
    return __import__("icl_ocr_ui_guard").metadata(result)



def normalize_chemical_name(value):
    return re.sub(
        r"[^a-z0-9]+",
        " ",
        str(value or "").casefold(),
    ).strip()



# ============================================================
# ADD CHEMICAL — MULTI-PHOTO EVIDENCE GROUPS
# ============================================================

def _group_text_similarity(value_a, value_b):
    a = normalize_chemical_name(value_a)
    b = normalize_chemical_name(value_b)

    if not a or not b:
        return 0.0

    return (
        SequenceMatcher(
            None,
            a,
            b,
        ).ratio()
        * 100.0
    )


def _join_group_label_fragments(*values):
    parts = [
        clean_text(value)
        for value in values
        if clean_text(value)
    ]

    if not parts:
        return ""

    result = parts[0]

    for part in parts[1:]:
        if result.rstrip().endswith("-"):
            result = (
                result.rstrip()
                + part.lstrip()
            )
        else:
            result = (
                result.rstrip()
                + " "
                + part.lstrip()
            )

    result = re.sub(
        r"\(\s+",
        "(",
        result,
    )
    result = re.sub(
        r"\s+\)",
        ")",
        result,
    )
    return re.sub(
        r"\s+",
        " ",
        result,
    ).strip()


def _candidate_identity_key(candidate):
    cas = clean_text(
        candidate.get("cas")
    )

    if cas:
        return (
            "cas",
            re.sub(
                r"\s+",
                "",
                cas,
            ),
        )

    return (
        "name",
        normalize_chemical_name(
            candidate.get("name")
        ),
    )


def _candidate_rank(candidate):
    source = str(
        candidate.get("source")
        or ""
    ).upper()

    exact_bonus = 0

    if "CAS EXACT" in source:
        exact_bonus = 500
    elif "CATALOG EXACT" in source:
        exact_bonus = 450
    elif "EXACT" in source:
        exact_bonus = 400
    elif "PUBCHEM" in source:
        exact_bonus = 250
    elif "FUZZY" in source:
        exact_bonus = 100

    if candidate.get("auto_select"):
        exact_bonus += 60

    if candidate.get("recommended"):
        exact_bonus += 30

    try:
        score = float(
            candidate.get("score")
            or 0.0
        )
    except Exception:
        score = 0.0

    return (
        exact_bonus + score
    )


def _merge_group_candidates(records):
    """
    Merge identity candidates from several photos of ONE physical bottle.

    Candidate identity remains CAS/name based. No bottle manufacturer/catalog
    metadata is copied into the identity candidate.
    """
    best_by_identity = {}

    for record in records:
        result = (
            record.get("result")
            or {}
        )

        for candidate in (
            result.get("candidates")
            or []
        ):
            key = _candidate_identity_key(
                candidate
            )

            if not key[1]:
                continue

            previous = best_by_identity.get(
                key
            )

            if (
                previous is None
                or _candidate_rank(
                    candidate
                )
                > _candidate_rank(
                    previous
                )
            ):
                best_by_identity[key] = (
                    dict(candidate)
                )

    merged = sorted(
        best_by_identity.values(),
        key=_candidate_rank,
        reverse=True,
    )

    return merged[:8]


def _consensus_value(
    values,
    normalizer=lambda value: clean_text(value),
):
    prepared = []

    for value in values:
        cleaned = clean_text(value)

        if not cleaned:
            continue

        key = normalizer(cleaned)

        if not key:
            continue

        prepared.append(
            (
                key,
                cleaned,
            )
        )

    if not prepared:
        return ""

    counts = Counter(
        key
        for key, _value
        in prepared
    )

    winner_key = counts.most_common(
        1
    )[0][0]

    # Preserve the first human-readable spelling for the winning normalized key.
    for key, display in prepared:
        if key == winner_key:
            return display

    return ""


def _group_label_name_from_records(
    canonical_name,
    records,
):
    """
    Recover one current-bottle label name from multiple photos.

    This is the key cross-photo behavior:
    photo A may contain the beginning of a long name, while photo B contains
    the rest. We reuse OCR evidence already returned by each image and test
    single fragments + cross-photo fragment pairs against the verified
    canonical identity.

    No additional GPU OCR call is made.
    """
    canonical_name = clean_text(
        canonical_name
    )

    if not canonical_name:
        return (
            "",
            0.0,
            "no_canonical_identity",
        )

    fragments = []

    for record_index, record in enumerate(
        records
    ):
        result = (
            record.get("result")
            or {}
        )
        metadata = (
            result.get("metadata")
            or {}
        )

        label_name = clean_text(
            metadata.get("label_name")
        )

        if label_name:
            fragments.append(
                {
                    "text": label_name,
                    "record_index":
                        record_index,
                    "kind": "label_name",
                }
            )

        for query in (
            result.get(
                "ocr_name_queries"
            )
            or []
        ):
            query_text = clean_text(
                query.get("text")
            )

            if query_text:
                fragments.append(
                    {
                        "text":
                            query_text,
                        "record_index":
                            record_index,
                        "kind":
                            "ocr_name_query",
                    }
                )

    # Deduplicate while keeping the record association.
    deduped = []
    seen = set()

    for item in fragments:
        key = (
            item["record_index"],
            normalize_chemical_name(
                item["text"]
            ),
        )

        if not key[1] or key in seen:
            continue

        seen.add(key)
        deduped.append(item)

    fragments = deduped[:24]

    candidates = []

    # Single-photo names/fragments.
    for item in fragments:
        candidates.append(
            (
                item["text"],
                f"photo_{item['record_index'] + 1}:"
                f"{item['kind']}",
            )
        )

    # Cross-photo pairs in BOTH orders. This supports a long product name split
    # over two photographs taken from different sides/angles of the same bottle.
    for left_index, left in enumerate(
        fragments
    ):
        for right in fragments[
            left_index + 1:
        ]:
            if (
                left["record_index"]
                == right["record_index"]
            ):
                continue

            joined_lr = (
                _join_group_label_fragments(
                    left["text"],
                    right["text"],
                )
            )
            joined_rl = (
                _join_group_label_fragments(
                    right["text"],
                    left["text"],
                )
            )

            candidates.append(
                (
                    joined_lr,
                    "cross_photo_pair",
                )
            )
            candidates.append(
                (
                    joined_rl,
                    "cross_photo_pair",
                )
            )

    canonical_norm = normalize_chemical_name(
        canonical_name
    )

    best_text = ""
    best_score = 0.0
    best_source = ""

    for candidate_text, source in candidates:
        candidate_norm = (
            normalize_chemical_name(
                candidate_text
            )
        )

        if not candidate_norm:
            continue

        length_ratio = (
            len(candidate_norm)
            / max(
                1,
                len(canonical_norm),
            )
        )

        if not (
            0.70
            <= length_ratio
            <= 1.50
        ):
            continue

        score = _group_text_similarity(
            canonical_name,
            candidate_text,
        )

        if score > best_score:
            best_text = candidate_text
            best_score = score
            best_source = source

    # Same safety principle as v16: only a high-similarity current-label name
    # is allowed to become the registration display name.
    if (
        best_text
        and best_score >= 90.0
    ):
        return (
            best_text,
            round(
                best_score,
                1,
            ),
            best_source,
        )

    return (
        "",
        round(
            best_score,
            1,
        ),
        "no_high_confidence_group_label_name",
    )



def _group_resolution_payload(
    successful,
    bottle_metadata,
):
    manufacturer = _consensus_value(
        [
            meta.get(
                "manufacturer"
            )
            for meta
            in bottle_metadata
        ],
        normalizer=lambda value:
            normalize_vendor_key(
                canonical_label_vendor(
                    value
                )
            ),
    )

    catalog_candidates = []
    cas_candidates = []
    name_queries = []

    seen_catalog = set()
    seen_cas = set()
    seen_name = set()

    for meta in bottle_metadata:
        for value in (
            [
                meta.get(
                    "catalog"
                )
            ]
            + list(
                meta.get(
                    "catalog_candidates"
                )
                or []
            )
        ):
            display = clean_text(
                value
            )

            key = re.sub(
                r"[^A-Z0-9]",
                "",
                display.upper(),
            )

            if (
                key
                and key not in seen_catalog
            ):
                seen_catalog.add(
                    key
                )
                catalog_candidates.append(
                    display
                )

        for cas in (
            list(
                meta.get(
                    "raw_cas_candidates"
                )
                or []
            )
            + (
                [
                    meta.get(
                        "cas"
                    )
                ]
                if meta.get(
                    "cas"
                )
                else []
            )
        ):
            cas = clean_text(
                cas
            )

            if (
                cas
                and cas not in seen_cas
            ):
                seen_cas.add(
                    cas
                )
                cas_candidates.append(
                    cas
                )

        for value in (
            meta.get(
                "verified_label_name"
            ),
            meta.get(
                "label_name"
            ),
        ):
            text = clean_text(
                value
            )
            key = (
                normalize_chemical_name(
                    text
                )
            )

            if (
                key
                and key not in seen_name
            ):
                seen_name.add(key)
                name_queries.append(
                    text
                )

    for record in successful:
        result = (
            record.get(
                "result"
            )
            or {}
        )

        for query in (
            result.get(
                "ocr_name_queries"
            )
            or []
        ):
            text = clean_text(
                (
                    query.get(
                        "text"
                    )
                    if isinstance(
                        query,
                        dict,
                    )
                    else query
                )
            )

            key = (
                normalize_chemical_name(
                    text
                )
            )

            if (
                key
                and key not in seen_name
            ):
                seen_name.add(key)
                name_queries.append(text)

    return {
        "manufacturer":
            manufacturer,
        "catalog_candidates":
            catalog_candidates,
        "cas_candidates":
            cas_candidates,
        "name_queries":
            name_queries[:24],
    }


def resolve_photo_group_identity(
    successful,
    bottle_metadata,
):
    payload = (
        _group_resolution_payload(
            successful,
            bottle_metadata,
        )
    )

    signature = hashlib.sha1(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
        ).encode(
            "utf-8"
        )
    ).hexdigest()

    cache = st.session_state.setdefault(
        "add_group_resolution_cache",
        {},
    )

    if signature in cache:
        return cache[
            signature
        ]

    url = get_ocr_api_url()
    key = get_ocr_api_key()

    if not url or not key:
        return {
            "ok": False,
            "status":
                "BACKEND_NOT_CONFIGURED",
            "candidates": [],
            "selected_candidate_id":
                None,
            "resolution_method":
                "unavailable",
        }

    try:
        response = requests.post(
            url
            + "/resolve-group",
            headers={
                "X-API-Key": key
            },
            json=payload,
            timeout=60,
        )

        if response.status_code != 200:
            result = {
                "ok": False,
                "status":
                    (
                        "GROUP_RESOLVER_HTTP_"
                        f"{response.status_code}"
                    ),
                "candidates": [],
                "selected_candidate_id":
                    None,
                "resolution_method":
                    "backend_error",
                "error":
                    response.text,
            }
        else:
            result = (
                response.json()
            )

    except Exception as exc:
        result = {
            "ok": False,
            "status":
                "GROUP_RESOLVER_ERROR",
            "candidates": [],
            "selected_candidate_id":
                None,
            "resolution_method":
                "backend_error",
            "error":
                str(exc),
        }

    cache[
        signature
    ] = result

    return result


def _clear_add_form_identity_for_group_change():
    """
    Avoid showing Ammonium bicarbonate (or any previous bottle) while the user
    is currently reviewing another evidence group.
    """
    defaults = {
        "add_name": "",
        "add_description": "",
        "add_cas": "",
        "add_manufacturer": "",
        "add_manufacturer_choice": "",
        "add_manufacturer_manual": "",
        "add_catalog": "",
        "add_lot": "",
        "add_initial_amount": 0.0,
        "add_remaining_amount": 0.0,
        "add_unit": "g",
        "add_formula": "",
        "add_mw": "",
        "add_purity": "",
    }

    for key_name, value in defaults.items():
        st.session_state[
            key_name
        ] = value

    st.session_state[
        "add_manufacturer_widget_needs_sync"
    ] = True


def merge_add_ocr_group(records, group_id):
    return __import__("icl_ocr_ui_guard").merge_group(globals(), records, group_id)


def build_add_ocr_groups(records):
    grouped = {}

    for index, record in enumerate(
        records
    ):
        group_id = int(
            record.get(
                "group_id",
                index + 1,
            )
            or index + 1
        )

        grouped.setdefault(
            group_id,
            [],
        ).append(
            record
        )

    return [
        merge_add_ocr_group(
            grouped[group_id],
            group_id,
        )
        for group_id
        in sorted(grouped)
    ]


def same_vendor(value_a, value_b):
    a = canonical_label_vendor(value_a)
    b = canonical_label_vendor(value_b)

    if not a or not b:
        return False

    return normalize_vendor_key(a) == normalize_vendor_key(b)


def same_catalog(value_a, value_b):
    a = re.sub(
        r"[^A-Z0-9]",
        "",
        str(value_a or "").upper(),
    )
    b = re.sub(
        r"[^A-Z0-9]",
        "",
        str(value_b or "").upper(),
    )

    return bool(a and b and a == b)


def should_recommend_new_physical_bottle(
    bottle_matches,
    label_manufacturer="",
    label_catalog="",
):
    if not bottle_matches:
        return False

    if label_catalog:
        if not any(
            same_catalog(
                label_catalog,
                row["catalog_number"],
            )
            for row in bottle_matches
        ):
            return True

    if label_manufacturer:
        if not any(
            same_vendor(
                label_manufacturer,
                row["manufacturer"],
            )
            for row in bottle_matches
        ):
            return True

    return False



# Shared OCR review for registration and storage checks (V8/V8.1 API).
def _icl_result_key(result):
    return hashlib.sha256(str((result or {}).get('review_token') or (result or {}).get('scan_id') or '').encode()).hexdigest()[:20]


def _icl_review_ready(result, candidate, acknowledged):
    return bool(candidate and candidate.get('candidate_id') and acknowledged
                and result.get('schema_version') == 'icl.ocr.v8'
                and result.get('review_token')
                and result.get('verification_status') != 'CONFLICT'
                and not result.get('conflicts'))


def _icl_eligible_candidates(result):
    """Candidates whose catalog agrees with the label (or that carry an explicit correction)."""
    guard = __import__('icl_ocr_ui_guard')
    label_codes = {str(x.get('value', '')).strip().upper() for x in (result.get('fields') or {}).get('catalog', [])} - {''}
    return [c for c in guard.candidates(result) if c.get('catalog_correction') or not c.get('catalog')
            or not label_codes or str(c['catalog']).strip().upper() in label_codes]


def _icl_review_product(result, scope):
    guard = __import__('icl_ocr_ui_guard')
    if not guard.is_v8(result):
        st.warning('Please analyze the photo again, or switch to manual entry.')
        return None, False
    state = result.get('verification_status')
    if state == 'CONFLICT' or result.get('conflicts'):
        st.error('The name, CAS or catalog number conflicts. Check the label before continuing.')
    elif state == 'PRODUCT_VERIFIED':
        st.success('Product details matched. Compare them with the label and select a product.')
    elif state == 'IDENTITY_VERIFIED':
        st.info('Chemical identity verified. Manufacturer and catalog below may come directly from the bottle label when the manufacturer product page is unavailable.')
    else:
        st.warning('Review required. If the candidates do not match the label, enter details manually or take another photo.')
    label_fields = result.get('fields') or {}
    label_codes = {str(x.get('value', '')).strip().upper() for x in label_fields.get('catalog', [])} - {''}
    label_vendors = {str(x.get('value', '')).strip() for x in label_fields.get('manufacturer', [])} - {''}
    choices = _icl_eligible_candidates(result)
    if len(choices) < len(guard.candidates(result)):
        st.warning('Candidates with a different catalog number were excluded. Rescan to refresh the lookup.')
    if label_codes:
        st.caption('Catalog read from label: ' + ' / '.join(sorted(label_codes)))
    key = scope + '_' + _icl_result_key(result)
    # When every candidate without a catalog correction is the same chemical
    # (one CAS), the first one is preselected; the label comparison checkbox
    # below still has to be ticked before anything is saved.
    plain = [c for c in choices if not c.get('catalog_correction')]
    same = {clean_text(c.get('cas')) or normalize_chemical_name(c.get('name')) for c in plain}
    preselect = choices.index(plain[0]) if plain and len(same) == 1 else None
    selected = st.radio('Select a product', list(range(len(choices))) + [-1], index=preselect,
        format_func=lambda i: 'None match · Find the bottle myself' if i == -1 else guard.candidate_label(choices[i]),
        key=key + '_candidate')
    candidate = choices[selected] if selected is not None and 0 <= selected < len(choices) else None
    acknowledged = False
    if candidate:
        with st.container(border=True):
            st.caption('Product details to apply')
            st.subheader(str(candidate.get('name') or 'Unnamed'))
            a, b = st.columns(2)
            label_manufacturer = next(iter(label_vendors)) if len(label_vendors) == 1 else ''
            label_catalog = next(iter(label_codes)) if len(label_codes) == 1 else ''
            display_manufacturer = candidate.get('manufacturer') or candidate.get('observed_manufacturer') or label_manufacturer
            display_catalog = candidate.get('catalog') or candidate.get('observed_catalog') or label_catalog
            manufacturer_source = '' if candidate.get('manufacturer') else (' · Label' if display_manufacturer else '')
            catalog_source = '' if candidate.get('catalog') else (' · Label' if display_catalog else '')
            a.write('CAS: ' + str(candidate.get('cas') or 'Not confirmed'))
            b.write('Manufacturer: ' + str(display_manufacturer or 'Not confirmed') + manufacturer_source)
            a.write('Catalog: ' + str(display_catalog or 'Not confirmed') + catalog_source)
            b.write('Formula: ' + str(candidate.get('formula') or 'Not confirmed'))
            url = str(candidate.get('source_url') or '')
            if url.startswith(('https://', 'http://')):
                st.link_button('Open product or chemical source', url)
            if not candidate.get('catalog') and len(label_codes) == 1:
                maker_text = (' and manufacturer ' + next(iter(label_vendors))) if len(label_vendors) == 1 else ''
                st.info('Chemical identity is verified independently. Catalog ' + next(iter(label_codes)) + maker_text + ' are bottle-label observations and will be retained after your review; the exact manufacturer SKU is not independently verified.')
                denied = [e for e in (result.get('lookup_events') or []) if e.get('vendor') and e.get('status') in ('access_denied', 'rate_limited', 'timeout', 'network_error', 'budget_exhausted')]
                if denied:
                    st.caption('Manufacturer product lookup was unavailable (' + ', '.join(dict.fromkeys(str(e.get('status')) for e in denied)) + '). This does not invalidate clear label evidence.')
            if candidate.get('catalog_correction'):
                st.warning('Suggested catalog correction: ' + ' / '.join(candidate.get('observed_catalogs', [])) + ' → ' + str(candidate.get('catalog')) + '. Open the manufacturer source and check the original label before applying.')
            # In Storage Scan an existing-inventory candidate is checked again when
            # the physical Bottle ID is confirmed, so the comparison is pre-ticked.
            from_inventory = (scope.startswith('scan_') and candidate.get('source') == 'inventory_unreviewed'
                              and not candidate.get('catalog_correction'))
            acknowledged = st.checkbox('I compared this candidate with the name, CAS and catalog number on the label.',
                value=from_inventory, key=key + '_confirm_' + str(candidate.get('candidate_id')))
            st.session_state[guard.review_key(result, candidate)] = bool(acknowledged)
    with st.expander('OCR text and lookup evidence'):
        observations = result.get('observations') or []
        texts = [str(o.get('text', '')) for o in observations if o.get('pass_id') == 'full']
        st.code('\n'.join(texts or [str(o.get('text', '')) for o in observations]))
        st.json({k: result.get(k) for k in ('fields', 'lookup_events', 'conflicts', 'warnings')})
    packages = (result.get('fields') or {}).get('package') or []
    if packages:
        st.caption('Package marking: ' + ' / '.join(dict.fromkeys(str(p.get('value', '')) for p in packages))
                   + ' · Enter the actual amount and remaining amount when registering.')
    return candidate, _icl_review_ready(result, candidate, acknowledged)


def _icl_exact_bottle_matches(candidate):
    # Catalog numbers are only unique within a manufacturer. Never match a
    # catalog from another vendor or ignore a conflicting recorded CAS.
    rows = find_existing_bottles(cas=candidate.get('cas', ''),
        chemical_name=candidate.get('name', ''), manufacturer=candidate.get('manufacturer', ''))
    found = []
    for row in rows:
        item = dict(row)
        expected = clean_text(candidate.get('cas'))
        actual = clean_text(item.get('cas_number'))
        if expected and actual and expected != actual:
            continue
        if not (expected and actual == expected):
            if normalize_chemical_name(item.get('chemical_name', '')) != normalize_chemical_name(candidate.get('name', '')):
                continue
        found.append(item)
    return found


def _icl_scan_save(record, candidate, bottle, storage, move):
    # Revalidate the signed candidate before making any database mutation.
    result = record['result']
    draft = __import__('icl_ocr_ui_guard')._post(globals(), '/draft', {
        'review_token': result.get('review_token'), 'candidate_id': candidate.get('candidate_id'), 'reviewed': True})
    evidence = dict(result)
    evidence['_user_selected_identity'] = draft['fields']
    return _icl_scan_write(record, bottle, storage, move, evidence, draft['fields'].get('cas'))


def _icl_scan_save_manual(record, bottle, storage, move, search_text, status=None, check_location=True):
    """Link a photo the OCR could not identify to a bottle the user picked from inventory."""
    evidence = dict(record.get('result') or {})
    evidence['_manual_identity'] = {'source': 'manual_inventory_pick', 'search': clean_text(search_text),
        'bottle_id': bottle.get('bottle_id'), 'chemical_name': bottle.get('chemical_name'),
        'cas': bottle.get('cas_number')}
    return _icl_scan_write(record, bottle, storage, move, evidence, None, status=status, check_location=check_location)


def _icl_scan_forget(record):
    """Drop cached inventory rows for a photo so the next rerun reads current bottle data."""
    for key in ('_auto', '_matches', '_searches'):
        record.pop(key, None)


def _icl_scan_write(record, bottle, storage, move, evidence, expected_cas, status=None, check_location=True):
    conn = get_connection()
    try:
        current = conn.execute('SELECT * FROM chemicals WHERE id=? FOR UPDATE', (bottle['id'],)).fetchone()
        if not current or current['status'] == 'Disposed':
            raise ValueError('This bottle was deleted or disposed of. Refresh the inventory before continuing.')
        expected = clean_text(expected_cas)
        if expected and clean_text(current['cas_number']) and expected != clean_text(current['cas_number']):
            raise ValueError('The bottle\'s CAS has changed. Review the record again.')
        if check_location and (current['storage_unit_id'], current['shelf_number']) != (bottle['storage_unit_id'], bottle['shelf_number']):
            raise ValueError('This bottle was moved after the photo was loaded (for example by another photo of the '
                             'same bottle). The row now shows its current location; save it again.')
        if not check_location:
            # Simple mode moves the bottle wherever it currently is; only skip the
            # UPDATE when it is already at the destination.
            move = (current['storage_unit_id'], current['shelf_number']) != (storage[0], storage[1])
        if status and current['status'] != status:
            conn.execute('UPDATE chemicals SET status=?, updated_at=? WHERE id=?',
                (status, datetime.now().strftime('%Y-%m-%d %H:%M'), bottle['id']))
            write_audit_log('EDIT_CHEMICAL', target_type='chemical', target_id=bottle['id'],
                bottle_id=bottle['bottle_id'], details=json.dumps({'status': [current['status'], status],
                'source': 'storage_scan'}, ensure_ascii=False), conn=conn)
        if move:
            conn.execute('UPDATE chemicals SET storage_unit_id=?, shelf_number=?, location=?, updated_at=? WHERE id=?',
                (storage[0], storage[1], storage[2] if storage[0] is not None else '',
                 datetime.now().strftime('%Y-%m-%d %H:%M'), bottle['id']))
            write_audit_log('MOVE_CHEMICAL', target_type='chemical', target_id=bottle['id'],
                bottle_id=bottle['bottle_id'], details=json.dumps({'location': storage[2]}, ensure_ascii=False), conn=conn)
        observation_id = save_inventory_observation(bottle['id'], record['filename'], record['bytes'],
            decode_ocr_label_crop(record.get('result') or {}), evidence, storage[0], storage[1], _conn=conn)
        conn.commit()
        return observation_id
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _icl_clear_add_identity():
    for key in ('add_name', 'add_cas', 'add_formula', 'add_mw', 'add_manufacturer', 'add_catalog',
                'add_lot', 'add_purity', 'add_description', 'add_owner', 'add_expiration'):
        st.session_state[key] = ''
    st.session_state['add_initial_amount'] = 0.0
    st.session_state['add_remaining_amount'] = 0.0
    st.session_state['add_unit'] = 'g'
    st.session_state['add_status'] = 'Unopened'
    st.session_state['add_manufacturer_widget_needs_sync'] = True
    for key in ('icl_v8_registration_draft', 'icl_applied_result_key', 'icl_applied_photos', 'add_source_scan_id'):
        st.session_state.pop(key, None)


def _icl_ocr_connection_panel(scope):
    if not get_ocr_api_key():
        st.warning('OCR_API_KEY is missing. Check the OCR connection settings in Streamlit Secrets.')
    with st.expander('OCR connection'):
        if st.button('Check server connection', key=scope + '_connection'):
            try:
                with st.spinner('Checking the OCR server…'):
                    info = __import__('icl_ocr_ui_guard')._post(globals(), '/health/deep', {})
                if info.get('ok') and info.get('cuda'):
                    st.success('Connected to the OCR server and GPU.')
                else:
                    st.warning('Please check the server response.')
                    st.json(info)
            except Exception as exc:
                st.error('Connection failed: ' + _icl_exception_message(exc))


def _icl_scan_done(record):
    return bool(record.get('saved_observation_id') or record.get('registered_bottle_id') or record.get('skipped'))


def _icl_label_image(image_bytes, result, flip=False):
    """Crop the detected label from the OCR text boxes and turn it upright.

    Box coordinates come from the EXIF-corrected original. Text direction is
    the weighted median angle of each box's long edge; vertical text is assumed
    to read bottom-to-top (the usual way a bottle is held), and ``flip`` turns
    the result 180 degrees when that guess is wrong. Returns JPEG bytes, or
    None when there is not enough box evidence.
    """
    import math
    boxes = []
    for o in (result or {}).get('observations') or []:
        box = o.get('box') or []
        if o.get('pass_id') != 'full' or len(box) < 4 or float(o.get('score') or 0) < 0.5:
            continue
        if len(str(o.get('text') or '').strip()) < 2:
            continue
        boxes.append([(float(x), float(y)) for x, y in box[:4]])
    if not boxes:
        return None
    try:
        image = ImageOps.exif_transpose(Image.open(BytesIO(image_bytes))).convert('RGB')
    except Exception:
        return None
    angles = []
    for b in boxes:
        e1 = (b[1][0] - b[0][0], b[1][1] - b[0][1])
        e2 = (b[2][0] - b[1][0], b[2][1] - b[1][1])
        dx, dy = e1 if math.hypot(*e1) >= math.hypot(*e2) else e2
        angle = math.degrees(math.atan2(dy, dx))
        while angle <= -90:
            angle += 180
        while angle > 90:
            angle -= 180
        if angle > 45:
            angle -= 180  # Vertical text: treat as reading bottom-to-top.
        angles.append((angle, math.hypot(dx, dy)))
    angles.sort()
    half, running, angle = sum(w for _, w in angles) / 2, 0.0, angles[0][0]
    for value, weight in angles:
        running += weight
        if running >= half:
            angle = value
            break
    # Group boxes that sit close together (a label is one dense block of
    # text) and keep the block with the most text area. Stray text on box
    # flaps or other bottles forms separate, smaller blocks.
    rects = [(min(p[0] for p in b), min(p[1] for p in b), max(p[0] for p in b), max(p[1] for p in b)) for b in boxes]
    heights = sorted(min(r[2] - r[0], r[3] - r[1]) for r in rects)
    gap = max(heights[len(heights) // 2] * 4, max(image.width, image.height) * 0.03)
    parent = list(range(len(rects)))
    def _root(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for i, r in enumerate(rects):
        for j in range(i + 1, len(rects)):
            q = rects[j]
            if r[0] - gap <= q[2] and q[0] - gap <= r[2] and r[1] - gap <= q[3] and q[1] - gap <= r[3]:
                parent[_root(i)] = _root(j)
    areas = {}
    for i, r in enumerate(rects):
        areas[_root(i)] = areas.get(_root(i), 0) + (r[2] - r[0]) * (r[3] - r[1])
    best = max(areas, key=areas.get)
    kept = [b for i, b in enumerate(boxes) if _root(i) == best]
    xs = [p[0] for b in kept for p in b]
    ys = [p[1] for b in kept for p in b]
    # Generous margin: low-contrast label text often is not detected at all.
    pad_x = (max(xs) - min(xs)) * 0.25 + image.width * 0.05
    pad_y = (max(ys) - min(ys)) * 0.25 + image.height * 0.05
    crop = image.crop((max(0, int(min(xs) - pad_x)), max(0, int(min(ys) - pad_y)),
                       min(image.width, int(max(xs) + pad_x)), min(image.height, int(max(ys) + pad_y))))
    if crop.width < 8 or crop.height < 8:
        return None
    crop = crop.rotate(angle + (180 if flip else 0), expand=True, resample=Image.BICUBIC, fillcolor=(255, 255, 255))
    if crop.width < 900:
        scale = min(3.0, 900 / crop.width)
        crop = crop.resize((int(crop.width * scale), int(crop.height * scale)), Image.LANCZOS)
    buffer = BytesIO()
    crop.save(buffer, format='JPEG', quality=88)
    return buffer.getvalue()


def _icl_show_label(record, key, caption_prefix=''):
    """Show the upright label crop with the original photo behind an expander."""
    result = record.get('result') or {}
    flip = bool(record.get('_label_flip'))
    cache = record.get('_label_view')
    token = (_icl_result_key(result), flip)
    if not cache or cache[0] != token:
        cache = (token, _icl_label_image(record['bytes'], result, flip))
        record['_label_view'] = cache
    label = cache[1]
    caption = caption_prefix + record['filename']
    if not label:
        st.image(record['bytes'], caption=caption, width=440)
        return
    st.image(label, caption='Detected label · ' + caption, width=640)
    left, right = st.columns([1, 3])
    if left.button('Flip 180°', key=key + '_flip'):
        record['_label_flip'] = not flip
        record.pop('_thumb', None)
        st.rerun()
    with right.expander('Original photo'):
        st.image(record['bytes'], width=440)


def _icl_scan_thumbnail(record):
    # Cached per record; data URLs render directly in st.data_editor image columns.
    if '_thumb' not in record:
        try:
            label = _icl_label_image(record['bytes'], record.get('result'), bool(record.get('_label_flip')))
            image = Image.open(BytesIO(label or record['bytes']))
            image = ImageOps.exif_transpose(image).convert('RGB')
            image.thumbnail((240, 160))
            buffer = BytesIO()
            image.save(buffer, format='JPEG', quality=70)
            record['_thumb'] = 'data:image/jpeg;base64,' + base64.b64encode(buffer.getvalue()).decode()
        except Exception:
            record['_thumb'] = None
    return record['_thumb']


def _icl_scan_target_location(storage_options, show_move=True):
    """Batch-wide location chosen once; individual review uses it as the default."""
    st.markdown('#### Where were these bottles found or moved to?')
    labels = [option[2] for option in storage_options]
    current = st.session_state.get('icl_scan_target_label')
    target_label = st.selectbox('Location for this batch', labels,
        index=labels.index(current) if current in labels else None,
        placeholder="Keep each bottle's current location", key='icl_scan_target_label')
    target = storage_options[labels.index(target_label)] if target_label in labels else None
    if not show_move:
        return target, target is not None
    move = st.checkbox("Also update each bottle's current storage location", value=True,
        key='icl_scan_target_move', disabled=target is None)
    return target, bool(target is not None and move)


def _icl_scan_location_for(bottle, storage_options, target, move):
    if target is not None:
        observed = target
    else:
        observed = storage_options[current_location_index(storage_options, bottle['storage_unit_id'], bottle['shelf_number'])]
    unchanged = (observed[0], observed[1]) == (bottle['storage_unit_id'], bottle['shelf_number'])
    return observed, bool(move and not unchanged)


def _icl_lot_key(value):
    return re.sub(r'[^A-Z0-9]', '', str(value or '').upper())


def _icl_scan_auto_match(record):
    """Prepare a photo for the quick-confirm table, or say why it needs review.

    Candidates that only differ in name but share one CAS are treated as the
    same chemical. All non-disposed bottles of that chemical become choices;
    a Lot number read from the label picks the bottle when it matches exactly
    one. Cached per OCR result so table reruns do not query the database again.
    """
    result = record.get('result') or {}
    key = _icl_result_key(result)
    cached = record.get('_auto')
    if cached and cached.get('key') == key:
        return cached
    auto = {'key': key, 'candidate': None, 'bottles': [], 'bottle': None, 'note': '', 'reason': ''}
    record['_auto'] = auto
    guard = __import__('icl_ocr_ui_guard')
    if record.get('error') or not result:
        auto['reason'] = 'Analysis failed'
        return auto
    if not guard.is_v8(result) or not result.get('review_token'):
        auto['reason'] = 'Analyze again'
        return auto
    if result.get('verification_status') == 'CONFLICT' or result.get('conflicts'):
        auto['reason'] = 'Conflicting label evidence'
        return auto
    choices = [c for c in _icl_eligible_candidates(result) if not c.get('catalog_correction')]
    if not choices:
        auto['reason'] = 'Catalog correction needs review' if _icl_eligible_candidates(result) else 'No product candidate'
        return auto
    identities = {clean_text(c.get('cas')) or normalize_chemical_name(c.get('name')) for c in choices}
    if len(identities) > 1:
        auto['reason'] = f'{len(identities)} different chemicals'
        return auto
    bottles = {}
    try:
        for choice in choices:
            for bottle in _icl_exact_bottle_matches(choice):
                bottles.setdefault(bottle['id'], bottle)
    except Exception as exc:
        auto['reason'] = 'Inventory lookup failed: ' + _icl_exception_message(exc)
        return auto
    bottles = list(bottles.values())
    if not bottles:
        auto['reason'] = 'No matching bottle (new bottle?)'
        return auto
    names = {normalize_chemical_name(b.get('chemical_name')) for b in bottles}
    auto['candidate'] = next((c for c in choices if normalize_chemical_name(c.get('name')) in names), choices[0])
    auto['bottles'] = bottles
    if len(bottles) == 1:
        auto['bottle'] = bottles[0]
        return auto
    lots = {_icl_lot_key(x.get('value')) for x in (result.get('fields') or {}).get('lot', [])}
    lots.add(_icl_lot_key((result.get('metadata') or {}).get('lot')))
    lots.discard('')
    lot_hits = [b for b in bottles if _icl_lot_key(b.get('lot_number')) in lots]
    if len(lot_hits) == 1:
        auto['bottle'] = lot_hits[0]
        auto['note'] = 'Lot matched'
    else:
        auto['note'] = f'{len(bottles)} bottles match · check Bottle ID'
    return auto


def _icl_bottle_option(bottle):
    """Bottle choice label: enough to tell same-chemical bottles apart by their label."""
    amount = bottle.get('initial_amount') or bottle.get('amount')
    try:
        amount = (f'{float(amount):g} ' + clean_text(bottle.get('unit'))).strip() if amount not in (None, '') and float(amount) > 0 else ''
    except (TypeError, ValueError):
        amount = clean_text(amount)
    lot = clean_text(bottle.get('lot_number'))
    parts = [bottle.get('bottle_id'), bottle.get('chemical_name'), clean_text(bottle.get('manufacturer')),
             amount, clean_text(bottle.get('catalog_number')), 'Lot ' + lot if lot else '',
             bottle.get('location_text') or 'Not assigned']
    return ' · '.join(str(x) for x in parts if x)


@st.fragment
def _icl_scan_quick_table(batch, storage_options, target, move):
    pending = [r for r in batch if not _icl_scan_done(r)]
    ready = [r for r in pending if _icl_scan_auto_match(r)['candidate']]
    st.markdown(f'#### Quick confirm · {len(ready)} of {len(pending)} photos')
    if not ready:
        st.caption('No photo matched an inventory bottle. Review the photos below one by one.')
        return
    st.caption('Check each photo against the Bottle ID, tick the rows that are correct, then save them together. '
               'When several bottles of the same chemical exist, pick the right Bottle ID in the table. '
               'Ticking a row confirms that you compared the label with the product and the physical Bottle ID.')
    for failure in st.session_state.get('icl_scan_quick_failures') or []:
        st.error('Not saved · ' + failure)
    select_all = st.checkbox('Select all rows', value=False, key='icl_scan_quick_all')
    # Fixed picks first, then give same-chemical photos different bottles by default.
    used = {r['_auto']['bottle']['id'] for r in ready if r['_auto']['bottle']}
    defaults = {}
    for record in ready:
        auto = record['_auto']
        pick = auto['bottle'] or next((b for b in auto['bottles'] if b['id'] not in used), auto['bottles'][0])
        used.add(pick['id'])
        defaults[record['scan_id']] = pick
    rows, options = [], []
    for record in ready:
        auto = record['_auto']
        pick = defaults[record['scan_id']]
        observed, will_move = _icl_scan_location_for(pick, storage_options, target, move)
        for bottle in auto['bottles']:
            label = _icl_bottle_option(bottle)
            if label not in options:
                options.append(label)
        rows.append({
            'Save': select_all,
            'Label': _icl_scan_thumbnail(record),
            'Chemical': str(auto['candidate'].get('name') or ''),
            'CAS': str(auto['candidate'].get('cas') or ''),
            'Bottle': _icl_bottle_option(pick),
            'Check': auto['note'] or 'One bottle matches',
            'After saving': (('Move to ' if will_move else 'Observed at ') + observed[2]),
            'File': record['filename'],
        })
    signature = hashlib.sha256('|'.join(r['scan_id'] for r in ready).encode()).hexdigest()[:12]
    edited = st.data_editor(
        pd.DataFrame(rows),
        hide_index=True,
        disabled=[c for c in rows[0] if c not in ('Save', 'Bottle')],
        column_config={
            'Save': st.column_config.CheckboxColumn('Save', width='small'),
            'Label': st.column_config.ImageColumn('Label', width='medium'),
            'Bottle': st.column_config.SelectboxColumn('Bottle', options=options, required=True, width='large'),
        },
        key=f'icl_scan_quick_{signature}_{select_all}',
    )
    chosen = [(ready[i], bottle_label) for i, (flag, bottle_label)
              in enumerate(zip(edited['Save'].tolist(), edited['Bottle'].tolist())) if flag]
    if st.button(f'Save {len(chosen)} selected', type='primary', disabled=not chosen, key='icl_scan_quick_save'):
        seen, saved, failures = {}, 0, []
        progress = st.progress(0.0)
        for index, (record, bottle_label) in enumerate(chosen):
            auto = record['_auto']
            bottle = next((b for b in auto['bottles'] if _icl_bottle_option(b) == bottle_label), None)
            if bottle is None:
                failures.append(f"{record['filename']}: {bottle_label} is not a {auto['candidate'].get('name')} bottle.")
            elif bottle['id'] in seen:
                failures.append(f"{record['filename']}: bottle {bottle['bottle_id']} was also chosen for {seen[bottle['id']]}.")
            else:
                seen[bottle['id']] = record['filename']
                observed, will_move = _icl_scan_location_for(bottle, storage_options, target, move)
                try:
                    record['saved_observation_id'] = _icl_scan_save(record, auto['candidate'], bottle, observed, will_move)
                    record['resolved_action'] = 'existing_bottle_observation'
                    saved += 1
                except Exception as exc:
                    _icl_scan_forget(record)
                    failures.append(f"{record['filename']}: {_icl_exception_message(exc)}")
            progress.progress((index + 1) / len(chosen))
        st.session_state['icl_scan_notice'] = (saved, [])
        st.session_state['icl_scan_quick_failures'] = failures
        remaining = [i for i, r in enumerate(batch) if not _icl_scan_done(r)]
        if remaining:
            st.session_state['_icl_scan_next'] = remaining[0]
        st.rerun()


_ICL_SEARCH_COLUMNS = ('c.bottle_id', 'c.legacy_id', 'c.chemical_name', 'c.cas_number',
                       'c.catalog_number', 'c.lot_number', 'c.manufacturer')


def _icl_search_bottles(terms, match_all=True):
    """Inventory search for photos the OCR could not identify.

    match_all=True: every term must appear in some column (typed search).
    match_all=False: any term may match (suggestions from label fields).
    """
    terms = [t for t in dict.fromkeys(clean_text(t) for t in terms) if len(t) >= 2][:8]
    if not terms:
        return []
    per_term = '(' + ' OR '.join(col + ' ILIKE ?' for col in _ICL_SEARCH_COLUMNS) + ')'
    where = (' AND ' if match_all else ' OR ').join([per_term] * len(terms))
    params = []
    for term in terms:
        params.extend(['%' + term + '%'] * len(_ICL_SEARCH_COLUMNS))
    conn = get_connection()
    try:
        rows = conn.execute(f"""
            SELECT c.*, {LOCATION_EXPR} AS location_text
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed' AND ({where})
            ORDER BY c.chemical_name, c.bottle_id
            LIMIT 40
            """, params).fetchall()
    finally:
        conn.close()
    return [dict(row) for row in rows]


def _icl_label_terms(result):
    """(label, value) pairs the OCR parsed from the label (catalog, lot, CAS, maker, label name)."""
    fields = result.get('fields') or {}
    meta = result.get('metadata') or {}
    pairs = []
    for field, label in (('catalog', 'Catalog'), ('lot', 'Lot'), ('cas', 'CAS'), ('name', 'Name'), ('manufacturer', 'Maker')):
        # fields['name'] holds single-word name fragments; only the joined label name is useful here.
        values = [meta.get('label_name')] if field == 'name' else [x.get('value') for x in fields.get(field) or []] + [meta.get(field)]
        if field == 'catalog':
            # Labels often print the package size after the catalog (C11081-100G);
            # inventory usually stores only the catalog number.
            values = [re.sub(r'(?i)-\d+(\.\d+)?\s*(G|MG|KG|UG|ML|L|EA)$', '', clean_text(v)) for v in values] + values
        for value in values:
            value = clean_text(value)
            if value and (label, value) not in pairs:
                pairs.append((label, value))
    return pairs


def _icl_ocr_fragments(result):
    """Every text piece the OCR read, in label reading order (rows top to bottom, left to right).

    Box corners follow the text direction (p0 -> p1 runs along the text), so the
    summed p0 -> p1 vectors give the label's orientation even on a rotated photo.
    """
    import math
    items = []
    for o in (result or {}).get('observations') or []:
        text = clean_text(o.get('text'))
        if o.get('pass_id', 'full') != 'full' or len(text) < 2:
            continue
        box = o.get('box') or []
        if len(box) >= 4:
            pts = [(float(x), float(y)) for x, y in box[:4]]
            items.append((text, pts))
        else:
            items.append((text, None))
    boxed = [(t, b) for t, b in items if b]
    if not boxed:
        return list(dict.fromkeys(t for t, _ in items))
    vx = sum(b[1][0] - b[0][0] for _, b in boxed)
    vy = sum(b[1][1] - b[0][1] for _, b in boxed)
    angle = math.atan2(vy, vx)
    cos, sin = math.cos(angle), math.sin(angle)
    placed = []
    for text, b in boxed:
        cx = sum(x for x, _ in b) / 4
        cy = sum(y for _, y in b) / 4
        height = math.hypot(b[3][0] - b[0][0], b[3][1] - b[0][1]) or 1.0
        placed.append((-sin * cx + cos * cy, cos * cx + sin * cy, height, text))
    placed.sort()
    row_gap = sorted(h for *_, h, _ in placed)[len(placed) // 2] * 0.6
    rows, current, last_y = [], [], None
    for y, x, h, text in placed:
        if last_y is not None and y - last_y > row_gap:
            rows.append(current)
            current = []
        current.append((x, text))
        last_y = y
    rows.append(current)
    ordered = [text for row in rows for _, text in sorted(row)]
    ordered += [t for t, b in items if not b]
    return list(dict.fromkeys(ordered))


def _icl_join_pieces(pieces):
    return ' '.join(clean_text(p).strip(' ,;:') for p in pieces if clean_text(p).strip(' ,;:'))


def _icl_scan_register_form(record, storage_options, target, scope, prefill, candidate, on_saved, default_status='Unopened'):
    """Register a bottle that is not in inventory without leaving Storage Scan.

    The photo is attached as the new bottle's first observation. When the
    details came from an OCR candidate, the signed candidate is revalidated
    with /draft before anything is written.
    """
    pieces = _icl_ocr_fragments(record.get('result') or {})
    if pieces:
        pick = scope + '_regpick_' + str((candidate or {}).get('candidate_id', 'manual'))
        st.caption('Fill fields from OCR pieces (optional)')
        n, c1, c2, c3 = st.columns([2, 1, 1, 1])
        name_pieces = n.multiselect('Name pieces, in order', pieces, key=pick + '_name')
        cas_piece = c1.selectbox('CAS', pieces, index=None, placeholder='—', key=pick + '_cas')
        catalog_piece = c2.selectbox('Catalog', pieces, index=None, placeholder='—', key=pick + '_catalog')
        lot_piece = c3.selectbox('Lot', pieces, index=None, placeholder='—', key=pick + '_lot')
        prefill = dict(prefill)
        if name_pieces:
            prefill['name'] = _icl_join_pieces(name_pieces)
        for field, piece in (('cas', cas_piece), ('catalog', catalog_piece), ('lot', lot_piece)):
            if piece:
                prefill[field] = _icl_join_pieces([piece])
    # The prefill is part of the key so picking other OCR pieces refreshes the form.
    key = scope + '_reg_' + _icl_result_key(record.get('result') or {}) + '_' + str((candidate or {}).get('candidate_id', 'manual')) \
        + '_' + hashlib.sha256(json.dumps(prefill, sort_keys=True, default=str).encode()).hexdigest()[:8]
    maker = clean_text(prefill.get('manufacturer'))
    with st.form(key + '_form', enter_to_submit=False):
        name = st.text_input('Chemical name *', value=clean_text(prefill.get('name')), key=key + '_name')
        a, b = st.columns(2)
        cas = a.text_input('CAS No.', value=clean_text(prefill.get('cas')), key=key + '_cas')
        maker = (canonical_label_vendor(maker) or maker) if maker else ''
        makers = list(ADD_CHEMICAL_MANUFACTURERS) + ([maker] if maker and maker not in ADD_CHEMICAL_MANUFACTURERS else []) + [OTHER_MANUFACTURER_OPTION]
        maker_choice = b.selectbox('Manufacturer', makers, index=makers.index(maker),
            format_func=lambda v: 'Select manufacturer' if v == '' else v, key=key + '_maker_choice')
        other_maker = b.text_input('Other manufacturer (when "' + OTHER_MANUFACTURER_OPTION + '" is selected)', key=key + '_maker')
        catalog = a.text_input('Catalog No.', value=clean_text(prefill.get('catalog')), key=key + '_catalog')
        lot = b.text_input('Lot No.', value=clean_text(prefill.get('lot')), key=key + '_lot')
        amount_col, unit_col, remaining_col = st.columns([2, 1, 2])
        units = ['g', 'mg', 'kg', 'mL', 'L', 'ea']
        amount = amount_col.number_input('Amount', min_value=0.0, value=float(prefill.get('amount') or 0.0), key=key + '_amount')
        unit = unit_col.selectbox('Unit', units, index=units.index(prefill['unit']) if prefill.get('unit') in units else 0, key=key + '_unit')
        remaining = remaining_col.number_input('Remaining amount', min_value=0.0, value=float(prefill.get('amount') or 0.0), key=key + '_remaining')
        c, d = st.columns(2)
        storage = c.selectbox('Storage location', storage_options,
            index=storage_options.index(target) if target in storage_options else 0,
            format_func=lambda x: x[2], key=key + '_storage')
        statuses = ['Unopened', 'In Use', 'Disposal Pending']
        status = d.selectbox('Status', statuses, index=statuses.index(default_status) if default_status in statuses else 0,
            key=key + '_status')
        confirmed = st.checkbox('I checked the name, CAS and amount against the label.', key=key + '_confirm')
        submitted = st.form_submit_button('Register new bottle with this photo', type='primary')
    if not submitted:
        return
    if not confirmed:
        st.error('Tick the confirmation checkbox first.')
        return
    if not name.strip():
        st.error('Enter a chemical name.')
        return
    if remaining > amount:
        st.error('Remaining amount cannot exceed the amount.')
        return
    manufacturer = other_maker if maker_choice == OTHER_MANUFACTURER_OPTION else maker_choice
    draft = None
    try:
        if candidate:
            draft = __import__('icl_ocr_ui_guard')._post(globals(), '/draft', {
                'review_token': record['result'].get('review_token'),
                'candidate_id': candidate.get('candidate_id'), 'reviewed': True})
        timestamp = datetime.now().strftime('%Y-%m-%d %H:%M')
        values = dict(chemical_name=name.strip(), description='', cas_number=cas.strip(),
            manufacturer=manufacturer.strip(), catalog_number=catalog.strip(), lot_number=lot.strip(),
            amount=amount, initial_amount=amount, remaining_amount=remaining, unit=unit,
            storage_unit_id=storage[0], shelf_number=storage[1],
            location=storage[2] if storage[0] is not None else '', owner='', expiration_date='', status=status,
            source_registered_date=datetime.now(ZoneInfo('Asia/Seoul')).date().strftime('%Y-%m-%d'),
            molecular_formula=clean_text(prefill.get('formula')), molecular_weight=clean_text(prefill.get('mw')),
            purity='', created_at=timestamp, updated_at=timestamp)
        bottle_id = _icl_register_new_bottle(values, storage, [record], draft or {'fields': {'source': 'storage_scan_manual_registration'}},
            record['scan_id'])
    except Exception as exc:
        st.error('Registration failed: ' + _icl_exception_message(exc))
        return
    record['registered_bottle_id'] = bottle_id
    record['resolved_action'] = 'new_bottle_registration'
    try:
        safety = fetch_pubchem_safety(cas.strip(), name.strip())
        if safety:
            save_pubchem_safety(int(bottle_id.rsplit('-', 1)[1]), safety)
    except Exception:
        pass
    on_saved()


def _icl_scan_manual_panel(record, storage_options, target, move, scope, on_saved):
    """Find the bottle by hand when the OCR has no usable product candidate."""
    result = record.get('result') or {}
    st.markdown('##### Find the bottle in inventory')
    pairs = _icl_label_terms(result)
    pieces = _icl_ocr_fragments(result)
    search_key = scope + '_manual_q'
    if pairs:
        st.caption('Read from label: ' + ' · '.join(label + ' ' + value for label, value in pairs))

    def use_pieces():
        st.session_state[search_key] = _icl_join_pieces(st.session_state.get(scope + '_manual_pieces') or [])

    if pieces:
        st.multiselect('Pick OCR pieces to search (click them in order, e.g. Carbon → tetrabromide)', pieces,
            key=scope + '_manual_pieces', on_change=use_pieces, placeholder='Pieces are listed in label reading order')
    query = st.text_input('Search inventory', key=search_key,
        placeholder='Chemical name, CAS, catalog, lot or Bottle ID')
    searches = record.setdefault('_searches', {})

    def search(terms, match_all):
        key = (tuple(terms), match_all)
        if key not in searches:
            searches[key] = _icl_search_bottles(terms, match_all=match_all)
        return searches[key]

    try:
        if query.strip():
            bottles = search(query.split(), True)
            if not bottles and len(query.split()) > 1:
                bottles = search([query], True)
            heading = f'{len(bottles)} bottles match "{query.strip()}"'
        else:
            codes = [v for label, v in pairs if label in ('Catalog', 'Lot', 'CAS')]
            bottles = search(codes, False) if codes else []
            heading = (f'{len(bottles)} bottles share a catalog, lot or CAS read from the label'
                       if codes else 'Type a search or pick text read from the label.')
    except Exception as exc:
        st.error('Inventory search failed: ' + _icl_exception_message(exc))
        return
    st.caption(heading)
    if bottles:
        mapping = {b['id']: b for b in bottles}
        lots = {_icl_lot_key(v) for label, v in pairs if label == 'Lot'} - {''}
        lot_hits = [b['id'] for b in bottles if _icl_lot_key(b.get('lot_number')) in lots]
        ids = list(mapping)
        default = ids.index(lot_hits[0]) if len(lot_hits) == 1 else (0 if len(ids) == 1 else None)

        bottle_id = st.selectbox('Select the physical bottle', ids, index=default, format_func=lambda i: _icl_bottle_option(mapping[i]),
            placeholder='Choose a bottle', key=scope + '_manual_bottle')
        if bottle_id is not None:
            bottle = mapping[bottle_id]
            key = scope + '_manual_' + str(bottle_id)
            default_observed, _ = _icl_scan_location_for(bottle, storage_options, target, move)
            observed = st.selectbox('Observed storage location', storage_options,
                index=storage_options.index(default_observed), format_func=lambda x: x[2], key=key + '_storage')
            move_here = st.checkbox("Also update the bottle's current storage location", value=True, key=key + '_move')
            unchanged = (observed[0], observed[1]) == (bottle['storage_unit_id'], bottle['shelf_number'])
            move_here = bool(move_here and not unchanged)
            st.write(f"**{bottle['bottle_id']}** · " + (('Move to ' if move_here else 'Observed at ') + observed[2]))
            if st.button('Link this photo to the bottle and save', type='primary', key=key + '_save'):
                try:
                    record['saved_observation_id'] = _icl_scan_save_manual(record, bottle, observed, move_here, query)
                    record['resolved_action'] = 'existing_bottle_observation'
                except Exception as exc:
                    _icl_scan_forget(record)
                    st.error('Could not save: ' + _icl_exception_message(exc))
                else:
                    on_saved()
    with st.expander('Not in inventory? Register it as a new bottle here'):
        values = {}
        for label, value in pairs:
            values.setdefault(label, value)
        typed_name = query.strip() if query.strip() and not any(ch.isdigit() for ch in query) else ''
        prefill = {'name': values.get('Name') or typed_name, 'cas': values.get('CAS', ''),
                   'catalog': values.get('Catalog', ''), 'lot': values.get('Lot', ''), 'manufacturer': values.get('Maker', '')}
        _icl_scan_register_form(record, storage_options, target, scope, prefill, None, on_saved)


def _icl_scan_simple_choices(record, target):
    """Inventory bottles this photo most likely shows, best first (cached per photo).

    Union of every OCR candidate's bottles and bottles sharing a catalog, lot or
    CAS read from the label. A lot match comes first, then a catalog match.
    """
    result = record.get('result') or {}
    key = _icl_result_key(result)
    cached = record.get('_simple')
    if cached and cached[0] == key:
        return cached[1]
    found = {}
    auto = _icl_scan_auto_match(record)
    for bottle in auto.get('bottles') or []:
        found.setdefault(bottle['id'], bottle)
    for candidate in _icl_eligible_candidates(result)[:6]:
        try:
            for bottle in _icl_exact_bottle_matches(candidate):
                found.setdefault(bottle['id'], bottle)
        except Exception:
            pass
    pairs = _icl_label_terms(result)
    codes = [v for label, v in pairs if label in ('Catalog', 'Lot', 'CAS')]
    if codes:
        try:
            for bottle in _icl_search_bottles(codes, match_all=False):
                found.setdefault(bottle['id'], bottle)
        except Exception:
            pass
    lots = {_icl_lot_key(v) for label, v in pairs if label == 'Lot'} - {''}
    catalogs = {str(v).upper() for label, v in pairs if label == 'Catalog'}

    def rank(bottle):
        lot_hit = _icl_lot_key(bottle.get('lot_number')) in lots
        catalog_hit = clean_text(bottle.get('catalog_number')).upper() in catalogs
        return (not lot_hit, not catalog_hit, str(bottle.get('chemical_name') or ''), str(bottle.get('bottle_id') or ''))

    choices = sorted(found.values(), key=rank)[:10]
    record['_simple'] = (key, choices)
    return choices


@st.fragment
def _icl_scan_simple(batch, storage_options, target, status):
    """One photo at a time: pick the bottle from a short list, search, or register."""
    pending = [i for i, r in enumerate(batch) if not _icl_scan_done(r)]
    last = st.session_state.pop('icl_simple_last', None)
    if last:
        st.success(last)
    if not pending:
        st.success('Every photo in this batch is done.')
        return
    if target is None:
        st.info('Choose where these bottles go (for example 261014 throw) above to start.')
        return
    current = st.session_state.get('icl_simple_index')
    if current not in pending:
        current = pending[0]
    position = len(batch) - len(pending) + 1
    record = batch[current]
    scope = 'simple_' + record['scan_id']
    st.markdown(f'#### Photo {position} of {len(batch)}')

    def done(message):
        record.pop('_label_view', None)
        later = [i for i in pending if i > current] or [i for i in pending if i < current]
        st.session_state['icl_simple_index'] = later[0] if later else None
        st.session_state['icl_simple_last'] = message
        st.rerun(scope='app')

    left, right = st.columns([1, 1.3])
    with left:
        _icl_show_label(record, scope)
        if st.button('Skip this photo', key=scope + '_skip', width='stretch'):
            record['skipped'] = True
            done('Skipped ' + record['filename'])
    with right:
        if record.get('error') or not record.get('result'):
            st.error('This photo could not be read: ' + str(record.get('error') or 'no result'))
            choices = []
        else:
            choices = _icl_scan_simple_choices(record, target)
        if choices:
            st.markdown('**1. Is it one of these bottles?**')
            for bottle in choices:
                here = (bottle.get('storage_unit_id'), bottle.get('shelf_number')) == (target[0], target[1])
                label = _icl_bottle_option(bottle) + ('  ✓ already here' if here else '')
                if st.button(label, key=scope + '_pick_' + str(bottle['id']), width='stretch'):
                    try:
                        record['saved_observation_id'] = _icl_scan_save_manual(
                            record, bottle, target, True, 'simple mode pick', status=status, check_location=False)
                        record['resolved_action'] = 'existing_bottle_observation'
                    except Exception as exc:
                        _icl_scan_forget(record)
                        record.pop('_simple', None)
                        st.error('Could not save: ' + _icl_exception_message(exc))
                    else:
                        done(f"Saved {bottle['bottle_id']} · {bottle.get('chemical_name')} → {target[2]}")
        else:
            st.markdown('**1. No matching bottle was found automatically.**')
        st.markdown('**2. Not in the list? Search the whole inventory**')
        pieces = _icl_ocr_fragments(record.get('result') or {})
        search_key = scope + '_q'

        def use_pieces():
            st.session_state[search_key] = _icl_join_pieces(st.session_state.get(scope + '_pieces') or [])

        if pieces:
            st.multiselect('Or pick words the OCR read', pieces, key=scope + '_pieces', on_change=use_pieces,
                placeholder='e.g. Carbon → tetrabromide')
        query = st.text_input('Search by name, CAS, catalog or lot', key=search_key)
        if query.strip():
            try:
                hits = _icl_search_bottles(query.split(), match_all=True)[:15]
            except Exception as exc:
                hits = []
                st.error('Search failed: ' + _icl_exception_message(exc))
            shown = {b['id'] for b in choices}
            hits = [b for b in hits if b['id'] not in shown]
            if not hits:
                st.caption('No other bottles match "' + query.strip() + '".')
            for bottle in hits:
                if st.button(_icl_bottle_option(bottle), key=scope + '_hit_' + str(bottle['id']), width='stretch'):
                    try:
                        record['saved_observation_id'] = _icl_scan_save_manual(
                            record, bottle, target, True, query, status=status, check_location=False)
                        record['resolved_action'] = 'existing_bottle_observation'
                    except Exception as exc:
                        st.error('Could not save: ' + _icl_exception_message(exc))
                    else:
                        done(f"Saved {bottle['bottle_id']} · {bottle.get('chemical_name')} → {target[2]}")
        with st.expander('3. Still not found? Register it as a new bottle'):
            result = record.get('result') or {}
            values = {}
            for label, value in _icl_label_terms(result):
                values.setdefault(label, value)
            top = next(iter(_icl_eligible_candidates(result)), {}) if result else {}
            typed = query.strip() if query.strip() and not any(ch.isdigit() for ch in query) else ''
            prefill = {'name': typed or top.get('name') or values.get('Name', ''),
                       'cas': top.get('cas') or values.get('CAS', ''),
                       'catalog': values.get('Catalog', '') or top.get('catalog') or '',
                       'lot': values.get('Lot', ''),
                       'manufacturer': values.get('Maker', '') or top.get('manufacturer') or ''}
            _icl_scan_register_form(record, storage_options, target, scope, prefill, None,
                lambda: done(f"Registered a new bottle → {target[2]}"), default_status=status or 'Unopened')


@st.fragment
def _icl_scan_review_one(batch, storage_options, target, move):
    pending_order = [i for i, r in enumerate(batch) if not _icl_scan_done(r)]
    jump = st.session_state.pop('_icl_scan_next', None)
    if jump is not None:
        st.session_state['icl_scan_selected'] = jump
    if st.session_state.get('icl_scan_selected') not in range(len(batch)):
        st.session_state['icl_scan_selected'] = pending_order[0] if pending_order else 0

    def advance(current):
        later = [i for i in pending_order if i > current and i != current]
        earlier = [i for i in pending_order if i < current]
        nxt = (later or earlier or [current])[0]
        st.session_state['_icl_scan_next'] = nxt

    st.markdown('#### Review one photo')
    selected_index = st.selectbox('Photo to review', range(len(batch)), key='icl_scan_selected',
        format_func=lambda i: f'{i + 1}. {batch[i]["filename"]}' + (' · Complete' if _icl_scan_done(batch[i]) else ' · Pending review'))
    record = batch[selected_index]
    scope = 'scan_' + record['scan_id']
    _icl_show_label(record, scope)
    if record.get('saved_observation_id') or record.get('registered_bottle_id'):
        st.success('Saved · ' + str(record.get('registered_bottle_id') or 'Observation linked to the existing bottle.'))
        return
    if record.get('skipped'):
        st.info('This photo was skipped.')
        if st.button('Review again', key=scope + '_unskip'):
            record.pop('skipped', None)
            st.rerun()
        return
    if record.get('error'):
        st.error('Analysis failed: ' + record['error'])
    retry, skip = st.columns(2)
    if retry.button('Retry this photo', key=scope + '_retry'):
        try:
            with st.spinner('Analyzing the original photo again…'):
                record['result'] = analyze_scan_memory_item(record, enable_pubchem=True)
            record.pop('error', None)
            record.pop('_auto', None)
        except Exception as exc:
            record['error'] = _icl_exception_message(exc)
        st.rerun()
    if skip.button('Skip this photo', key=scope + '_skip'):
        record['skipped'] = True
        advance(selected_index)
        st.rerun()
    result = record.get('result') or {}
    if not result or record.get('error'):
        return
    auto = _icl_scan_auto_match(record)
    if auto['reason']:
        st.caption('Needs review · ' + auto['reason'])
    def saved():
        st.session_state['icl_scan_notice'] = (1, [])
        advance(selected_index)
        st.rerun(scope='app')

    candidate, ready = _icl_review_product(result, scope)
    if candidate is None:
        _icl_scan_manual_panel(record, storage_options, target, move, scope, saved)
        return
    if not ready:
        if result.get('verification_status') == 'CONFLICT' or result.get('conflicts'):
            # The OCR server refuses /draft for conflicting evidence, so the candidate
            # cannot be saved. Linking to a bottle picked from inventory still works.
            reasons = ', '.join(str(c).split(':')[0].replace('_', ' ') for c in result.get('conflicts') or []) or 'conflicting evidence'
            st.warning('This candidate cannot be saved because the label evidence conflicts (' + reasons
                       + '). Pick the bottle from inventory below instead.')
            _icl_scan_manual_panel(record, storage_options, target, move, scope, saved)
            return
        st.caption('Select a product and confirm the label match to continue.')
        return
    # Inventory lookups are cached per photo and candidate; saving re-checks the bottle under a row lock.
    match_cache = record.setdefault('_matches', {})
    match_key = _icl_result_key(result) + ':' + str(candidate.get('candidate_id'))
    if match_key not in match_cache:
        match_cache[match_key] = _icl_exact_bottle_matches(candidate)
    matches = match_cache[match_key]
    action_options = ['Existing bottle · Add an observation', 'New bottle · Create a Bottle ID']
    action = st.radio('How would you like to handle this bottle?', action_options,
        index=0 if matches else 1,
        key=scope + '_' + _icl_result_key(result) + '_' + candidate['candidate_id'] + '_action')
    if action.startswith('New bottle'):
        label_fields = result.get('fields') or {}
        meta = result.get('metadata') or {}
        lots = [x.get('value') for x in label_fields.get('lot') or []] + [meta.get('lot')]
        quantity = {}
        for pack in label_fields.get('package') or []:
            m = re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*(g|mg|kg|mL|ml|L|l)\s*', str(pack.get('value', '')))
            if m:
                unit = {'ml': 'mL', 'l': 'L'}.get(m.group(2), m.group(2))
                quantity = {'amount': float(m.group(1)), 'unit': unit}
                break
        prefill = {'name': candidate.get('name'), 'cas': candidate.get('cas'),
                   'manufacturer': candidate.get('manufacturer') or candidate.get('observed_manufacturer'),
                   'catalog': candidate.get('catalog') or candidate.get('observed_catalog'),
                   'lot': next((clean_text(v) for v in lots if clean_text(v)), ''),
                   'formula': candidate.get('formula'), 'mw': candidate.get('mw'), **quantity}
        _icl_scan_register_form(record, storage_options, target, scope, prefill, candidate, saved)
        return
    if not matches:
        st.info('No bottle with this chemical was found automatically. Search the inventory below, or register a new bottle.')
        _icl_scan_manual_panel(record, storage_options, target, move, scope, saved)
        return
    st.caption('Several bottles may contain the same chemical. Check the physical Bottle ID before linking this photo.')
    mapping = {m['id']: m for m in matches}
    lot_pick = auto.get('bottle') if auto.get('candidate') else None
    ids = list(mapping)
    default_index = 0 if len(ids) == 1 else (ids.index(lot_pick['id']) if lot_pick and lot_pick['id'] in mapping else None)
    if lot_pick and lot_pick['id'] in mapping and len(ids) > 1:
        st.caption('Lot number on the label matches ' + str(lot_pick['bottle_id']) + '.')
    bottle_id = st.selectbox('Select the physical bottle', ids, index=default_index,
        format_func=lambda i: _icl_bottle_option(mapping[i]),
        key=scope + '_' + candidate['candidate_id'] + '_bottle')
    if bottle_id is None:
        return
    bottle = mapping[bottle_id]
    key = scope + '_' + _icl_result_key(result) + '_' + candidate['candidate_id'] + '_' + str(bottle_id)
    default_observed, _ = _icl_scan_location_for(bottle, storage_options, target, move)
    observed = st.selectbox('Observed storage location', storage_options,
        index=storage_options.index(default_observed), format_func=lambda x: x[2], key=key + '_storage')
    move_here = st.checkbox("Also update the bottle's current storage location", value=True, key=key + '_move')
    unchanged = (observed[0], observed[1]) == (bottle['storage_unit_id'], bottle['shelf_number'])
    move_here = bool(move_here and not unchanged)
    st.write(f"**{bottle['bottle_id']}** · " + (('Move to ' if move_here else 'Observed at ') + observed[2]))
    if st.button('Confirm Bottle ID and save', type='primary', key=key + '_save'):
        try:
            record['saved_observation_id'] = _icl_scan_save(record, candidate, bottle, observed, move_here)
            record['resolved_action'] = 'existing_bottle_observation'
        except Exception as exc:
            _icl_scan_forget(record)
            st.error('Could not save: ' + _icl_exception_message(exc))
        else:
            saved()


def _icl_scan_analyze(items, previous):
    """Analyze photos concurrently; the OCR request helper never touches session state."""
    old = {r['scan_id']: r for r in previous}
    batch = [old.get(item['scan_id']) or dict(item) for item in items]
    todo = [r for r in batch if not r.get('result') and not r.get('saved_observation_id')]
    progress = st.progress(0.0, text=f'Analyzing 0 / {len(todo)}')
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(_ocr_api_request_bytes, r['filename'], r['bytes'],
                                   r.get('mime') or 'application/octet-stream', True): r for r in todo}
        for done, future in enumerate(as_completed(futures), start=1):
            record = futures[future]
            try:
                record['result'] = future.result()
                record.pop('error', None)
                record.pop('_auto', None)
            except Exception as exc:
                record['error'] = _icl_exception_message(exc)
            progress.progress(done / len(todo), text=f'Analyzing {done} / {len(todo)}')
    return batch


def _icl_storage_scan():
    st.header('Storage Scan')
    _icl_ocr_connection_panel('scan')
    st.caption('Read labels → Choose a location once → Save matched bottles together → Review the rest')
    if not get_ocr_api_key():
        st.error('OCR is not configured. Ask an administrator to set OCR_API_KEY.')
    uploads = st.file_uploader('Bottle photos or ZIP', type=['jpg', 'jpeg', 'png', 'webp', 'zip'],
        accept_multiple_files=True, key='icl_scan_uploads')
    st.caption('Each photo is one review item. To combine multiple views of one bottle, use Add Chemical.')
    items, warnings = _icl_expand_uploads_cached(uploads, '_icl_scan_upload_cache')
    for warning in warnings:
        if warning.startswith('Image limit reached'):
            st.warning(warning + ' Only ' + str(len(items)) + ' photos are in this batch; upload the rest as a second batch after finishing this one.')
    if warnings:
        with st.expander('Upload notes'):
            for warning in warnings:
                st.write(warning)
    signature = hashlib.sha256('|'.join(i['scan_id'] for i in items).encode()).hexdigest()
    previous = st.session_state.get('inventory_scan_batch') or []
    if items and previous and st.session_state.get('icl_scan_upload_signature') != signature:
        st.info('Your upload list has changed. Analyze photos to update the review queue.')
    if st.button('Analyze photos', type='primary', disabled=not items or not get_ocr_api_key()):
        st.session_state['inventory_scan_batch'] = _icl_scan_analyze(items, previous)
        st.session_state['icl_scan_upload_signature'] = signature
        st.session_state.pop('icl_scan_selected', None)
        st.session_state.pop('icl_scan_quick_failures', None)
        st.rerun()
    batch = st.session_state.get('inventory_scan_batch') or []
    if not batch:
        st.info('Upload photos to review product details and link them to your inventory.')
        return
    notice = st.session_state.pop('icl_scan_notice', None)
    if notice:
        saved, failures = notice
        if saved:
            st.success(f'Saved {saved} photo(s).')
        for failure in failures:
            st.error('Not saved · ' + failure)
    done = sum(_icl_scan_done(r) for r in batch)
    st.progress(done / len(batch), text=f'Completed {done} / {len(batch)}')
    if done == len(batch):
        st.success('Every photo in this batch is complete.')
    storage_options = get_storage_options()
    mode = st.radio('Mode', ['Simple · one photo at a time', 'Advanced · table and full review'],
        horizontal=True, key='icl_scan_mode')
    target, move = _icl_scan_target_location(storage_options, show_move=not mode.startswith('Simple'))
    if mode.startswith('Simple'):
        throw_like = target is not None and re.search(r'throw|dispos|폐기', str(target[2]), re.I)
        pending_status = st.checkbox('Also mark these bottles as Disposal Pending', value=bool(throw_like),
            key='icl_scan_simple_pending_' + str(target[0] if target else ''))
        _icl_scan_simple(batch, storage_options, target, 'Disposal Pending' if pending_status else None)
        return
    _icl_scan_quick_table(batch, storage_options, target, move)
    _icl_scan_review_one(batch, storage_options, target, move)


def _icl_add_photos():
    st.header('Add Chemical')
    _icl_ocr_connection_panel('add')
    st.caption('Review the label, apply the product details, then enter the amount and storage location.')
    if st.session_state.get('icl_saved_notice'):
        st.success(st.session_state['icl_saved_notice'])
        if st.button('Return to Storage Scan'):
            st.session_state['_pending_navigation'] = 'Inventory Scan'
            st.rerun()
    mode = st.radio('Entry method', ['Photo entry', 'Manual entry'], horizontal=True, key='icl_add_mode')
    if st.session_state.get('icl_previous_add_mode') != mode:
        if st.session_state.get('icl_previous_add_mode') is not None:
            _icl_clear_add_identity()
            st.session_state.pop('icl_scan_handoff', None)
        st.session_state['icl_previous_add_mode'] = mode
        if not st.session_state.get('icl_scan_handoff'):
            st.session_state.pop('icl_add_context', None)
    if mode == 'Manual entry':
        st.caption('Enter the chemical name and CAS exactly as shown on the label.')
        return None, [], True
    if st.session_state.get('icl_scan_handoff'):
        result = st.session_state.get('add_ocr_result') or {}
        records = st.session_state.get('icl_applied_photos') or []
        st.info('Product details from Storage Scan are ready. Review the registration form below.')
        if st.button('Start with another photo'):
            st.session_state.pop('icl_scan_handoff', None)
            st.session_state.pop('icl_add_batch', None)
            _icl_clear_add_identity()
            st.rerun()
    else:
        nonce = st.session_state.get('icl_add_upload_nonce', 0)
        uploads = st.file_uploader('Bottle photos or ZIP', type=['jpg', 'jpeg', 'png', 'webp', 'zip'],
            accept_multiple_files=True, key='icl_add_upload_' + str(nonce))
        with st.expander('Take a photo'):
            camera = st.camera_input('Photograph the front of the label', key='icl_add_camera_' + str(nonce))
        items, warnings = _icl_expand_uploads_cached(list(uploads or []) + ([camera] if camera else []), '_icl_add_upload_cache')
        if warnings:
            with st.expander('Upload notes'):
                for warning in warnings:
                    st.write(warning)
        groups = {}
        if items:
            grouping = st.radio('How are these photos grouped?', ['Each photo shows a different bottle', 'All photos show the same bottle', 'Assign groups manually'],
                key='icl_add_grouping', horizontal=False)
            if grouping == 'Assign groups manually':
                with st.expander('Assign the same group number to photos of one bottle', expanded=True):
                    for i, item in enumerate(items):
                        groups[item['scan_id']] = int(st.number_input(item['filename'], min_value=1, max_value=120,
                            value=i + 1, step=1, key='icl_group_' + item['scan_id']))
            else:
                groups = {item['scan_id']: (1 if grouping.startswith('All photos') else i + 1) for i, item in enumerate(items)}
            st.caption(f'{len(items)} photos · {len(set(groups.values()))} bottles to review')
        too_many = any(list(groups.values()).count(g) > 12 for g in set(groups.values()))
        if too_many:
            st.warning('A bottle can have up to 12 photos.')
        if st.button('Analyze photos', type='primary', disabled=not items or too_many or not get_ocr_api_key()):
            batch = []
            progress = st.progress(0.0)
            for i, item in enumerate(items):
                record = dict(item, group_id=groups[item['scan_id']])
                try:
                    with st.spinner(f'{i + 1}/{len(items)} · {item["filename"]} · Analyzing'):
                        record['result'] = analyze_scan_memory_item(item, enable_pubchem=True)
                except Exception as exc:
                    record['error'] = _icl_exception_message(exc)
                batch.append(record)
                progress.progress((i + 1) / len(items))
            st.session_state['icl_add_batch'] = batch
            st.session_state.pop('icl_add_group_selected', None)
            st.session_state.pop('icl_saved_notice', None)
            st.rerun()
        batch = st.session_state.get('icl_add_batch') or []
        if not batch:
            st.caption('To register without a photo, select Manual entry above.')
            return None, [], False
        group_ids = sorted(set(r['group_id'] for r in batch))
        group_id = st.selectbox('Bottle to review', group_ids, key='icl_add_group_selected',
            format_func=lambda g: f'Bottle {g} · {sum(r["group_id"] == g for r in batch)} photos')
        records = [r for r in batch if r['group_id'] == group_id]
        if any(r.get('error') for r in records):
            for r in records:
                if r.get('error'):
                    st.error(r['filename'] + ': ' + r['error'])
            if st.button('Retry failed photos'):
                for r in records:
                    if r.get('error'):
                        try:
                            r['result'] = analyze_scan_memory_item(r, enable_pubchem=True)
                            r.pop('error', None)
                        except Exception as exc:
                            r['error'] = _icl_exception_message(exc)
                st.rerun()
            return None, [], False
        merged = __import__('icl_ocr_ui_guard').merge_group(globals(), records, group_id)
        if merged.get('error'):
            st.error('Could not resolve the photo group: ' + merged['error'])
            return None, [], False
        result = merged.get('result') or {}
    context = _icl_result_key(result)
    if st.session_state.get('icl_add_context') != context:
        _icl_clear_add_identity()
        st.session_state['icl_add_context'] = context
    if records:
        _icl_show_label(records[0], 'add_' + str(records[0].get('scan_id', '')))
        if len(records) > 1:
            with st.expander(f'Additional views of this bottle: {len(records) - 1} photos'):
                for record in records[1:]:
                    st.image(record['bytes'], caption=record['filename'], width=360)
    candidate, ready = _icl_review_product(result, 'add')
    if st.button('Apply selected product to form', type='primary', disabled=not ready):
        load_add_form_from_ocr(result, records[0]['bytes'] if records else None,
            records[0]['filename'] if records else 'bottle.jpg', selected_candidate=candidate, evidence_photos=records)
        st.session_state.pop('_pending_navigation', None)
        st.session_state['icl_applied_result_key'] = context
        st.session_state['icl_applied_photos'] = list(records)
        st.rerun()
    applied = st.session_state.get('icl_applied_result_key') == context
    if applied:
        fields = (st.session_state.get('icl_v8_registration_draft') or {}).get('fields') or {}
        st.success('Applied to form: ' + str(fields.get('name', '')))
        if candidate and candidate.get('name') != fields.get('name'):
            st.warning('You selected a different candidate. Apply it to update the form.')
    return result, records, applied


def _icl_register_new_bottle(values, storage, photos, draft, source_scan_id):
    # Bottle, audit entry and all photographs share one transaction.
    conn = get_connection()
    try:
        columns = list(values)
        cursor = conn.execute('INSERT INTO chemicals (' + ', '.join(columns) + ') VALUES ('
            + ', '.join('?' for _ in columns) + ')', tuple(values[k] for k in columns))
        chemical_id = cursor.lastrowid
        bottle_id = f'ICL-{chemical_id:06d}'
        conn.execute('UPDATE chemicals SET bottle_id=?, registered_by=? WHERE id=?',
            (bottle_id, CURRENT_USER['email'], chemical_id))
        write_audit_log('ADD_CHEMICAL', target_type='chemical', target_id=chemical_id, bottle_id=bottle_id,
            details=json.dumps({'chemical_name': values['chemical_name'], 'cas_number': values['cas_number'],
                'location': storage[2], 'ocr_selected_fields': (draft or {}).get('fields'),
                'source_scan_id': source_scan_id}, ensure_ascii=False), conn=conn)
        for photo in photos:
            evidence = dict(photo.get('result') or {})
            evidence['_registration_selection'] = (draft or {}).get('fields') or {}
            save_inventory_observation(chemical_id, photo['filename'], photo['bytes'],
                decode_ocr_label_crop(evidence), evidence, storage[0], storage[1], _conn=conn)
        conn.commit()
        return bottle_id
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def load_add_form_from_ocr(result, original_bytes, filename, selected_candidate=None, source_scan_id=None, evidence_photos=None):
    return __import__("icl_ocr_ui_guard").apply_result(globals(), result, original_bytes, filename, selected_candidate, source_scan_id, evidence_photos)


def load_add_form_from_existing_bottle(existing_bottle, result, original_bytes, filename, source_scan_id=None):
    return __import__("icl_ocr_ui_guard").apply_result(globals(), result, original_bytes, filename, source_scan_id=source_scan_id)


def ocr_candidate_label(candidate):
    return __import__("icl_ocr_ui_guard").candidate_label(candidate)


def ocr_selected_candidate(result):
    selected_id = result.get("selected_candidate_id")
    if not selected_id:
        return None
    for candidate in result.get("candidates", []) or []:
        if candidate.get("candidate_id") == selected_id:
            return candidate
    return None


def decode_ocr_label_crop(result):
    encoded = str(result.get("label_crop_base64") or "").strip()
    if not encoded:
        return None
    try:
        return base64.b64decode(encoded)
    except Exception:
        return None


def _image_bytes(value):
    if value is None:
        return None
    if isinstance(value, memoryview):
        return value.tobytes()
    return bytes(value)


def compress_observation_image(raw_bytes, max_size=(1600, 1600), quality=82):
    if not raw_bytes:
        return None
    try:
        image = Image.open(BytesIO(raw_bytes)).convert("RGB")
        image.thumbnail(max_size)
        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=quality, optimize=True)
        return buffer.getvalue()
    except Exception:
        return raw_bytes


def save_inventory_observation(
    chemical_id,
    original_file_name,
    original_image_data,
    label_image_data,
    ocr_result,
    observed_storage_unit_id,
    observed_shelf_number,
    _conn=None,
):
    now = datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y-%m-%d %H:%M:%S")
    original_image_data = compress_observation_image(original_image_data)
    label_image_data = compress_observation_image(
        label_image_data,
        max_size=(1400, 1400),
        quality=86,
    )

    conn = _conn if _conn is not None else get_connection()
    cursor = conn.execute(
        """
        INSERT INTO inventory_observations (
            chemical_id,
            original_file_name,
            original_image_data,
            label_image_data,
            ocr_result_json,
            observed_storage_unit_id,
            observed_shelf_number,
            captured_at,
            uploaded_by
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            chemical_id,
            original_file_name,
            original_image_data,
            label_image_data,
            json.dumps(ocr_result or {}, ensure_ascii=False),
            observed_storage_unit_id,
            observed_shelf_number,
            now,
            CURRENT_USER["email"],
        ),
    )
    observation_id = cursor.lastrowid

    chemical = conn.execute(
        "SELECT bottle_id, chemical_name FROM chemicals WHERE id=?",
        (chemical_id,),
    ).fetchone()

    write_audit_log(
        "INVENTORY_SCAN",
        target_type="chemical",
        target_id=chemical_id,
        bottle_id=chemical["bottle_id"] if chemical else None,
        details=json.dumps(
            {
                "observation_id": observation_id,
                "file_name": original_file_name,
                "storage_unit_id": observed_storage_unit_id,
                "shelf_number": observed_shelf_number,
            },
            ensure_ascii=False,
        ),
        conn=conn,
    )
    if _conn is None:
        conn.commit()
        conn.close()
    return observation_id


def get_inventory_observation(observation_id):
    conn = get_connection()
    row = conn.execute(
        """
        SELECT
            io.*,
            su.storage_name,
            c.bottle_id,
            c.chemical_name
        FROM inventory_observations io
        JOIN chemicals c ON c.id = io.chemical_id
        LEFT JOIN storage_units su ON su.id = io.observed_storage_unit_id
        WHERE io.id=?
        """,
        (observation_id,),
    ).fetchone()
    conn.close()
    return row


def get_inventory_observations(chemical_id, limit=8):
    conn = get_connection()
    rows = conn.execute(
        """
        SELECT
            io.*,
            su.storage_name,
            c.bottle_id,
            c.chemical_name
        FROM inventory_observations io
        JOIN chemicals c ON c.id = io.chemical_id
        LEFT JOIN storage_units su ON su.id = io.observed_storage_unit_id
        WHERE io.chemical_id=?
        ORDER BY io.id DESC
        LIMIT ?
        """,
        (chemical_id, limit),
    ).fetchall()
    conn.close()
    return rows


def render_observation_card(observation, expanded=True):
    if not observation:
        return

    storage_text = observation["storage_name"] or "Not Assigned"
    if observation["observed_shelf_number"] is not None:
        storage_text += f" / Shelf {observation['observed_shelf_number']}"

    with st.expander(
        f"📷 Scan #{observation['id']} · {observation['captured_at']} · {storage_text}",
        expanded=expanded,
    ):
        image_col1, image_col2 = st.columns(2)

        original = _image_bytes(observation["original_image_data"])
        label = _image_bytes(observation["label_image_data"])

        with image_col1:
            if original:
                st.image(original, caption="Original bottle", width="stretch")
            else:
                st.caption("Original image unavailable")

        with image_col2:
            if label:
                st.image(label, caption="Detected label", width="stretch")
            else:
                st.caption("Label crop unavailable")

        st.caption(
            f"Bottle: {observation['bottle_id'] or '-'} · "
            f"Uploaded by {observation['uploaded_by'] or '-'}"
        )

        try:
            ocr_data = json.loads(observation["ocr_result_json"] or "{}")
        except Exception:
            ocr_data = {}

        if ocr_data:
            with st.expander("OCR evidence"):
                st.json(
                    {
                        "status": ocr_data.get("status"),
                        "metadata": ocr_data.get("metadata"),
                        "stages_run": ocr_data.get("stages_run"),
                        "elapsed_sec": ocr_data.get("elapsed_sec"),
                    }
                )


def render_bottle_scan_history(chemical_id, limit=8):
    rows = get_inventory_observations(chemical_id, limit=limit)
    if not rows:
        st.info("No physical bottle scan has been saved for this bottle yet.")
        return
    st.caption(f"{len(rows)} most recent scan observation(s)")
    for observation in rows:
        render_observation_card(observation, expanded=False)


def find_existing_bottles(catalog="", cas="", chemical_name="", manufacturer=""):
    catalog = clean_text(catalog)
    cas = clean_text(cas)
    chemical_name = clean_text(chemical_name)
    manufacturer = clean_text(manufacturer)

    conn = get_connection()

    if catalog:
        rows = conn.execute(
            f"""
            SELECT c.*, {LOCATION_EXPR} AS location_text
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed'
              AND LOWER(TRIM(COALESCE(c.catalog_number,'')))=LOWER(?)
            ORDER BY c.id
            """,
            (catalog,),
        ).fetchall()
        if rows:
            conn.close()
            return rows

    if cas:
        rows = conn.execute(
            f"""
            SELECT c.*, {LOCATION_EXPR} AS location_text
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed'
              AND TRIM(COALESCE(c.cas_number,''))=?
            ORDER BY c.id
            """,
            (cas,),
        ).fetchall()
        if rows:
            conn.close()
            return rows

    if chemical_name:
        like = f"%{chemical_name}%"
        params = [like]
        maker_sql = ""
        if manufacturer:
            maker_sql = " AND c.manufacturer ILIKE ? "
            params.append(f"%{manufacturer}%")
        rows = conn.execute(
            f"""
            SELECT c.*, {LOCATION_EXPR} AS location_text
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed'
              AND c.chemical_name ILIKE ?
              {maker_sql}
            ORDER BY c.id
            LIMIT 30
            """,
            params,
        ).fetchall()
        if rows:
            conn.close()
            return rows

    # No usable Catalog / CAS / name evidence.
    # Do not surface arbitrary inventory rows as false "matches".
    conn.close()
    return []



# ICL UX v2 — inventory detail popup.
def _icl_inventory_detail_body(detail):
    st.subheader(detail["chemical_name"] or "Chemical")

    top1, top2 = st.columns([3, 1])
    with top1:
        if detail["description"]:
            st.caption(detail["description"])
    with top2:
        st.markdown(f"**{detail['status'] or '-'}**")

    latest_scans = get_inventory_observations(
        detail["id"],
        limit=1,
    )

    if latest_scans:
        latest_scan = latest_scans[0]
        original_image = _image_bytes(
            latest_scan["original_image_data"]
        )
        label_image = _image_bytes(
            latest_scan["label_image_data"]
        )

        if original_image or label_image:
            st.markdown("#### Bottle image")

            if original_image and label_image:
                image_col1, image_col2 = st.columns([1.35, 1])

                with image_col1:
                    st.image(
                        original_image,
                        caption="Bottle photo",
                        width="stretch",
                    )

                with image_col2:
                    st.image(
                        label_image,
                        caption="Detected label",
                        width="stretch",
                    )
            elif original_image:
                st.image(
                    original_image,
                    caption="Bottle photo",
                    width="stretch",
                )
            else:
                st.image(
                    label_image,
                    caption="Detected label",
                    width="stretch",
                )

            storage_text = latest_scan["storage_name"] or "Not Assigned"
            if latest_scan["observed_shelf_number"] is not None:
                storage_text += (
                    f" / Shelf {latest_scan['observed_shelf_number']}"
                )

            st.caption(
                f"Latest scan · {latest_scan['captured_at']} · {storage_text}"
            )
    else:
        st.caption("No bottle image has been saved for this bottle yet.")

    st.markdown("#### General Information")
    c1, c2, c3 = st.columns(3)

    with c1:
        st.caption("Bottle ID")
        st.write(detail["bottle_id"] or "-")
        st.caption("CAS No.")
        st.write(detail["cas_number"] or "-")
        st.caption("Legacy ID")
        st.write(detail["legacy_id"] or "-")

    with c2:
        st.caption("Manufacturer")
        st.write(detail["manufacturer"] or "-")
        st.caption("Catalog No.")
        st.write(detail["catalog_number"] or "-")
        st.caption("Lot No.")
        st.write(detail["lot_number"] or "-")

    with c3:
        st.caption("Storage")
        st.write(detail["location_text"] or "-")
        st.caption("Owner")
        st.write(detail["owner"] or "-")
        st.caption("Storage temperature")
        st.write(detail["storage_temperature"] or "-")

    st.markdown("#### Inventory")
    i1, i2, i3, i4 = st.columns(4)
    i1.metric(
        "Initial",
        f"{detail['initial_amount']} {detail['unit'] or ''}"
        if detail["initial_amount"] is not None else "-",
    )
    i2.metric(
        "Remaining",
        f"{detail['remaining_amount']} {detail['unit'] or ''}"
        if detail["remaining_amount"] is not None else "-",
    )
    i3.metric("Purchased", detail["purchase_date"] or "-")
    i4.metric("Expires", detail["expiration_date"] or "-")

    with st.expander("Chemical information", expanded=False):
        p1, p2, p3 = st.columns(3)
        with p1:
            st.caption("Molecular formula")
            st.write(detail["molecular_formula"] or "-")
            st.caption("Molecular weight")
            st.write(detail["molecular_weight"] or "-")
        with p2:
            st.caption("Purity")
            st.write(detail["purity"] or "-")
            st.caption("Concentration / density")
            st.write(detail["concentration_density"] or "-")
        with p3:
            st.caption("Opened")
            st.write(detail["opened_date"] or "-")
            st.caption("Registered")
            st.write(
                detail["source_registered_date"]
                or detail["created_at"]
                or "-"
            )

    with st.expander("Safety", expanded=False):
        render_safety_panel(
            signal_word=detail["safety_signal_word"],
            h_codes=detail["safety_h_codes"],
            hazard_statements=detail["safety_hazard_statements"],
            pictograms=detail["safety_pictograms"],
            source_url=detail["safety_source_url"],
            checked_at=detail["safety_checked_at"],
        )

    if detail["comments"]:
        with st.expander("Notes", expanded=False):
            st.write(detail["comments"])

    with st.expander("Physical bottle / scan history", expanded=False):
        render_bottle_scan_history(detail["id"], limit=8)


if hasattr(st, "dialog"):
    @st.dialog("Chemical inventory", width="large")
    def _icl_inventory_dialog(detail):
        _icl_inventory_detail_body(detail)
else:
    def _icl_inventory_dialog(detail):
        _icl_inventory_detail_body(detail)


# ICL UX v3 — storage detail popup.
def _icl_storage_detail_body(storage_id):
    conn = get_connection()
    storage = conn.execute(
        """
        SELECT id, storage_name, notes
        FROM storage_units
        WHERE id = ?
        """,
        (storage_id,),
    ).fetchone()
    shelves = conn.execute(
        """
        SELECT shelf_number
        FROM storage_shelves
        WHERE storage_unit_id = ?
        ORDER BY shelf_number
        """,
        (storage_id,),
    ).fetchall()
    conn.close()

    if not storage:
        st.error("Storage unit not found.")
        return

    st.subheader(storage["storage_name"])
    if storage["notes"]:
        st.caption(storage["notes"])

    shelf_options = ["All Locations", "Unassigned"] + [
        f"Shelf {shelf['shelf_number']}" for shelf in shelves
    ]
    selected_shelf = st.selectbox(
        "Shelf",
        shelf_options,
        key=f"storage_dialog_shelf_{storage_id}",
    )

    conditions = ["c.storage_unit_id = ?", "c.status != 'Disposed'"]
    params = [storage_id]
    if selected_shelf == "Unassigned":
        conditions.append("c.shelf_number IS NULL")
    elif selected_shelf.startswith("Shelf "):
        shelf_number = int(selected_shelf.replace("Shelf ", ""))
        conditions.append("c.shelf_number = ?")
        params.append(shelf_number)

    where_sql = " AND ".join(conditions)
    conn = get_connection()
    chemicals = conn.execute(
        f"""
        SELECT
            c.id,
            c.bottle_id AS "Bottle ID",
            c.chemical_name AS "Chemical Name",
            c.cas_number AS "CAS No.",
            c.manufacturer AS "Manufacturer",
            COALESCE(c.remaining_amount, c.amount) AS "Remaining",
            c.unit AS "Unit",
            CASE
                WHEN c.shelf_number IS NOT NULL THEN 'Shelf ' || c.shelf_number
                ELSE 'Unassigned'
            END AS "Shelf",
            c.status AS "Status"
        FROM chemicals c
        WHERE {where_sql}
        ORDER BY c.shelf_number, c.chemical_name
        """,
        params,
    ).fetchall()
    conn.close()

    st.caption(f"{len(chemicals):,} active bottle(s)")
    if chemicals:
        display_data = []
        for chemical in chemicals:
            remaining = chemical["Remaining"]
            try:
                value = float(remaining)
                if value.is_integer():
                    value = int(value)
                remaining_text = f"{value} {chemical['Unit'] or ''}".strip()
            except Exception:
                remaining_text = (
                    f"{remaining} {chemical['Unit'] or ''}".strip()
                    if remaining is not None else "-"
                )
            display_data.append(
                {
                    "Bottle ID": chemical["Bottle ID"],
                    "Chemical Name": chemical["Chemical Name"],
                    "CAS No.": chemical["CAS No."],
                    "Manufacturer": chemical["Manufacturer"],
                    "Remaining": remaining_text,
                    "Shelf": chemical["Shelf"],
                    "Status": chemical["Status"],
                }
            )
        st.caption("Click a row to see the bottle's photo and details below.")
        event = st.dataframe(
            display_data,
            width="stretch",
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            key=f"storage_dialog_table_{storage_id}_{selected_shelf}",
        )
        picked = event.selection.rows if event else []
        if picked:
            conn = get_connection()
            try:
                detail = conn.execute(
                    f"""
                    SELECT c.*, {LOCATION_EXPR} AS location_text
                    FROM chemicals c
                    LEFT JOIN storage_units su ON c.storage_unit_id = su.id
                    WHERE c.id = ?
                    """,
                    (chemicals[picked[0]]["id"],),
                ).fetchone()
            finally:
                conn.close()
            if detail:
                st.divider()
                # A dialog cannot open another dialog, so the detail renders here.
                _icl_inventory_detail_body(detail)
    else:
        st.info("No chemicals found in this location.")

    if st.button("Add shelf", key=f"storage_dialog_add_shelf_{storage_id}"):
        conn = get_connection()
        max_shelf = conn.execute(
            """
            SELECT MAX(shelf_number)
            FROM storage_shelves
            WHERE storage_unit_id = ?
            """,
            (storage_id,),
        ).fetchone()[0]
        next_shelf = (max_shelf or 0) + 1
        conn.execute(
            """
            INSERT INTO storage_shelves (storage_unit_id, shelf_number)
            VALUES (?, ?)
            """,
            (storage_id, next_shelf),
        )
        write_audit_log(
            "ADD_SHELF",
            target_type="storage_unit",
            target_id=storage_id,
            details=f"Added Shelf {next_shelf} to {storage['storage_name']}",
            conn=conn,
        )
        conn.commit()
        conn.close()
        st.rerun()


if hasattr(st, "dialog"):
    @st.dialog("Storage", width="large")
    def _icl_storage_dialog(storage_id):
        _icl_storage_detail_body(storage_id)
else:
    def _icl_storage_dialog(storage_id):
        _icl_storage_detail_body(storage_id)


# ICL UX v3.6 — in-memory camera upload adapter.
class _ICLCameraUpload:
    def __init__(self, record):
        self.name = record["name"]
        self.type = record.get("mime_type") or "image/jpeg"
        self.group_id = int(record.get("group_id") or 1)
        self._data = record["bytes"]

    def getvalue(self):
        return self._data


# ICL UX v4.0 — native one-step browser back buffer.
#
# Streamlit writes query-parameter changes to the top-level app URL. We change
# one harmless parameter exactly once per Streamlit session, which gives mobile
# browsers one same-app history entry before the page that opened this app.
#
# Important: after Back removes/reverts this parameter, session_state remains
# seeded, so we DO NOT immediately add it again and trap the user in the app.
if not st.session_state.get("_icl_back_guard_seeded", False):
    _icl_existing_guard = str(
        st.query_params.get("_icl_guard", "")
        or ""
    )
    st.session_state["_icl_back_guard_seeded"] = True

    # Always change the value, even if somebody bookmarked a guarded URL. That
    # guarantees a fresh same-app history entry for this browser session.
    st.query_params["_icl_guard"] = (
        "2"
        if _icl_existing_guard == "1"
        else "1"
    )


# ============================================================
# HEADER / SIDEBAR
# ============================================================

st.markdown(
    """
<div class="icl-shell-header">
  <div class="icl-shell-copy">
    <div class="icl-shell-eyebrow">ICL · INORGANIC CHEMISTRY LABORATORY</div>
    <div class="icl-shell-title">Chemical Manager</div>
  </div>
  <div class="icl-shell-pills" aria-label="Workspace areas">
    <span>Inventory</span><span>Storage</span><span>Safety</span>
  </div>
</div>
    """,
    unsafe_allow_html=True,
)

# Render the original repository PNG directly. Streamlit preserves the aspect
# ratio, so the laboratory logo is never vertically compressed.
st.sidebar.image(
    BytesIO(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAABDMAAAFeCAYAAACYZVSYAABbg0lEQVR42u3df2wcd3rn+c9T4oyNswzpsFioq3W3ZNaKk92bBTl3i90LxDlxboGRQhXP7bUCCwg24lw0GMxaG1OzOxjdwbBkx9g4mM2Is2fvIBgFphIEsC5yTIMtRhpgb2gMdTnsDZAmbv7I+qQLiYW7W38cIMHywfZQ9dwf3ZJJieyq/sn+8X4B3s3Yxeqqp75VXc/T3x8SAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAKCdjBAA2CQTTTzy7wLbG8jH0u4ijm1V8tVH/kM5v0SAAQAAADSLYgbQ7zK5EWl9ZGNBwt3HJNtbfQqMybVnB54+d+QqSJJZpfgRu9+WW0EShQ8AAAAANdIJAL2vWrAIAk3ItdelMUkjkoZ7/AlVLXj4bTMrxLKC7N6qPlwscNEBAACAwUUxA+g1mWhC5mOBBWPuPiLp0IBGYsXMCpKvxrGW9PhQQavzt2kgAAAAQP+jmAF0s/2TY4ptolq4GJM0SlBqWjGzQhxrSbvuFejBAQAAAPQnihlAN6kWL0yakNnEjsxl0V9PuDtyX3JpSf6FeZXnVwkKAAAA0A+v+gB2TiY3IvtFjuJFx6yZ2VIszau4ME84AAAAgN5EMQPotP2TY0EcTLs0IYaN7DB/zz2Y1+O75plvAwAAAOgdicWMIBudc9dZQvVQ4EyvxMX8OSKBVLJTuUDKuTxH74tuRWFjUPE9x/ccAADoPUOEAGiT/ZNjge+aqRQwfI8TkW5P3Z4x82f06fpbCo9WChvlhTniAgAAAHQfihlAK2VyI0GwPu2uacUadlHC6E3VwkY2mjXZfGz3ZlkZBQAAAOgeFDOAVshMTZv5tLR+yKlf9A/XHpefMA9OKIxW3G2WYSgAAADAzqOYATTqfi8MaUbuzIPR/0bN/C19tj5rYTQX+9AsS70CAAAAOyMgBECdMtFEkJ2aM1v/G3edZULPAePa49KLZut/E2Sn5rR/coygAAAAAJ1FMQNIKxNNWBgtmekn7n6CgMDdT1gc/JWF0ZIy0QQRAQAAADqDYgaQJDM1bWG0aqafSDpEQLCFQ2b6CUUNAAAAoDMoZgDbeVDE8LckDRMQpEBRAwAAAOgAJgAFHpadypn7rOQUMNCoQ2b6icLoffehaSYKBQAAAFqLnhnAfffnxHB/V/TEQGscejBR6EhuL+EAAAAAWoNiBpDJjVh4dJ45MdAu7n7CPltfVXh0hmgAAAAAzaOYgcE1ktsbZKNzZut/I9kzBARt5dpjsvMWRqvMpwEAAAA0h2IGBlN2KmefrhfcdZZgoMOGzfQThp4AAAAAjaOYgcGSyY0wLwa6wYOhJ9mpHNEAAAAA6kMxA4MjPDpjwXpBzIuBbuHaY+7vVpZyzY0QEAAAACAdihnof/snxyyMCiY7L9ceAoIudMiC9QK9NAAAAIB0KGagv4VHZywO/krSKMFAV3vQS+PoPHNpAAAAALVRzEB/uj83huw8wUBvsWfs0/UCK54AAAAA26OYgf6TncoxNwZ6XHXFk+gcoQAAAAAeRTED/WMktzfITs2Z+7vMjYF+4K6zFkZLDDsBAAAANqOYgf6wf3LMPl1fcvcTBAN95pB9tr6q/ZNjhAIAAACooJiB3pedypkHS2KST/Qr1x6Lg79SZmqaYAAAAAAUM9DrDTiMZhlWgkFh5m8F2ak5IgEAAIBBN0QI0JNGcnvt01/MufQMwcAgcfcTFkYj/thQTqvzt4kIAAAABhE9M9B7qvNjSEYhA4PqkH26vsQ8GgAAABhUFDPQW/ZPjjE/BiBJGjUPKGgAAABgIFHMQO/ITE1bHPwV82MAVa495sGSslM5ggEAAIBBQjEDvSEzNW3mbxEI4CGuPeb+LiudAAAAYJBQzED3N9Ls1ByFDKA2M3+LggYAAAAGJk8kBOjqBpqdmnP3E0QCSEZBAwAAAAOTKxICdG3jpJAB1I2CBgAAAAYiXyQE6MqGSSEDaBgFDQAAAPR9zkgI0FVGcnspZADNo6ABAACAfkYxA92VgH26Pk8hA2jR/URBAwAAAH2KYga6pzFmp+YkHSISQOtQ0AAAAEBf5o+EAF3REBlaArSNmb+l/ZNjRAIAAAB9k0MSAux4I6SQAbSdebBEQQMAAAB9k0cSAuxoA6SQAXSGa4/FwbxGcnsJBgAAAHo+lyQE2DHh0RkKGUBHDdun60sUNAAAANDrKGZgZ2Smpk12nkAAHTcafHZvljAAAACgl1HMQOdlogkzf4tAADvD3U8oPDpDJAAAANCrKGags/ZPjlmgeQIB7CyTnVcmmiASAAAA6EUUM9A5I7m95sGSXHsIBrDzLBATggIAAKAnUcxAZ4zk9tqn6xQygG7i2mOfrs8TCAAAAPQaihnoTEOrTDg4SiSArnMoyEbnCAMAAAB6KsckBGh7I8tG51iCFehe7jrL/BkAAADoqTyTEKCtslM5d50lEEB3M9Mc82cAAACgV1DMQPvsnxwz+RyBAHrCcHU4GAAAAND1KGagPUZyey0O5pjwE+gd7n5C2akckQAAAEC3o5iB9jQsJvwEepLJGW4CAACArjdECNBymalpJvzsKyuSbj+a9Kogq/x7d01skRWP0TOnB7n2BJ+un4ulGYIBAACAbkUxA621f3LM3GflhKInmO7IVbhfmIhlBcV+W48PFbQ6f7t2zpvSSG6vPlkfU2B7A/mYu49JNiJ67nQtl15UJppXOb9ENAAAANCNKGagdUZye+3T9TmJX+O71JrkBTMrxLGWNhYs2lp7qnzGkiTF0vym/5aJJmQ+ZtKEzCboydE9zDTr0hiRAAAAQDeimIGWCT5dP+f82t5F2ajumGw+jrUk7VpSeX5VUnd1mqn88r/k0qwkaf/kmGKbMFlO0iEu4o4aDbLRubiYP0coAAAA0G0oZqA1MtGESy8SiB23Yqb52OJ5fbhY6LnRPh8uFiQVXJqtDE+5lzOLc5I9w6XtPJdmNJKbTRpyBAAAAHQaxQw0byS31z5bn2eejB2zZtJ87EOzKs+v9s1lqCTQcy7NKZMbkf0iZ7IZScNc8g5x7Qk+uzcbS9MEAwAAAN2EpVnRNPv0F3PMdbADcTe76K6veik/EpfyM/eHkfSl8vyqSldmvZQfcddXJb1PC+gMdz+h/ZNjRAIAAADdhGIGmpOdyjEEoINMd8z0ivvQL8XFhemBXG2inF/yUn7CfeiXzOwijaIDzS4OZokCAAAAugnFDDRuJLfX5HMEohPZZLWI8cWhkbiYP9fXvTDSKs+vxsWFaYoaHXFImWiCMAAAAKBbUMxA443ns3uzDC9ps4eLGEzE+KgNRQ0x/KR9TdF0jigAAACga/JRQoCGZKIJdz9BINqaPFLEqEd5frUy/ERflbRGQFqO3hkAAADoGhQz0GiiPUcU2ub9ypwYFDEaUplTY8RMrxCMlt/354gCAAAAugHFDNTfaLLRObE8ZhsyRd1xs2e9lJ9gTozmxcX8OQ/iL0taIRotQ+8MAAAAdEdeSghQl0xuxF1nCUSr+Xv+xaERFRfmiUULfbhY8FJ+jF4arUPvDAAAAHQDihmoM5H5xSxRaGVAdcfdvu6lKzmGlLTPhl4azKXRPHpnAAAAYMdRzEB6mWhCsmcIRMuseDw0pvLCHKHogA8XC/7Y0Jjk7xGM5pj5DFEAAADATqKYgToSGCb9bF0s7aI/NsTcGJ22On/bS1dyLj9NMJpqwc8okxshDgAAANgpFDOQTmZqWkz62RLu9vW4uDDNsJIdVLoy62bPynSHYDT45WHrM0QBAAAAO/Y+SgiQaCS31wKfJRBNqq5WwrCSLlFcmHeLJyhoNMZN0xrJ7SUSAAAA2AkUM5DcSD5bn5FrD5FogumOWzzBaiVd5sPFgn9xaEQs31o/1x59ci9HIAAAALAjeSohQE0jub0uzRCIpqy5xRP6cLFAKLrQ6vxtf2xoQhQ06sZEoAAAANgpFDNQu4F8un6OXhlNWfHHhsYoZHQ5ChqNGtX+yTHCAAAAgI7nqoQA28rkRlx6kUA0bMUfG5pgos8eQUGjsS+ROJgmCgAAAOj4eyghwLaNI7h3jig0yHSHQkYPoqBRNzdNEwUAAAB0PF8lBNhSJjfi7icIRAPuT/ZJIaM3UdCoj2uPslM5AgEAAIBOopiBrRsGvTIac7+QwRwZvW11/rb7UI5lW1N/keSIAgAAADr8Dgo8ZCS3l14ZjXHZNIWMPlGeX3WLJyhopGj37ic0kttLJAAAANApFDPwaKP4bH2GKDSQ0MlPq7gwTyT6yIeLBY+N+yGNT+7lCAIAAAA6lrcSAmwyktvrEslbnczsokpXZolEHyovzJnZRQKR8GUSaIIoAAAAoGPvn4QAm3z6i2m59hCIuqzEX9w1Qxj6V1xcmBYTgtbk8hxRAAAAQKdQzMAmJrrU1xmwOx7E06xcMgDJehBPM39GrQCxqgkAAAA6h2IGPpeZmpY0TCDqyN9im2HCzwHx4WLB3c8RiBpfKO4TRAEAAAAdefckBLjPzKeJQj3xsosqL8wRiQFSmRflfQKxNWeJVgAAAHQIxQxU7J8ck3SIQKS2xjwZA5qwBzHXfXvDyuRGCAMAAADajWIGKg3BSczrSmhdzJMxqD5cLJj0AwKxDftFjiAAAACg7TksIUBlOVZWIkidq0k/UDm/RCQGV/zY0DkmA932/pggCgAAAGg3ihmQPrmXYznW1Nbix4bOEYYBtzp/26RZArEFswmCAAAAgHajmAGZ+QxRSMfNZhheAkmKi/lzktaIxMM3ifZU5+ABAAAA2oZixqCrJB2jBCKV91VcmCcM+Dxv91misIXYJggCAAAA2olixqA3gDiYJgopE1cfIlbY7LEvzDF3xqOYNwMAAABtz2UJwYAn6CYS9DTJmekVledXiQQ2Ye6M7W6YCYIAAACAdqKYMciyU0z8mSox0534i0MkrNhSHA/NEYWHuPYokxshEAAAAGgXihmDffFzRCFFXhYz6SdqKM+vSv4egXj4AXNvjCAAAACgjfksBjZJl+eIQqI1lRfmCANq3ksW0EYe+XLxMaIAAACA9r1vYjAxxCRdkup2jiggUXFhnolAH753mAQUAAAA7UMxY3AvfI4oJKJXBlIz2TxR2BSQMYIAAACANua0GEQMMUkRI3ploA6xNE8UNj1kmAQUAAAAbUMxYxAxxCQNemWgPsWFeYLwsPURYgAAAIB2oJgxmBc9RxRqo1cGGvQ+IdjwrAmYNwMAAABty2sxcIk6Q0xqM93R47vmCQTqbjqmJaKwKSIjxAAAAADtQDFj0OyfHGOISUL6Jc1qdf42kUC94phixkbuPkIUAAAA0A4UMwbtgnuQIwpJCenQHFFAQ8r5JYKwASuaAAAAoF25LSEYLO7Ml5EQofdUnl8lDmjCCiG4fzvRCwwAAADtQTFjkIzk9koaJRA1ci8L5ogCmmFmBaKwQSaaIAgAAABoNYoZg+STezmCUNMay2uieb5KDAAAAID2opgxSBebZRJrMtMcUUCzmASU5w4AAADab4gQDA53J6momYQy8Sda4PGhgn+y/lUCcf+5M7RKFAAAANBqFDMGRSY3Iq0PE4htrTDxJ1qisqzvEoEAAAAA2odhJgPj3gQx2J7L54gCAAAAAPQGemYMiCDQhDtx2NZjX5gjCAAAAABq5lXZqTl3PyJp3+f/1q+7279WOb9IhDp4LQjBYGC+jJrerw4NAAAAAIAtWRhddfcT2lTIkCQ7aKYryky9TJQ6h2LGIBjJ7ZXEfBnbcPk8UQAAAACwbeKcnZqTdLjWNmb+r5SNxolWh64JIRgAnzFfRk3+hXmCAAAAAGDblMH9WIrNngxkJ4lWZ1DMGIiL7GNEYVusYgIAAAAgyRNpNnIn9+oUJgAdAO6aIApbM2mJeVEBAEDL3zGy0YJc/2XNdzTT76mYv0S0AKB+FDMGI2MfExn7lmLXPFEAAAAtVylkjNbexkMCBfSMj5Wud8YNQtUZDDPpd/snx+TaQyC2YLqjcn6JQAAAAABIsJxmIzfNEqrOoJjR73zXCEHYLja+RBAAAAAAJKYOir8p6YNa25jZRRXzy0SrMyhm9P0FZgKa7R9IWiIKAAAAABKVFtdc8dckXdviv95y2bfj4sI0geoc5szo94SdyT+3F9AzAwAAAEBKpcU1l44onByW7FlJklmJiXx3BsWMfsfkn9vF5Y4+XCwQCAAAAAB1KS2uScyNsdMYZtLPRnJ7mfxzG64CQQAAAACA3kTPjH72yfqYjDBsxUxLdFgBAAB9LRuNy/0fBh5kFCijWOXY4rLMfrYjkxRWu+Y/OJ777h+X/N3qL95dF7/Y/IZK+QuNn/vRmYfPO3a/Ifldla7MdtX5mh2QbPfnxxkXdqzN1Hf8z8s9vH/8D+LbC8dez/0ivxu73+jZ86KYgXQZu4+JasaW4pjJPwEAQHcLMtHbbnq+1jYuP70pGc5G4+Z6SdKYXPskk5tLLskkk1X+7zD6WNKyS5ebStJTJJjm+g1J45L2Sfr8eB68s1aPS3ZeYXRL0rKb/qzZeQgsjP5a0q/UjF8pbw8n9OaakTS+MX6B2cVYulDf9Zt63S0el+zgVudt9//fMDov+XWXXVUp/1q9bcBcl+Jy/nijBYxAdtLdjzw4X6+0rM+P005U28z9azP7cBLdUKxbcb3C6KRJx6rX64mNx/8gvgnHnr64cHTGZOcTtvqPXsr/ahP3S6X9VaYKeGrL+0UP38deMA+W4/LCGYoZ6J8vQNkIvQ+28fhQgSAAAIDufpmzT6qZWYpEa3LYFPyhXIdT7v0JSYdNOqwwOuWmUy39lbeSZJ6Sa7TOv9wn6TlzPacw+p9ceqOJYssndYU7OzXn7v9U0pPNnfvUaZN/1+X70v+waAdNOqhsNO2xfkfl/GLqNhDYJ/Uf4+SwKfgDuZ7z9BPsPbg2lp26GPu9sxt60nzSglaTfh+ZaNJM/1pK3b42tqtrrvibbeoF1Fgc7hch79+/6ZO4JyQ76OYHLYymzW1ukIoazJnRx1waIwpbWtPq/G3CAAAA+kImmjQFP5dSFzIeNmquRWWmXm6+iDE5bGFUMOlHdSSa2x+X9CMLjy5Xuty3SfWY3f1EU4WMcHLYwqPLJv9+NXlu5AX+KTNdCbJTc21tLxb8paTnGs4z3E+Ygh8rE012vr1PvWymK020r8Om4Oc7cuxbJeTZ6A1z/bSJ+/e+fW7+XctGN5SNxgfh0Ucxo7+NEIItH78FYgAAAPqByY5UE7vdTe7qSTN/pamCxudFldEWn+XBtiWflR4KP276mB/spzKkpOm3VfcTFkZX21YIcIUt2NvTZnq7o0WBzNTLZv5KC/a0u+PHvlXLDqOr7nqhtamOnqoUJ7ujWNNOFDP62zAh2OKhYVYgCgAAoE8cbuXOzPwVhdHJ+pPMaNJMb6v5okpHk09T8IeSnm5qJ58XRJ5u9bV19/+hZXsLp063qBCw0ZPV6/5fdOY9vqXHXzn2HerFYGH081bfv4+cW58XNChm9Kv9k2MEYWtM/gkAAFArYdSrdQ3r+LyQ8WSbD60dCVrTyaTJ/lStL2Tc95+3ZC/ZaNzU8kLGg+si6W/1aHN/0lxvdPweq/S6+a/afm6BLrR1iNYOo5jRr+4FewnCdoZWiQEAAMA2XGFgu9IlvuHksAW6oPYXMj5P0EznuyVBC7LRG60aWtLW5LmSsD9J497SaEO9kRoVTp1W+3pkPHIvV4tt/ZnV0Xb7FMuybq88v0oQAABAn7plZldjjwsPEm4P/q6bPy3poFIOA3H3Y5KmE185KytipJ9/wXTTZMuKVY4tLgdu/8BNe7Rh6dYUng5s1ytxiuNrb1I6OeyuE3X+1V3JV8yCG4r9cQX2ibvvrvP86xJkpl53eb1zgtyStGxmdz8/zvhAFxZuVsy0HLvf2Nze4/+6nmM16ZTXufRuo23G5K/WdX5uV2OLy42fnx1UGJ1s6xLMFDPQ0oeW2V5nXdatvE8IAABAH/rITH8cF/OnHn4FjDclUvanKZOgJxRGL6mUf23bLbLRuDz1ihgfuPQ9FfMXfKtjU6WXg7t+Syl6ELj7CWWjCy1dTraW2B9/9H171yvVQkTK62N/fn85U9+u6BD4b7Vocs7PY2U+nXpj0013vX4/8fUtkvHqeZ/Y6fbu0rdV2tyeNrWpyvCn80o3BGhU2Wi83e2pjjaTfH6V5Y+/n+Z+Mfl0R4o1nc55ee73J3cfIwpbRuY2MQAAAP337qfjcTF/quZGpcU1L10Zl3QtXV6rmu+T5nop5dFdd8VfS/plOC7mT7nruKQPUh1f6s9P7ZaZ3nTTcS/lbeM/cTl/fIv37SMp9/tB5fosTKu0uLbt+ZcXzrjHv5b2+qQSRi8pfY+Pa17MH6h5nUqLa3FxYdpdRyV9tGOFDNfxxJ4G5fyiK/5a2vYUxBZ1IEc71rLzK+UvVO+XFNfBDvbjcq0UM/qW7SUGW0SFlUwAAED/vd9cVDm/mDqhUvxNSXdTbHpg+yR5cljpxv1/4PLfrJXEP5KAmn47ZaI83roY6k0v5TNxMX9KxfylFhYJPnLTb6e+PqXFNS/lj0h+vSXnlTpGfr3yuSmV84vpE+kWFwSkb9cVT9fpVPs1P9DWA6+0mScSr5nbv0t9fuX8orn9u5SJ//F+e/ZRzOhfI4TgUXFsq0QBAAD0kbux3ztb11+UFtfMdDHFltvOsxD4rm+lShBdp1MXMu4r5pdTJmhPVBPE5hJ+s4uJvVoeLRKMpUu87WwjQxdc/ptKV3BKkqaYcbf6efUp5xfN9Mcdbe2mm3XP/VApDKTp7XKgvYee4lqYbsblhTN15TeV7f+DpJVa/7jUd6uaMGdG/xomBFt+NawSAwAA0Eeu110skBRLb5v0QsNvVJZqSPO1enqMPJygWRhNK6H3g0ljTU4TdysuLkw38HdpEt8VlRbON3RUpcU1y06909TcFNnoeXmqngBvern+NiRVhgZZGB1TmyYv3eJ6X23keru0bJ1aQWR7ifeMxXa5ofMr5f/xID786JnRj0ZyewnCNh4fKhAEAADQL8x0o6E/rPQW+LiJrDIxmXep2ckU0/z9gQ58xlYSVwcxt6vNHFjdPW4evQC/nOpz7N4Pm2yDlzvV3jeuWlJne13qQFtKkljwiQPP81RLj2JGP/pkfYwgbGN1/jZBAAAA/aLh5K6i8b91PZXi6P6kuVzc0xQaRjvwGZtV5gtpe5Gg2uNmpeFEzyxl75HGemU8OE/p7Y41eLOfNfR36Yb6PNG24w6PzqTY6uOOrc7TJyhmYJCsEQIAAIAOaTJJlvzd7kyO7dlUiWnT5y+pmYJTqtNvweT4nUzA+zvZvyHUhWJGX15VVjLZxiohAAAAaFI2ej7FVitNf05rigE7lRy3JDE1s7s90ipucWOg42kvIejHi5pqQiYAAACgfu4hQcBDyoQAnc97gQFh6Sb+AQAAQO2XqjRDMzIDHqVBO/8D3BjoNIoZAAAAANJLNzSj+aU60w1n6TyzUkfOX5J73NYigbu3qujyBDdGs+1KuwlCfShm9CF3TRCFR8Xut4kCAABAhzRZjAhi+3KKzTo/V0MxfynVdulWsEjKcJsoZnia+TbGmj7Elpxnn0vTmynVCkHY9IwgBBgY3oLZmgEAACClmOAzkL7S3KtbPJ5is52aqyGxiGKyI019QqUY1HAPj5TL9u5TGJ1sKk+XjXM7JEg70SyFoboMEQIAAAAAdbohabTWBu46JulUg4n8uFwHExNp07LvyOn7DcmSCg2HlY3GG10xxVxfb+oQS1dmFUbnEz9HOuXShYY+I5wclnSY2yGVlaR7xmTjLs02cL88nzgxr1kpda+iHkHPDAAAAAD1pfJSIcVm+4LM1OuNJfKe6u9i6ac7cf7mwXIrz+PRIkF0siVFAtPNFFuNKpw63dju7U8l5npIFStL1Uv8cLVAVEdbmRw219smO1/zn2aLY12IYkZf3iktGPvWjx4fKhAEAACAFijlX1OKoRZu/s+ViSbrSlCy0RuSJfbKkOnmTv3SHNu9H6Z8MT9oYXS1rp1nokmTvt+atEBX023nr9R7nZSZejnVdUKlzcj/IsVmu6sFojqucfAHabZzabnfYkoxox+59hCELazO3yYIAAAArWFmaRLlJ830dtp5GYLs1Jy7Xkj1+bFd3rGTLy2uSX495daHLYyupvrFPZw6baa3JT3ZkgTa4++l3LR6nVL00Agnhy2MLpv5K9wFdagU3lZStOz0BbDM1MuSnkux5cfVAmRfoZgBAAAAoIFE+d5ZSWlWzHjSpB9ZGF3dcoWTcHI4yEy9bmFUdvcTqT7cdDMuL5zZyfN32Vwdmx82C/4yyEZvbFnUCKOXLIwKJv++WlTIkFQtuuhayq2fNPn3LYwKCqOXHvmv2Wg8yEZvmAV/mTKBxiNtRm+kbi+V67B1ETAbjVsYXU1bUDLbwcJfGzEBKAAAAICGEmXLRhfT9qSQdNhchxVGP9CDVUh8j2QjbvVN4+mu13f+/PMXFB6dTj3UwhW69IIpeEFhdKsSg8r5tzeBjr9pCn6u9HNbjJo0qjD6XX3ekyAr1992Wn0L2kx0SgkTgW64Dj9SGL0mqWxmBXf/ezL9rbqWcTWVqoXHvkPPDAAAAAANiYv5U0rVdX6TfdVkbrSRRN7MLqqUv9AN5+9mZyR91MCf7mv0/OtPoBfXzHSxwb+uXif9bVp7q9qMTtXZZvZJGq32WvpHdRUyJLnb96o9dPoOxQwMihVCAAAA0IbkTPEzkj7o0Mddi4sL011z8sX8ssu6/lfvStEp9Rwf6J82845KC+f7NZQUMzAobhMCAACANigtrrnir6n9BY1rXsof6b7zXzjv3pbk9AOZ8q3amct/s03X6Jr44bBb2szD98uxfg4jxYx+s39yjCAAAACgs8nZg4JGW5JaM7vYlYWM+8oLr7Y4Of3ATb9tsv+3y6/Rta6+Lt2svPCqy76txoYpJXlnEK4LxYx+cy/YSxAAAADQcaXFNS/lx8z0ptKtcpLGiruOdtXQklrJqeuopFvN7civu+KvqZhfbt81sovN7spMb1LIaPZ6LJx306RaVWAylVz6Rr/3yLiPYgYAAACAlomL+VOu+EvVhLnRxH6lmpSNqZxf7JmTL+cXvZTPVAs69Z276WblnK+Mt3vCxri4MO2mryj9sq0bXXPTV6qTv6JZxfyyl/JjLn1DjRc1bpnpTff417plctxOYGlWAAAAoMXM7aoCFWpt49LPaiac7suBmZrZR1PHGPvjDS/FWVpci6VpSVJ4dCawYMzdxyQdkPTEVom8XDfMrRAHnm9Fr4S2nl9isSB/StIphdFLJt0/762W41yRdMNNf6Zi/lLHk2jpiLLReCAdr1wfO6DK6hmbEmXJb5hZIZbe3uratCLW7b5e5rqkwD5peP9mPzMl9GiJVW78nslfcOnCputhltlm9ZJbqi7XGsv/QsX8pUFcNteSNgiy0Tl3nRUeast6JS7mz3XdgWWiCTP9hCv0iPe9lJ8gDAD4nuvx7zkAGKjvqKm56pKcNZ7XdrEnhuEArb4/CAEAAAAAdB933524UawykcIgopgBAAAAAN3pQNIGscUUMzCQmDMDAAAAAOoUZKZeV6BMrW1iv3e24ck8w8lhbT3PxmZmP+NqYBBRzAAAAACAOrnF43I7WGubwIK7sdTQqh+B7XrFPXFax1ttWcIV6AEMMwEAAACAOplZIWkbd/2WstF43TvPRJPu/k9TbFngSmBQUcwAAAAAgDrF0k9TbPakuf5ImWgy9Y4z0aSZzkt6MmlTl+iVgYFFMaPflPNLBAEAAABos2L+kqRbKbZ82kxvB9mpueo8GFsLJ4eD7NScmd6W9HSK/d5SKf8aFwKDijkzMCj2EgIAAAC0krnNufl3U2z6pLufMAUnFEYrkm6Y2d37/9Hd/56kv59qKdaNn80lwACjmIFBMUoIAAAA0EpxeeGMZaNjcj1V53vpaIrJPbdnuhmXFs5wBTDIGGYCAAAAAA1yaVrSRx38yI+qnwkMNIoZAAAAANCoYn7ZXcfVmYLGR+46znKsAMUMAAAAAGhOOb/oruMyldr4KZVCRjm/SMABihn9yXSHIGxhJLeXIAAAAKAtyvlF9/jXJF1rw95X3DRJIQP4HMWMfuQqEIQtfLI+RhAAAADQNqXFNS/lj7j0DbWiqGG66bJveyk/xtASYDNWMwEAAACAVirlL7h0QeHkcOC7vuUWj0t2QNK+xL813ZSr4KY/UzF/iWACW6OYAQAAAADtUFpci6XNS6iGR2e23NasRPECSI9iBgZGEGgilpaIBAAAAHZM6cosQQBakN8Rgv5jRsIOAAAAAOhfFDMwOFx7CQIAAAAA9D6KGRgYLo0RBQAAAADofRQz+lAsKxAFAAAAAEC/YgLQfhT7bRlh2MIhQgAAAAB0QDg5LNmzrNKCdqGYAQAAAABojTA6adIpSaOSJJcURn8kadkVf1OlxTWChFZgmEk/KueXCMI29k+OEQQAAACgDTJTL5v0I90vZHzuCUmHTcGPKz02gObRMwOD5V6wlyAAAADUEEYnA7PxWpvE8gsq5pcJFj5vN5PDJv9OwlZPm4I/dOkIAUOzKGb0rxU9WhEdeEGgiVhaIhIAAADbvC+Zjbv7idpbeUESxQxsaDfBd9y1O8Wmh4kWWoFiRv+6TQi2+t7VXoIAAAC6Xnh0ppIg2gHJHiSIsfsNye9K/i5zD6C72O7KBBkp23fpyiwxQzMoZvTro8Rs1d1ZveMhLo0RBQAA0HWy0fPm+g2ZxuR66sG7i1ffYO6/4z34f+28wuiWpIK5FWK790OKG0BrBZmp1xUos+0GsT8el/PHidTOoJjRv2n7KjHYglHMAAAAXSKcHK52zT8m177KK1xde9gn6bCbHzYFLyiMrrvpNeaywA7lH3dTb9ojvTLc/Ii8xtD9SnWRYsYOYTWTPhXHtkoUtnoiaY9GcnsJBAAA2NGX8MzU66bg5+56QaoWMpqzW9Jhc/3UwugyK0ag4/mHx9+TlKagcY1ooSXPUULQt1n7KjHYxifrYwQBAADsiGw0bmFUcPPvSqkmS2zEc6bg5wqjkwQcHVNaXHO37yVs9YEr/ibBQitQzOhXjw8VCMI2jT7QBFEAAAAdl4kmzbWozqw4t9ukHwXZ6A0Cj44pL7zq0jdUWVlxo48lveOKv8bcLmgV5szoV6vztxVGxGFLNkIMAABAR2WmXjbzVzr9se56wcLogJfyR7gI6IhS/oJLFypDnexZmZVUzF8iMGg1ihn97X1JrGjyyJe6jxEFAADQMZlocicKGRscDrJTc3FxYZqLgY6p9MCYJRBoF4aZ9HfafpsYbGmUSUABAEBHZKNxM73d0N+abkpaMbOLZnZR0n+o/rv63wrdTyicOs0FAdAv6JnRx8ys4K5niMQWKpOALhEItFUmNyKtjxCIqseHClqdv00gAAzU+5g0J+nJOv7gpsV2ObZ7P1SxMrfAI6u1hpPDge/6lpsfUR3zb5j8VQ8n/7ztcxZko+eD2L6sQJmN/zr2uCD5ux2dM6E61CGwYGyr/xx7XJDZz7piOdts9LxcvxyYHWjLMYZHZyr7tg0Tz/rd2P1Gx69Lc3Eal2vi4TjtaDvbyRjE/nhczg/k8rAUM/pYLCtYnYuVD4og0ERMMQNtb2f3zrnrBJGovi59ol+SRDEDwOB8D2Sn5tz9qZSb3zW3N+PSwpnEt7fS4losnZF0JshMve7m/1zpCia7TcEfutT6+TPC6KRJxySNy/WEmz9ShTHZCcnOKxvdNOlq7PH32pJwZqNxc83INCbXU5LkvnVUTXZCLimMbkladtNsM0UDC6Oky/cfvZT/1UeOVToi1xNbHeuDY2wkbmF00uTTkh2s7Ft6+MKYTJKdV3j0unmwHJcXzjRx/n8t6Vdqvg/IT6t0Zbaha1ppX/sSr2mD7SzN8dd5zTdf7wfX5eiMyc7X3K/r0sNFiur9Pr0pBiZJOh5kpv6tm/+L5APWT7yY/++berZVjuO7ie9+pq+0s1DIMJN+ZvdWCcI2N5azogk60c58hChsUJ7nmQRgcISTw+7+XMqtP3DFX2okiYzLC2fcdVzSRyn/5LCy0XhLCwdhVDDpR5IOS5WEvPYXpJ5y1wuV5WNbOPSlcixXzfVTSc/dL2SktE/Sc+b6qYXR1ZbGaLNPHiRi2am5B8daZ9yC7NRczW0z0eTn16VSyEiRPBx08+9aGJWbuC6ftPo+2nRNK9cp5YvYxnilXtXnk3Zd7/ozdfv8b8PopIVRuVpA2LfNs+B3VFk1Jiku/6jpd1yL09wfK+3u8UQxo599uFggCNs9qzVGENABFDMAYGBfNYI/kLQ7ZSGjueUqy/nFegoa5nqpNYnm1Okml5rdbfLvWxhdbTqpyUZvVI/lcAvO7LC5ftq2ZW3DyWHLRjfcvdHem7vd/YRloxuVYTRbxMJ0pYnrsq9V16UpmWjSFPy8Bdd0d2VVn6PLW8Wr62WmXq4WC9MUctJcsyeCzNTrDR9PNhpPUyBz6XK7Q0Mxo/+9Twi2vLv2aP8kBQ202zAh4FkEYABVEqY0CdhHbvrtlgy1KOcXzfTHaZP1ZpM6U/Csyb+veuYDqVU8aCJxtjC66q4XWnQsn78uVpa1bXFC73tMwY/r7DWy3fvsU6bgxxuv5YZYaKevS3P3UHSyWpDZ3bqd2kGz4N/31LOkziWd3dKtHpOyZ8V2BYQ083PcUin/WrvDQzGjz5nZKlHYRmwTBAHt+/KJaF8AMKAC3/WtNEmYmf64ld2w42L+lKT/IGkl+R97tslM+r9rcdgON9IToppsH27j5WxxQm8jkp5u4fE9fT9Br8bvcBvO/3JHb6BsNG7S99uyb9dTHT+fhg/Vx+te0rnyPFlJ0Q4PNjqUKs2cO2bWkSIYE4D2e77ucaEyCQ4euclkOWfta7SvhY2ICXg33G8qEA0Ag6Lyq6clbXarWnxo7WeX8v+4Z+PmekHZ6O20BZ7qBKuHO3Bohy2MLnspf6xLs96nLDy67K6DbfqE5xRGJ1XKX+jIO0NlGNSTbfyIjp5PM9e1wWLDZUsxxCiQjsdSfcXUyoo7iccV+72znQgRPTP6/9u0QBC2dYgQoG0P14DJPze/mbCKCYCBeuglJpVmukyctk1kk4XRyQbmnLilz3un/H8NJcA93Oaa+xrXqY6cRqW3wOGUB3VT0jtmdlGmvFL1SOjw+eyEyvCOW4lpYgOrGpnrN1Jsdq1Ty+LSM6PflfNLCiPisO0Dcyqn4sI8gUCrsWLOZnHMkDcAAyI8OpPquSi93WdnvmKm5dj9hkpXZivzONizgQVj7n5E6VehqMznkZAMmelMyg6Qd83snVh+4ZEeH9no+WpylmrVGZNec6kNv+b7dbPgRuxxQZIC2T6XflnSuOpZvaPm/q0Qu9+QpMCDv+vmT0s6qPRzUox2ojdDIB1PcVk/cunbKlaOxR8qhpj76ymKO6PKRuMPtwk3/Z7cww3X/MXq0KAaRQE/XaOhlnbiZjSzq4nFPtdTykbPq5i/VMeuEwsgbnqrU+dJMWMwrKjx2Yz7WuA+EUsUM9AOI4Rg01cbxQwAg/FuYXbAk7Ix0812L1nYQR+57KxKC+c3nXalGDEb3z/lyjwFqYoGge/6Vixtu0xtkJl63d3TdMH/wBV/zYvbFEaK+UsuXaokwPojJc9lsS/ITL3eyBK62x6f67TKVxY3xi5++FzN/7kaGnbh193sjIpXlrfcfzg5HNiuV9L2cDHpWHuKORuO2JVJ3Eb69rZFlWJ+2aVxC48uJxU0gtiiR4ZZPJzYh9F04kGXrsx2vlqhm3Ldrf6vxx/+z7H8gkmJ19Vcv+FSumJGGL2kpCWEK8+2S50KA8NMBoBJS0Rh24dhjiig9WWM3F6xkslmjw8VCAKAAXnzSv6l23Wjb96lXMdVWjifuF0pf8zMLqbaZ8JKCx54mrkr0i95W8wvu+KvSfog+dj8SItCVzm+cn6x1kZxeeFMPcvubnDNS1fGaxbNSotrcXFh2t3Szm8w1oEmdSC5eJDcO8TNzqS4lgd66Fa7a2YX3XTcS3nzYv6Al/Jj1X9+das2Lfn1FPtNPQmoKbntW2wdHT5HMWMAxMybUcswS7Si5T67N0EQHrI6f5sgABgE7nFigmRm5X44VzO7mJSMb3onLS5MV+c5SNrz9jHMRuNpJiCse8nb0uKau06n2HK00VUgNiamqQstUr3L7krSLS/l0xddyguvpkx89zW7pG9LpIl/pYjznxK2+lKP3GofuOIvxcWF6Xp6PbhsLt01jZLnqakMG0sauvNxC3stpUIxYyDsWiIGNW6COJgmCmjtg9XHiMIm7xMCAIPDdg/KmTayYkHKX2731Uhejqf4+2sNDeOpFGauJb87WlMT0pnZO/VOkFhd+eZWqv17qiR2c+KboidDde/PtvXuseQfYc399TRFFS/l/46X8lbjn1/t/seJSnUVvjaq9GC5lfwRyb0zAgu+k+ITr3Y6PBQzBkF5flWmOwRimwcdQ03Q6jbF5J8PR+Q2MQCAjUWAuND7Z+HXG0mwUv9yu02y6p78g4HLG06qLEWPZrfmfrSI5X/R4J8utzTGGxXzy+l6zbT75lCKXkt20Cz4yyAbvdEVPUXayGL742ZWBkm5atLhpDh6uok/ZzsdHyYAHZzsakmyZwjElipDTT5cLBAKtAjL/m76IrWCEwYAeCDwIBP3/LM9uNHws70yeWHCUBF7VtoqObJ08xykXFXm0aTNfyXFZpmmgtfgBInmdsMtMeorDR+XqyDVvi6B2YF2tt3Y7v3QFHw3xbGGLr1gCl5QeLSyWov0dh9NrCu1YNhG7PH3TMELyc+kGpPuZqPnk+9Xv67ilY7HnmLGgHBpySSKGdvdwHEwHUszRAJNy0QTBOGhL1Ixbw8AbH7xaDIZ7o63y7tNvJjebeKDE5cqNdn5Np9849evid4PscVlk7XtpMzsricvxdPeYVSlxTWF0TtKufJN9ZgOuuugSS8ojG5JKri0LMV/0kyvhi5wo0XxvCbpcM1bsjKx7ZbFjOryxQm3tF3diQAxzGRgvjR9iSDUugEZaoJWvQkwX8ajMbm3ShAADJDEBMTde35ejdj9xgBf430N/2VzhZwBeS+P/6VSrCxT49ocNul3TcGqhUeXg8zU6/36LEkVT9NbKTarNbFt0hCTWyrlX9uRFJfbZUB8uFhg3oyahpWdyhEGNJ23y2hHWz1/AGBQvgfM0iSrB4gUsI3S4por/lpr5vCwg27+XQujcq8VNVI+S5IV85fSxDKQnXzkX1ZWOnmi9nHq8k7FiGLGYCVZ80ShRnw8niYKaAHmy9iMlUwADJSUPRZGiRRQQ2lxzYv5A2Z2sUV73Fcpahxd7vdJQ7fMc1KsIuTuRx7NHz2pV8bd2OPv7dR5UcwYpC/XWEtEoeZt/owyuRHigIbRu+fRu8pslSgAGKwHX8r3rcovnuhNtwhBh/KX4sK0m75SLWq0IO520BT8eNAKGrHd+6GUOMRpn7LR858/oyaHJTuY8DfXd3JeEooZA2XXEjFIuCGC9WmigCYeqDmi8NCXZ18sPwgAdaisppCYdJl0jGD1rNuEoLP3VFxcmPZSPuOm45LeaXIIytOm4A8GKoaVgsP1xOeS6+sP3mst+E7S9inn42gbVjMZJOX5VYXRiujauP0N6ZqWdI5IoLH24xNE4ZFvuQJBADCAlpW8GsOowuikSvkLrfzgIBu9kWbFiVj+F40uEbrDPlbCGH4v5Y0m2KeK+UsuVdptNhoPYou8Mvn6eFK7eMhzykbP9+g90OArmV4zr72qSTWOle2TJ/5c2en4UcwYMCYtOcWMWoaVmZpWeWGOUKAu+yfHFGuYQDyknF8iCAAGjZtmzZOXljTTGZdaV8wIo5PuekHy5G1NF3o0vDeU9C6bjcarPWTQz4r55bhSOLx/3Z+v9iw4KCmxoGeurz8ojAxIvFL8sP1EkJl6PQ48L9dTteNnV32HT4lhJgMmDmKS9MQXC58mCqj7YRoHtJtHrRACAIOaZKV6BrqesjC62pLPDCeHTfpOyq1XejXZN0vu8RdIx2mEA3nfXfJS/ogr/pKka8mNafBWFTK3xOeNWzxurpmEzT6Oywtndvp86JkxaD5cLCgb3ZFrD8HY1iFlogl+UUY9nPkytnhH0JITBgCD+73whkk/SrHp4SAbvREX86eae+ban0p6Om1C06vP51j+FyadqBl71zFJ9ccznDodWEKvj9gfj8t5iiVtEGSmXnervXqGuV1NTKJLi2suHbEwKqhWL4SEngf9KC4vnLEwOqWaQ3LsoBLm/TGzy93wDKFnxkAmGCzRmhgjY94M1CETTUgMMXn0fY/5MgAMsFL+guTX02zqrhcsjK42tMJCODlcSdoSVx24/yJ4sxt+UW1YZYx+0gSr+yyMLte132w0bvJX3P1EzX8C/UMad1uN1vonqdjxUMLdmfeQjSuA9IY0vcH21XzHk3fFMDWKGYOYYEjzRCHRoWqCCiQ/SAObJgpb2bVEDAAMMjc7I+mjlJsfNgU/DzJTr6cqaoSTw0E2esMU/Fx1zIfmrtd7Pa5mliYZe64yGWoK4eSwuf5I0pOJnx3bZVp2m3KUwP8qxWaViXNTtXXPdCahtl/vreeSZpvcw/VuGabGMJNBVFyYZ6hJmi9KnXNpgkigppHcXv9sPUcgHrGm8vwqYQAw2O9c+WUPp86a/Psp/2K3m3/XFLyg8OiKmRVi9xubEie3f+Cm/ZIOuidPcviQd1q9esqOJL1+76wpeE4JkzxWe7yMu+nUtslXGJ0006tyhSk++m5s935Iw27b/XJJYfQDJfQKMOn7Hkaq2ZYzUy9LnrRyx60UR1VWQrHQ3U8Emejx2Pz/2CKhKHXdiinF/LLCo9dT9+Z6+Hxlc91yKhQzBjVRl827/ASRqIm5M5Dsk3s5GYXBR7+7jfkyAECSSgvnFUYHlbxU66aihmQH3XXQtHmVUW900VHTTS/mj/VHTBfXLBtdrKzckmjUXD9VNropV8HM7lYSMv87cv19SfuU8gvL3N708uIajbqt7w9X3RNzlCdN+pHC6JRLl2X6v1XMX1I2el6uXzZpPEUhQ5IKKa55wS15X2563mTPb5H5r3Tjiikuu2qVVV/qdaubCqIMMxlQDDVJ+0Bl7gwktRGfIQpbPGNiLREFAKgmDqX8MaVZXaF9PnCP/0lffc8U86fSzklSzd6ekvTc/bkv5PqqEnoAPLSD6z0910ivXFe/d1amUsrNR036XXO9bWHk5nrbpN+VdDhdk1DikKE48HxfBrqUf03peqZsfu/17umVIVHMGFzFhXmZ7hCIRMydge1V2sYogdjKriViAAAbEqdS/oikd3bgoz9wxV9Tqf96FLj8NyV90JkY+m/SijuRZC+uuevlDnzSSqoeBsX8cl1Fsx6Scu6ZjbpumBXFjAHGqiZpb/RmJ8lB/7YNemVsg/kyAGCr5LuUP+ZuZ5V+UtBmXevXQsaDxFfx19TegsYHfR3Drryu+QvV+6Rd7rriZ1Lft/VN5NszGliR5Hq33QcUMwZYbPdI0tMZVWZqmjBgk0xuRLJnCMSjjGFsALC98sKrbpqUtNLeZM2+7aX8kb5Pwh8UNNry6/kKhYwdvE9k325DEeGuu56v65oW88vu9m/6LsaVSXFTD39z02vddgoUMwbZh4sFSTyc0yRngc9qJLeXSODBwzO4d44obC02WyIKAJCQHJXyYy59Q60tanxsZhdd8ZdUWjg/MPEsLa556cq4uf2+pLutSHjN9KaX8mMUMnbyui6cb23hz6+74i+pnF+s+0/LC6+69I065vPoCWnmDala6ZblWDe9j3OXDHiSbpojCqnu9D3Bp+skr6jI5EZSzLQ9uL64a4kgAECaZC1/wUv5MTcdV2U+jVuN7civm9vveym/Oy4uTA9qAh6XF8644i+Z6c0GY/mgGFSZYBQ7blPhr+HeNysufcNLV8abujdK+Qvu8a+Z6U2ZbvbLMyhVzuh1z6/RESzNOuDieGjObP0skUjxmiC9qEw0z1KtCIJ755x1R7e7U97T6vxt4gAAdSVslx4s35iNxoPYIgXKuMcHJNv9SGJhVpD8bux+Q6Urs9X3lNa9H7ovB2ZJ70U/a3T/5nZVQe1lMRvef2lxLZZOSTqlbPR8IH3F3ceqcTwg6YmNSe79eMYeF1S6MttsHM11SYF9sn1wVW74M8x+ZrKLtS9e4/tPc91j9+Udubal/AWXLiicHA5817fc/ED1ema0eVWaW5LKksrmVogDz7e0R8Hn7UsKJ4cle1aSAgvG6roeKa5lUqxb8k6bmXrdk58et7p1JZ/ElaqDbHTOXSS7j7Q/vRIX8+f64lzCo/OM/U9txUv5McIwwDK5EbP1vyEQ272g+On7L9a9gu+5/v+eAwAAW3zXZ6Mb1WWLa7wP2MW4uDDdle9wXEK4BXNEIbXRIBvxcj/QCd4vZolCrQfKF+YJAgAAQJcLo5NJhQypoVVPOoZiBqTiwryYCDR9ruY6q/2TY0RiAGWiCXox1bTCkqwAAADdz6RjKTa71o0Tf95HMQOVxsxEoPXFKw7mWN1kIO+TWaKwPZfzHAEAAOh22Whc0uHkd7vUq53sCIoZkCTFXxwiSavPKKubDJjw6IykUQJR6xuPISYAAADdLpCdTNzIdDPtaic7dx6AJK3O3zZLmB0Zm/M26UVlp3JEYgBkciNmdo5A1MQQEwAAgF7IY9wTh5hYbJe7/TwoZuCB2O7NEoX6mHyO+TMG4Drb+pxce4hEjS9FhpgAAAB0vSAz9bo2L1G8lbux3fth158LlxMPfLhYkPQ+gagrg9vD/Bl9rjK85BCBSLoXGGICAADQ9a9s5tMpNrum0mLXLxBBMQObG7frHFGo26h9+os5wtCHGF6S9snxHkNMAAAAulwYnZS0L/HNrkcmvaeYgc3K+SWxTGsD7JkgjGaJQ59dVVufZ3hJMvdgnigAAAB0+butUvXKWOnm5Vg3opiBLRITfoluKG7Si8pMTROJ/lAtTrF6SfK34h09vmueQAAAAHSxbDQu2cEUOc3lXjklihl4VHlhTvTOaCyvM3+LFU76QGZq2qUXCUSKNi+b1+r8bSIBAADQxe9srpkUm91SKf9ar5wTxQxsid4ZzSR3rHDS0/ZPjlngswQiHVZBAgAA6HLh5LCkw4l5jNnVXjotihnYGr0zGufaYx4sUdDoQSO5vebBEvNkpLZSXQUJAAAA3Zr0+65vSdqdtF3s98721HlxabFtTk7vjCaCR0Gj54zk9tqn6xQy6ntGzBIFAACALn9n66PlWDca4tJiW+WFOYXROUnDBKORp4b2mIIl3z85wa/XXe5+IYMJP9Mz3VFpYY5AAAAAdLFsNG5KHj4Syy/02qlRzEDtfNxsxtzfJRKNBrBa0MhOTau4ME9AujQv/3R9XhQy6ouZNOuEAQAAoLsV88uxtNyPp8YwEyQ0/oV5Se8TiCa49pj7uyzb2oVGcnuD7NScpEMEoz7xF4dmiQIAAAB2CsUMJOfirnNEoXlm/laQjYhlt6gOLXH3EwSj3rZsF1mOFQAAADuJYgaSlfNLkr9HIJrnrrMWHp3XSG4v0dhBzJHRlDjedY4oAAAAYCdRzEDKJPwLM0ShVewZ+3SdlU52yv7JMQoZTbRes4sqz68SCQAAAOwkihlIpzy/aqZXCETLjJoHS8yj0WGZaMI8WBKFjIbFdm+WKAAAAGCnUcxA+iSmMuHfGpFoEdceM3+LYScdEh6dMdNP5NpDMBr2PssMAwAAoBtQzEB6q/O33WyGQLSaPWOfra8qO5UjFm0wkttr4dF5k50nGM1hMmAAAAB0C4oZqA9LtbYpS6ws30ovjRbLRBP22fqqZM8QjOZU5srILxEJAAAAdIMhQoC6824fmjZb/xsi0ZaU8Rn7bH3Vw6PnVLoySzwaNJLbG3y6fs6lF+WEoxVYwQQA0ApBZup1Bcps/4Xjj8fl/HEiBSAJxQzUrzy/atnoFXedJRht4NpjsvMKoxk3m6n2hkFa2amcfbY+52JujFYxs4teYgUTAEALXnPMj8hrTMRtkiSKGQASMcwEDYmL+XOSVohEWw1Xhp5ES8pEE4QjQSaasDBaMvd3meSzhUx36JUBAACAbkPPDDTMg3ja4uCviETbHTLTTxRG77vZLD01HpLJjQTBvXPufoJgtJ5Js17u714ZcWyrZs5cQI885PV3g2x0jkCg7fdg5QeS/nt+htFfS/qV2reZn2ZYKdCEbDQeyE7WfsioHJcXzhCs/kMxA437cLHAcJOOOmTuhxRGa+52To/vmtfq/O2Bjcb+ybHAd824r59w5sVol7Xqksz9rbww59Icl/shYbQk1z8jEOiAc316Xp9waYE2c/+HLtX+Qcu0IoliRh9imAmawnCTHTFs5m/ZZ+urQXZqTvsnxwbq7LNTOQujJYuDv6I3RrvfD+zcQBfMAAAA0LXomYHmEx6Gm+xQ4LXH5SfMgxMKozWT5uMgntOHi4W+O9dMbiSw9RmXcnIf5uJ3xPsqL8wRBgAAAHQjihloHsNNusGwSy9aHLz4oLDhmlc5v9SzZ5TJjch+kTPZtLQ+ykiSzvIgniEKAAAA6FYUM9AScTF/zsJoQtIhorHjKoUN04sKI0n+npkV4lhLXV3cGMnt1Wf3JgL3CZdy0vrw/fXZ0FlmesX7sYcPAAAA+gbFDLSM+9C0BesFlsXsutT0GXc9Y6azleKG3jez1djjgtwKenyosCPzImSiCZmPBRaMufuYPl0flSR6YOy4wZj0EwAAAD2NYgZapzy/6tmpaZO/SzC62iF3P2SyEzJJn67rfpFDksy0JEmxrKDYb1cfFauqZ3nO/ZNjuhfsrfwPGwkCH5FsxN1HJI1Iqs57YXKWIukq7ppm0k8AAAB0O4oZaK3iwryF0Q9cepFg9JxD1WT2kCSZfMMojwcFj3RibfhbV6VeQdGi25n0A+/leVYAALWFk8NS8M8CswPunpGUqf6XG2Z2N3a/IcV/otLiWvcevz1b7dW5W9KB6n8pm1k5dr8h05KK+eWOHVM2GpdrohrTDcfkd82CG5JX4lq6MtsVMcxG40FskYLKtY/lFxqOV3h0ZrtroVjlOPB8R6/FjrbNSiwq79LxAcl2P3JvdbptpngOVN/by3F5oSeXrqWYgZaLS/mZ6vwZo0QD6Blr8WND5wgDAPSfIDP1upsfuf9utkWvyFF3r/4OEfyuwqPXXTanUv7Cxo0sjBJ/mfBSvvUTXoXRSZOOSTqcePwuKRvdtNguty1By0bj5pqRNC7Xvq2PyeTuByv/l0lhdF7y6+bBcrPHleI6/Ecv5X91izYwLdc+N9/wG5MXJC23+lrIJHN9d7trYWH015J+pWZbkp9+pAgUHp0x2fk6QzZaK2Zbfk7q+yoel+zg5ljY9m0zjG6Z2dWmikiSgkz0tpueryt+mWjSTL/zyLUzrUg6Y+HR5fvnUrP9uf1+0204G/1vcn219sNERS/n91PMQMf5Y0MT9tn6KvNnAD1yzzK8BAD6TxidNNMZd3+qzlTjoEkHFUan3HRqx35NDieHTfankg7W9Xeup9z8uxZG0y699HBRpskixkvySjLYQAp30M0PWhhNm9tcG38N/2RTAhvo39bfBrY6d3+94WuRjY55rN9ROb/4yDH2mPuFIZfva2Cy+n3ufsKkEwqja674mw31hArsE9UxVNvC6LKk52peKrP/xTz5+lYKOM3d14mFDEkm+5OkMwx4yqMtVudve6wcgQC6n0k/EMNLAKCvBNmpOZN+JFczSeyouRYVRic7fgKZaNIU/DzNL8W1EkeTflRN5Jo8nqmXzbUoNVrIeCihrRRbCpUu/20STp0209tNtoHKflyLTV0L11NmuqLM1Ms9e1Nlo3ELo4Kbf1eq9Mhp0mFT8HOFU6fb1wYmhy0b3UgqZEiSivlLMt1M8eZ4UNmo4YJG4Lu+lWKzu7Hd+2HivnjUo23K+SWXnyYQQFdbiUv5GcIAAP3Dwuiqu59o0e6eNOlHHU1CK93h35a0u0V7fM7C6Goz8TTzVyQ92eIzHTUFP1cmmmxLDOXfb/qYM1Mvt2Q/92Np/kpLikudlokmq8WsVg+j323y77crJmbBv6+nmGWxpTqOQDre6DF54MdSbHYtTY8Vihlor9KVWTO7SCCAbnzb1R0P4mkCAQB99GivJEWHW75f83/VsXMwXWpD4eBwIwWN6t8cbuPp7jbT260taPieagybE06drhZxWu05Sb/cMzdVZbjWlTa0yU0xaabgtvVrnp2pt1dOdejTx4ktzHWsoYOqTJibeExumk2zO4oZaLu4uDAtaYVIAN3FY5vRh4sFIgEAfaLSXf25Nu39yQ6eye427fdwkI3eSJ0MtqkwtFVszfR2M133HzrykaZjmI3GTW0pZNz3n/XEPZWJJk36foc+7XCQnZpr4f4aHQpzNdW+s9HzdbfMysS5SVbSztNDMQOdSZoeG5qQtEYkgO5gZhdVXpgjEgDQJ8LJYZO/SiAS3kldv5WqaNDewtBWnjSpa76XzfWGOlvA6s57ynS+k3Fw9xM7MkfNxmNI2SvCXL/RwO4T7z2XUg+5YTUTdMbq/G3fP5kzD5ZY4QTYcSvxF3fNEAYA6B+B7XrF3ev5Nf6upOsuLUt+98F+3P5bN02oNRMcttItScsuf/CLbSDb59KXVVlhI+25P2mul1w6UjuJ9e8o7WIRppJc/7vLl2X2M3n8nyR7NjA74NKR1F39XU8F2am5aq/mHUzip05LXufcENVlZy0ub2hL/8BN/41aOs+Ev+ubQm+HpKRFB3zVpR/U2meNe+rphmJQujKr8OiMZLutksCn7uFj0msuXdix61/MLyuMVlJctyP1tavopRTPlVsq5V+jmIHu8+FiwTNRzkw/IRjADjHd8XgoxzKsANBHwslhd6+nF8E7rvhfbjXBXnw/kassP/mC2jfkI62PzPTHcTF/artjVTg5XE080056eljZaHy7ruzVfYWpvlbNLsZ+7+wWsZyNNxQHqkM2En/hd/fnFE6ebWi5zmQfS7rx0AmUHj0nf0HpV/1cqSzfe2XZt70+lSWCm15VRVI1LrMP4hUelckSihl2R6X8bF2fk43G3f2fptz6A3edVvnK4qYYlK7MSqqE8sHStqlWhNkXZKZeb9vSvaabct3d8L//0yObuF11SyxoPaEweilt8cFS9Mow02Wv41QoZqCzyvklz0x93czfIhhA53msnMrzq0QCAPpHYMF33NMVHdztrMoLicNR4vLCGYWTPzQFP5b09A6d2kfuOu6l/GJSghtL08pM/T9pJ6w010ylV8qWBYVjaWPppeRYqrRw3rPR/2muP0oRy92BBd+JpVOtiqGZ/Xksv5BqHoJs9HwdRYdrXson/zpfyl9w6UIHJlNt3T0lO+nyJ1PFQPE3VU4oPhXzyy6Np42Bm09LamUxY8Wly2kLD3F54YyF0bQSelKY/IhLyfusLEGceN6xx9+r7zoBHS9oLMy529cJBNDhQobb11XOLxEJAOiz53vK7t5mejNNIWNjkcAVf03a8CtuhwsZKicUMja/Y77qbmdTbr31r8SVrvBPtDyWxfyyu05L+ij5+7rBlSIe9YG7jsfFhem0EyqmnwfBr6cqZGz8i1L+iEw3e+OdyY+kiq/ib9bTi6YSM7+eYtN91bbYNDO96aX8WD3DNyp/ZykmArWD1UJF7aKD7/pWio+8Vm+PJIoZ2LGCBku2Ap3DhJ8A0KfCyeGUv6SvbDVUI11Bw17ege+tP6+rkLGhoJE6WdxiNYY0XeEl3WooluX8opn9eaPHVn9CrtN1x9A0lmKruy7/zYaOSZru+nuqEvt9Kc7le40MB6rG7m7ypUh1LRLf/xpqq5Jieap5O9IUKtySi0NuqrvnPsUM7Ji4uDBNQQPoBH9vxycTAwC0K+1/NmUS+UbDH1FaON/hX9Q/buZ7y81Sdc8PYvvyFv86MYE0t7mG33/93tk0iWwgfaXJGF5rqBiUojBmZu80PKdHpYfItW6+owLZryffdrqpUr6xSTorsbuW4jPGmjyVu9X2psavVXJh0IOEYVmV1YNGE+NZzF+q/1oBO4iCBtB2K/7YF6YJAwD0p8CCNAnPrYYTrwe5hq528LSa+6xifjlV8SVQZtP/rnSXT/xFPrZ7P2z42EqLa5KvJCaIrgPNhMDl9ccwPDqT6v095S/22x5bA7/Ad5J7fCD5frDlJq/PcoqNmp0w9XqzE8m6Ugw1cT1Va7njQJa41KzFdrmR46OYgR1HQQNomzV/bGiClUsAYOAVmn5fk37asWQyTaLXgmTT3cce+qtnU+z4ZrMJollwI8VmmaYCUF1Jo+Uqv6A3d30a+AW+syxxMt3Y40JHrk8Tw43MrNx0KCrzbNxK/CzXTI37LGmIyceNrtxCMQNdgYIG0PKXjTsexCzBCgB9LtWvyK1IajqZgJr9rOl3S/cb7Ql485OhNp0IJ7vVyB+l6uXjLZsMdqWLb6sUvWL83Y5cp5RLBLeznZkpTa+J7SbUPank3k4N98SimIGuQUEDaNVLoO64xRP6cLFAMACg7x/6KX5FblNi3y7N/vJfyQLbtQJLp2LZzDCTcttam9kgvFskrmbTbO+cdl+nluZo6ZZL3VctXDz0Spq8Mo+bZhs9NooZ6K6bhYIG0ORbBoUMAAAAtEjKCUsfKVxU5qA5WPuv/HozxUuKGeg6FDSABlHIAIABlNwDITA70FOnVEmCmv1S3N2mo+tULLuyN82j84wMqJa00SbnRenkdVeqoSYHN8alumTr7tr7bXxlIIliBroUBQ2g3nc2ChkAMJCP/xSTSbp780lTjdUK2nBWzza7h3QFnLYNRWnBse3A+3eaORZMrSoSdXOBbaUTbVQpVs7pGqX8hRQrBO2uFjAqd5clTvzZ9CpLFDPQtShoAGnf+ShkAABqGms6AY8t6lyyHzR9vC5PLL48UghKN/HoaNPHlm7Z1e6cTyFhGc5UKn//RC/fUE0XpFIug9u2VWkaed1MsXyqWzy+4RqP1r7/dLnp68CzHd0sLi5Mu9vXiQSw3TcBhQwAGOh3pXQrFmw5OV9dOWzyr6yty5e9yc/KRuNyPZUcvIcKBpWx+x8nJ6LRS40nsWnmEZDMdmCYScrEOZA11Zaa/fu2v1qlmOS02eE2JktTELrVTXGJ7d4PpaTVbOygstF4imt8N+XEogltCeh25YU5ChrAlt+EFDIAYOClWyLSpFONJ+DRSbWgR0Id9gWZqdcb/np0T/W3ceD5LeJZSBHLY40eW2DBd6TkoRqx9NMdere4mdji3J9reM6IcHLY3Z/r5jsq3eo/laS90RhIOpxiy+56v6tMBHo9RYHheIqC5PVWrAhDMQO9obww50H8ZZnuEAxAFDIAAJ8nGCkSUEmjQTZ6o5HEy6TvdPq03PyFhhLmcOq0ZAdTbHlrq1UUUi49OtpQT5dsNO6u30p5bJd2pD15qgR6t8n+tLHXF/tTSbu7+57Kv6YUPXTSFs0ajYFLy90WGje9lbiN67eVMB9Imv2kQTEDvePDxYJbPKFUk/IAfW3Fvzg0QiEDAFBJjnQ1VSLiekGZqZfTFwYmh6uJ19M7cFq7TcGP6ypoZKJJk7+ScustE8W0Xd9N+r4y0WRdsXT9kaQnE/dtdnWn2pKb/ixlqztoYVTXcVa2T1Vo6gbLaWIQZKfm6tprZurl1MW2SlGlu1SKbEm52ONJ77GtKtZRzEBv+XCx4I8NTUh6n2BgQK34Y0MTWp2/TSgAABsS8FQrc5j5K6kSsEw0aQp+vMPJ59NmwV+m6QURZKM3zHQlTbGgmrTPbvkfKl3fr6XYxZNmejtVD40HsUxVFLob+72zO5qspuvpI0mHLTy6nFhwCieHLTy6rHRDK9phtN4hIWl7Drj7CQujq+liEF02S1ds28mCVuKxeXPH1uzfbzTE4x89Z3X+tksTQRjNuvQiAcGgMLOLcXFhmkgAAB5OwC079Y67n6gjATtipsux9NMHv5KGk8NS8M9MGt/BxPOhg1Vo0o8URqfM7WoceP7B8JBs9HwgfcVdx9zrWuby2lZDTDYksq+Zpzr/J+8fm0uXpfhPHswD0GAszewdLzY/l0BT7xuxXXbz76bc+qDJfq4wuuby5U2TiIZHZ0x2RJUJT9sztKR0ZVZhdD45gdZfWCb6X2Pz/2uLK/7uI/M3FPOXFB79FymLeYfNgr+0bPTnsfT2praVjZ4PZL9enSckXQxMpR0taCWIywtnLIxOqbEVaT6OywtnKGZg4MWl/IwyUwULfFauPUQE/V3I0CtxceEckQCAPnq2Sy8qjKbrzu9Nv/dwN+3Y7501BekTJmmfu14w6QWF0ds9EK5RNx8113cVRvcLHfL69/ORm2p33y/mly07dTFtcUjSqEmjUvC7D46tsQZxsxt+tKgmq0eUftLX3ZKeM9lzaQoLbbiRbqZYvWa3m/5Hkz16P1X+v9lH7zM7Y65Fpent4wpdD+6nTTv3Olupu15uxeSYbXZVUt0TuZrZZW/hQTDMBL2tvDBXnUdjjWCgT99077jZs3ExTyEDAPrvIT9STRjr+8c9fGRXpcU1l71MTJOSKf1xrV4ZDxL64sJ0HcMtWuEjj/U73RInd/3Pkj7qiYvqbVr1o5hfdrd/0+GzeUel/IWuD/l2w7SS7it5S8+NYgZ6X2UejTHJ3yMY6DMrbvGEigvzhAIAkKi0cF7p5nsYVNfiYj71ErXu8T+R9EFn8nF9W+X8YtdEqpxf3IFEvqOJdbo4LLwq6Z1OtU8v5Y/1xJ1UKQjWuyjDtTSFxHpQzEB/WJ2/7aUrOZefZvlW9Ad/zx8bYulVAEB93x6l/BG1oaDhbh0bw9+WzzLddMXfrOtvSotrrvhram9B4yOXvtGVv8aXF15t03W/plauTljML5vZxTbeU8fU/oLGteq92zvPGulyO7dPg2IG+kvpyizLt6LnX0Tlp710JceKJQCABpOvI61MvtztbPUX6l5Noq95MX+goXkIHhQ0/HrLz9N0013Hu3pYQTuuRRuS9upcI9faeE8dq8ah1UNvPjLTm71WyKjcG/nXJN1KufWtdrRzihnoPx8uFryUHzPpBwQDPWbNg/jLm2YCBwCg0eRL9m2lXLK1RrJ9tKOFjE1JtI7WkSxt5W5LEsXS4pqXroxX43mrBWd318wuejF/oKuGliRdi+bmELnrsm+3M2n3Uv5IO3toqLzwqpsmW1fY8utumqxn6FO3SbuErLnNtePzKWagb8Wl/Iy7viomB0VvvHa+548NjTGsBADQMqWF8674S9UEr54k/Ja5/f6OJ9vl/KKX8hlz+/06j/9jM7voir/U0kSxtHD+wfE0ktibbprpzcpx9dhS6+X8ohfzBxo49wfXojqnS3vf/4sL0276ipldbMsErsX8speujLvpuCo9QT6ucw8fS7rmpuNeujLe6jkkOp5vud9Isdnd2O79sB2fz9Ks6G/l/JKP5MaCT9fPufQiAUHXMd3x2GZUzs8RDADoH276vS1XHWnN3t9Nn4AvrsVSJXHORs8Hsl93jw9I9tASrn7XLLgRy/9CxfylrZZPdPnpnYhLXF44I+lM7eOXzKwQe1xQ6cqst/Habjie8SC2yM0PSDqwzeY3zO1GHHhexfxyK44r8TqYlTpw7s8H0lfcdUBSZutr4TdUyr/mDbeDOtr5QwWHWPq8SBAenWn551TukUv39x9YMLZdu3xwb1XbZkuug/yC5IW2xK+u11g/oi2Wu33I9XYtNZv4yUE2Oueujk340zP5h+kVlkrsMZlowiqzHY8SDHSJ992HplWeXyUU6LrvuTBaknSISKDtSX8pb0QBAHpMNho3108Tn/Gmr7SrBwrDTDA4yvklL+XHzPQKK55gZ7NE3alM8pmfoJABAACAnnuddc2k2GylnUNpKGZg4MTF/DmPh8Ykf49ooPP8PY+HxpjkEwAAAD0pnByWdDhpM/N0E4Q2imIGBlN5ftVLV3LVCUJZxhWdsOZmz3rpSo7eGAAAAOhVge/6lqTdCZvdqs6x0r7j4FJgoFWHnrjb1xl6gnYx0yv+2NCYigvzRAMAAAC9zM2nk99/29srQ6KYAVSUF+b8i0MjzKeBFj/q33Mf+qW4mD+n1fnbxAMAAAA9LYxekrQvabPKiivtRTEDuG91/nZczJ+jqIEWeN9dX2VICQAAAPqJScdSbHatnRN/3kcxA3gYRQ00rjovRn5C5fwS4QAAAEDfyEbjkkaTNnPpcicOh2IGsB2KGkhvzd2+7qX8CPNiAAAAoB+Z66XkjXRTpfyFThwPxQwgycNFDWmNoKDq8yJGeWGOcAAAAKAvVZZjPZi0mcV2uVOHNMRVAVJanb8dS+cknVNmatoqs/geIjAD6X13nWMoCQAAAAZBYMF33BOXY70b270fduyYuCxAA8oLc17KT3gQf9nMLjIEZTCY2UUP4i8zJwYAAAAGiXvKiT9Lix3rxU7PDKAZHy4WYmlaI7m9+uRezsxnlGJSHPSUNTPNxV8cmnWWVwUAAMCACTJTr7s8cTlWN8128rgoZgCtUEly51ya0/7JsSAOpl3KSRomOL3K33ML5lRcmHeCAQAAgEF9KzY/kmKzlU4sx7oRxQyg1Sq9NWYkzSg7lQuknMtzcu0hOF1vxeVzeuwLc6IXBgAAAAZdNnpe3j3LsW6UWMyIYy0FzKyxZVyIAhIVF+Zjab7yIKgWNtwnRI+NblIpYPgX5lWeXyUcQPWlxG0uCJzvOrS/rRECAOjmh/STZnaxdnLsj3s5/1qnD824OsAO2D85FniQc1dOzLGxE953+TwFDAAAAKA3UcwAdlp18tAg0AS9NtpmzcyW4lhLenzXPENIAAAAgN5GMQPoNpnciHRvolrcGBM9Nxp5st2R+5JLSwp8SR8uFggKAAAA0E+v/AC620hurz5ZH6sUNzQh0xiTiT5ixcwKlblsdi0xdAQAAADobxQzgF60ocAh2Yi7j0g6NABPrDtyFUwqxPJVuRVUzi/RIAAAAIDBQjED6CfVIofMxwKzvZVhKra3B3tzvC9JZlqK3W/LraDHhwrMdQEAAABAopgBDJZMbkRaH6nc/ZWChyQ9KHp8bkStnYh0TdLqhgdPQabbkhTLCor9tnbFt5nbAgAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAYOf9/9KJZACC1w2YAAAAAElFTkSuQmCC')),
    width=184,
)
st.sidebar.markdown(
    '<div class="icl-sidebar-product">CHEMICAL MANAGER</div>',
    unsafe_allow_html=True,
)

st.sidebar.divider()

# ICL UX v3.2 — hierarchical navigation.
_pending = st.session_state.pop("_pending_navigation", None)

_pending_map = {
    "Chemical Search": ("Inventory", "Chemicals"),
    "Inventory Management": ("Inventory", "Chemicals"),
    "Inventory Scan": ("Inventory", "Scan"),
    "Storage Management": ("Storage", None),
    "Excel Import": ("Excel", "Import"),
    "Excel Export": ("Excel", "Export"),
    "Add Chemical": ("Add Chemical", None),
    "SDS": ("SDS", None),
    "User Management": ("User Management", None),
    "Activity Log": ("Activity Log", None),
    "Settings": ("Settings", None),
}

_root_items = [
    "Inventory",
    "Storage",
    "Excel",
    "Add Chemical",
    "SDS",
]
if CURRENT_USER["role"] == "admin":
    _root_items.extend(["User Management", "Activity Log"])
_root_items.append("Settings")

if _pending in _pending_map:
    _pending_root, _pending_sub = _pending_map[_pending]
    st.session_state["icl_nav_root"] = _pending_root
    if _pending_root == "Inventory" and _pending_sub:
        st.session_state["icl_inventory_sub"] = _pending_sub
    elif _pending_root == "Excel" and _pending_sub:
        st.session_state["icl_excel_sub"] = _pending_sub

if st.session_state.get("icl_nav_root") not in _root_items:
    st.session_state["icl_nav_root"] = "Inventory"
if st.session_state.get("icl_inventory_sub") not in ["Chemicals", "Scan"]:
    st.session_state["icl_inventory_sub"] = "Chemicals"
if st.session_state.get("icl_excel_sub") not in ["Import", "Export"]:
    st.session_state["icl_excel_sub"] = "Import"


def _icl_nav_item(label, root, sub, key):
    active = st.session_state.get("icl_nav_root") == root
    if root == "Inventory":
        active = active and st.session_state.get("icl_inventory_sub") == sub
    elif root == "Excel":
        active = active and st.session_state.get("icl_excel_sub") == sub
    if st.sidebar.button(label, key=key, type="primary" if active else "secondary", width="stretch"):
        st.session_state["icl_nav_root"] = root
        if root == "Inventory":
            st.session_state["icl_inventory_sub"] = sub
        elif root == "Excel":
            st.session_state["icl_excel_sub"] = sub
        st.rerun()


def _icl_nav_section(title):
    st.sidebar.markdown('<div class="icl-nav-section">' + title + '</div>', unsafe_allow_html=True)


st.markdown("""
<style>
[data-testid="stSidebar"] .icl-nav-section {
  margin: 24px 0 8px; padding: 0 12px;
  font-size: 10px; font-weight: 750; letter-spacing: .14em;
  color: #718096; text-transform: uppercase;
}
[data-testid="stSidebar"] button[kind="secondary"] {
  justify-content: flex-start !important; text-align: left !important;
  min-height: 40px; border: 1px solid transparent !important;
  border-radius: 7px !important; padding: 8px 14px !important;
  color: #475569 !important; background: transparent !important;
}
[data-testid="stSidebar"] button[kind="secondary"]:hover {
  color: #102c57 !important; background: #edf2f8 !important;
}
[data-testid="stSidebar"] button[kind="primary"] {
  justify-content: flex-start !important; text-align: left !important;
  min-height: 40px; border: 1px solid #d8e3f1 !important;
  border-left: 3px solid #173d70 !important;
  border-radius: 7px !important; padding: 8px 12px !important;
  background: #e9f0fa !important; color: #102c57 !important;
  font-weight: 650 !important; box-shadow: none !important;
}
[data-testid="stSidebar"] button p { text-align: left !important; }
</style>
""", unsafe_allow_html=True)

_icl_nav_section("Inventory")
_icl_nav_item("All chemicals", "Inventory", "Chemicals", "nav_chemicals")
_icl_nav_item("Add chemical", "Add Chemical", None, "nav_add")
_icl_nav_section("Storage")
_icl_nav_item("Storage locations", "Storage", None, "nav_locations")
_icl_nav_item("Storage scan", "Inventory", "Scan", "nav_scan")
_icl_nav_section("Data & safety")
_icl_nav_item("Import Excel", "Excel", "Import", "nav_import")
_icl_nav_item("Export Excel", "Excel", "Export", "nav_export")
_icl_nav_item("Safety data sheets", "SDS", None, "nav_sds")
if CURRENT_USER["role"] == "admin":
    _icl_nav_section("Administration")
    _icl_nav_item("Users", "User Management", None, "nav_users")
    _icl_nav_item("Activity log", "Activity Log", None, "nav_audit")
_icl_nav_section("Workspace")
_icl_nav_item("Settings", "Settings", None, "nav_settings")

st.sidebar.divider()
with st.sidebar.expander("Account · " + str(CURRENT_USER['name'])):
    st.caption(CURRENT_USER['email'])
    st.caption("Role: " + CURRENT_USER['role'].title())
    st.button("Log out", on_click=st.logout, width="stretch", key="nav_logout")

_root_menu = st.session_state["icl_nav_root"]

if _root_menu == "Inventory":
    # Chemicals is the unified searchable + management workspace.
    # Scan remains the only child with a separate workflow.
    menu = (
        "Inventory Scan"
        if st.session_state["icl_inventory_sub"] == "Scan"
        else "Inventory Management"
    )
elif _root_menu == "Storage":
    menu = "Storage Management"
elif _root_menu == "Excel":
    menu = (
        "Excel Import"
        if st.session_state["icl_excel_sub"] == "Import"
        else "Excel Export"
    )
else:
    menu = _root_menu

if "notice" in st.session_state:
    st.success(st.session_state.pop("notice"))


# ============================================================
# CHEMICAL SEARCH
# ============================================================

if menu == "Chemical Search":

    st.header("🔍 Chemical Search")

    st.caption(
        "Search and browse chemicals registered in the laboratory."
    )

    # ========================================================
    # FILTER OPTIONS
    # ========================================================

    conn = get_connection()

    storage_rows = conn.execute(
        """
        SELECT storage_name, COALESCE(hidden_from_lists, FALSE) AS hidden
        FROM storage_units
        ORDER BY storage_name
        """
    ).fetchall()
    hidden_storages = [
        row["storage_name"]
        for row in storage_rows
        if row["hidden"]
    ]

    manufacturer_rows = conn.execute(
        """
        SELECT DISTINCT manufacturer
        FROM chemicals
        WHERE manufacturer IS NOT NULL
          AND TRIM(manufacturer) != ''
        ORDER BY manufacturer
        """
    ).fetchall()

    conn.close()

    storage_options = [
        "All Storage"
    ] + [
        row["storage_name"]
        for row in storage_rows
    ]

    manufacturer_options = [
        "All Manufacturers"
    ] + [
        row["manufacturer"]
        for row in manufacturer_rows
    ]

    status_options = [
        "Active Only",
        "Unopened",
        "In Use",
        "Disposal Pending",
        "Disposed",
        "All"
    ]

    # ========================================================
    # SEARCH BAR
    # ========================================================

    search = st.text_input(
        "Search",
        placeholder=(
            "Chemical name, CAS No., catalog number, "
            "manufacturer, Bottle ID..."
        ),
        label_visibility="collapsed"
    )

    # ========================================================
    # FILTERS
    # ========================================================

    filter_col1, filter_col2, filter_col3 = st.columns(
        [1.4, 1.4, 1]
    )

    with filter_col1:

        selected_storage = st.selectbox(
            "Storage",
            storage_options
        )

    with filter_col2:

        selected_manufacturer = st.selectbox(
            "Manufacturer",
            manufacturer_options
        )

    with filter_col3:

        selected_status = st.selectbox(
            "Status",
            status_options
        )

    show_hidden = False
    if hidden_storages and selected_storage == "All Storage":
        show_hidden = st.checkbox(
            "Show hidden locations ("
            + ", ".join(hidden_storages)
            + ")",
            key="search_show_hidden_storage"
        )

    st.write("")

    # ========================================================
    # BUILD QUERY
    # ========================================================

    conditions = []
    params = []

    # --------------------------------------------------------
    # Search text
    # --------------------------------------------------------

    if search.strip():

        like = f"%{search.strip()}%"

        conditions.append(
            """
            (
                   c.bottle_id ILIKE ?
                OR c.legacy_id ILIKE ?
                OR c.chemical_name ILIKE ?
                OR c.description ILIKE ?
                OR c.cas_number ILIKE ?
                OR c.manufacturer ILIKE ?
                OR c.catalog_number ILIKE ?
                OR c.lot_number ILIKE ?
                OR c.owner ILIKE ?
                OR su.storage_name ILIKE ?
            )
            """
        )

        params.extend(
            [like] * 10
        )

    # --------------------------------------------------------
    # Storage filter
    # --------------------------------------------------------

    if selected_storage != "All Storage":

        conditions.append(
            "su.storage_name = ?"
        )

        params.append(
            selected_storage
        )

    elif hidden_storages and not show_hidden:

        # Locations marked "hide from lists" (e.g. a throw list) only
        # appear when selected explicitly or shown on request.
        conditions.append(
            "COALESCE(su.hidden_from_lists, FALSE) = FALSE"
        )

    # --------------------------------------------------------
    # Manufacturer filter
    # --------------------------------------------------------

    if (
        selected_manufacturer
        != "All Manufacturers"
    ):

        conditions.append(
            "c.manufacturer = ?"
        )

        params.append(
            selected_manufacturer
        )

    # --------------------------------------------------------
    # Status filter
    # --------------------------------------------------------

    if selected_status == "Active Only":

        conditions.append(
            "c.status != 'Disposed'"
        )

    elif selected_status != "All":

        conditions.append(
            "c.status = ?"
        )

        params.append(
            selected_status
        )

    # --------------------------------------------------------
    # WHERE statement
    # --------------------------------------------------------

    if conditions:

        where_sql = (
            "WHERE "
            + " AND ".join(conditions)
        )

    else:

        where_sql = ""

    # ========================================================
    # COUNT RESULTS
    # ========================================================

    conn = get_connection()

    total_count = conn.execute(
        f"""
        SELECT COUNT(*)

        FROM chemicals c

        LEFT JOIN storage_units su
            ON c.storage_unit_id = su.id

        {where_sql}
        """,
        params
    ).fetchone()[0]

    # ========================================================
    # GET RESULTS
    # ========================================================

    rows = conn.execute(
        f"""
        SELECT
            c.id,

            c.bottle_id
                AS "Bottle ID",

            c.chemical_name
                AS "Chemical Name",

            c.cas_number
                AS "CAS No.",

            c.manufacturer
                AS "Manufacturer",

            COALESCE(
                c.remaining_amount,
                c.amount
            )
                AS "Remaining",

            c.unit
                AS "Unit",

            {LOCATION_EXPR}
                AS "Storage",

            c.status
                AS "Status"

        FROM chemicals c

        LEFT JOIN storage_units su
            ON c.storage_unit_id = su.id

        {where_sql}

        ORDER BY
            c.chemical_name,
            c.id

        LIMIT 300
        """,
        params
    ).fetchall()

    conn.close()

    # ========================================================
    # RESULT COUNT
    # ========================================================

    count_col1, count_col2 = st.columns(
        [3, 1]
    )

    with count_col1:

        st.markdown(
            f"### {total_count:,} chemicals"
        )

    with count_col2:

        if total_count > 300:

            st.caption(
                "Showing first 300 results"
            )

    # ========================================================
    # SEARCH RESULTS
    # ========================================================

    if rows:

        display_data = []

        chemical_ids = []

        for row in rows:

            chemical_ids.append(
                row["id"]
            )

            remaining = row["Remaining"]

            if remaining is None:
                remaining_text = "-"

            else:

                try:
                    remaining_number = float(
                        remaining
                    )

                    if remaining_number.is_integer():
                        remaining_number = int(
                            remaining_number
                        )

                    remaining_text = (
                        f"{remaining_number} "
                        f"{row['Unit'] or ''}"
                    )

                except Exception:

                    remaining_text = (
                        f"{remaining} "
                        f"{row['Unit'] or ''}"
                    )

            display_data.append(
                {
                    "Chemical Name":
                        row["Chemical Name"],

                    "CAS No.":
                        row["CAS No."],

                    "Manufacturer":
                        row["Manufacturer"],

                    "Remaining":
                        remaining_text,

                    "Storage":
                        row["Storage"],

                    "Status":
                        row["Status"]
                }
            )

        # ----------------------------------------------------
        # Clickable dataframe
        # ----------------------------------------------------

        selection = st.dataframe(
            display_data,
            width="stretch",
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            key="chemical_search_table"
        )

        # ====================================================
        # SELECTED CHEMICAL DETAILS
        # ====================================================

        if selection.selection.rows:

            selected_row_index = (
                selection.selection.rows[0]
            )

            selected_id = chemical_ids[
                selected_row_index
            ]

            conn = get_connection()

            detail = conn.execute(
                f"""
                SELECT
                    c.*,
                    {LOCATION_EXPR}
                        AS location_text

                FROM chemicals c

                LEFT JOIN storage_units su
                    ON c.storage_unit_id = su.id

                WHERE c.id = ?
                """,
                (
                    selected_id,
                )
            ).fetchone()

            conn.close()

            if detail:
                _icl_inventory_dialog(detail)

    else:

        st.info(
            "No chemicals match the selected filters."
        )


# ============================================================
# INVENTORY SCAN
# ============================================================

elif menu == "Inventory Scan":
    _icl_storage_scan()

elif menu == "Add Chemical":
    if st.session_state.pop('icl_reset_after_save', False):
        _icl_clear_add_identity()
        for k in ('icl_scan_handoff', 'icl_add_context', 'add_ocr_result', 'add_ocr_original',
                  'add_ocr_group_photos', 'add_ocr_group_id'):
            st.session_state.pop(k, None)
    add_ocr_result, active_photos, can_register = _icl_add_photos()
    st.divider()
    st.subheader('Bottle details')
    if not can_register:
        st.info('Review and apply a product, or select Manual entry above.')

    defaults = {
        "add_name": "",
        "add_description": "",
        "add_cas": "",
        "add_manufacturer": "",
        "add_manufacturer_choice": "",
        "add_manufacturer_manual": "",
        "add_catalog": "",
        "add_lot": "",
        "add_initial_amount": 0.0,
        "add_remaining_amount": 0.0,
        "add_unit": "g",
        "add_formula": "",
        "add_mw": "",
        "add_purity": "",
        "add_owner": "",
        "add_expiration": "",
        "add_status": "Unopened",
    }
    for key_name, default_value in defaults.items():
        st.session_state.setdefault(key_name, default_value)

    if st.session_state.pop(
        "add_manufacturer_widget_needs_sync",
        False,
    ):
        sync_add_manufacturer_widget()

    storage_options = get_storage_options()
    registration_date = datetime.now(ZoneInfo("Asia/Seoul")).date()

    with st.form("add_chemical_form_" + str(st.session_state.get("icl_add_upload_nonce", 0)) + "_" + str(st.session_state.get("icl_add_context", "manual")), clear_on_submit=False, enter_to_submit=False):
        chemical_name = st.text_input("Chemical name *", key="add_name")
        left, right = st.columns(2)
        with left:
            cas_number = st.text_input("CAS No.", key="add_cas")
            manufacturer_choice = st.selectbox(
                "Manufacturer",
                ADD_CHEMICAL_MANUFACTURERS + [OTHER_MANUFACTURER_OPTION],
                key="add_manufacturer_choice",
                format_func=lambda value: "Select manufacturer" if value == "" else value,
            )
            if manufacturer_choice == OTHER_MANUFACTURER_OPTION:
                manufacturer = st.text_input("Manufacturer (manual)", key="add_manufacturer_manual")
            else:
                manufacturer = manufacturer_choice
            catalog_number = st.text_input("Catalog No.", key="add_catalog")
        with right:
            ac, uc = st.columns([2, 1])
            with ac:
                initial_amount = st.number_input("Amount", min_value=0.0, key="add_initial_amount")
            with uc:
                unit = st.selectbox("Unit", ["g", "mg", "kg", "mL", "L", "ea"], key="add_unit")
            storage = st.selectbox(
                "Storage location", storage_options,
                format_func=lambda option: option[2], key="add_storage",
            )
        remaining_amount = st.number_input("Remaining amount", min_value=0.0, key="add_remaining_amount")
        full_bottle = st.checkbox("Unopened bottle · Set remaining amount equal to the amount above", value=False)
        if full_bottle:
            remaining_amount = initial_amount
        with st.expander("More details", expanded=False):
            description = st.text_input("Description / grade", key="add_description")
            dl, dr = st.columns(2)
            with dl:
                lot_number = st.text_input("Lot No.", key="add_lot")
                owner = st.text_input("Owner / person in charge", key="add_owner")
                expiration_date = st.text_input("Expiration date", placeholder="YYYY-MM-DD", key="add_expiration")
            with dr:
                molecular_formula = st.text_input("Molecular formula", key="add_formula")
                molecular_weight = st.text_input("Molecular weight", key="add_mw")
                purity = st.text_input("Purity", key="add_purity")
                status = st.selectbox("Status", ["Unopened", "In Use", "Disposal Pending"], key="add_status")
            st.date_input("Registration date", value=registration_date, disabled=True, format="YYYY-MM-DD")
        registration_confirmed = st.checkbox('I checked the name, CAS, amount and storage location. Register this as a new bottle.', value=False)
        submitted = st.form_submit_button("Register new bottle", type="primary", disabled=not can_register)

    if submitted:
        if not can_register or not registration_confirmed:
            st.error('Review the registration details and select the confirmation checkbox.')
        elif not chemical_name.strip():
            st.error('Enter a chemical name.')
        elif remaining_amount > initial_amount:
            st.error('Remaining amount cannot exceed the initial amount.')
        else:
            timestamp = datetime.now().strftime('%Y-%m-%d %H:%M')
            values = dict(chemical_name=chemical_name.strip(), description=description.strip(),
                cas_number=cas_number.strip(), manufacturer=manufacturer.strip(), catalog_number=catalog_number.strip(),
                lot_number=lot_number.strip(), amount=initial_amount, initial_amount=initial_amount,
                remaining_amount=remaining_amount, unit=unit, storage_unit_id=storage[0], shelf_number=storage[1],
                location=storage[2] if storage[0] is not None else '', owner=owner.strip(),
                expiration_date=expiration_date.strip(), status=status,
                source_registered_date=registration_date.strftime('%Y-%m-%d'), molecular_formula=molecular_formula.strip(),
                molecular_weight=molecular_weight.strip(), purity=purity.strip(), created_at=timestamp, updated_at=timestamp)
            draft = st.session_state.get('icl_v8_registration_draft') if add_ocr_result else None
            photos = st.session_state.get('icl_applied_photos', []) if add_ocr_result else []
            source_scan_id = st.session_state.get('add_source_scan_id')
            try:
                bottle_id = _icl_register_new_bottle(values, storage, photos, draft, source_scan_id)
            except Exception as exc:
                st.error('Registration failed: ' + _icl_exception_message(exc))
            else:
                for record in st.session_state.get('inventory_scan_batch', []):
                    if record['scan_id'] == source_scan_id:
                        record['registered_bottle_id'] = bottle_id
                # Remove the completed group; preserve other bottles awaiting review.
                completed_hashes = {p.get('scan_id') for p in photos}
                st.session_state['icl_add_batch'] = [r for r in st.session_state.get('icl_add_batch', [])
                    if r.get('scan_id') not in completed_hashes]
                st.session_state.pop('icl_add_group_selected', None)
                st.session_state['icl_saved_notice'] = f'{chemical_name} · {bottle_id} registered ({len(photos)} photos attached)'
                try:
                    with st.spinner('Checking safety information…'):
                        safety = fetch_pubchem_safety(cas_number, chemical_name)
                    if safety:
                        save_pubchem_safety(int(bottle_id.rsplit('-', 1)[1]), safety)
                    else:
                        st.session_state['icl_saved_notice'] += ' · Please check the safety information in SDS.'
                except Exception:
                    st.session_state['icl_saved_notice'] += ' · Safety lookup failed. Please check SDS.'
                st.session_state['icl_reset_after_save'] = True
                st.session_state['icl_add_upload_nonce'] = st.session_state.get('icl_add_upload_nonce', 0) + 1
                st.rerun()

elif menu == "Excel Import":
    st.header("📥 Excel Import")
    st.caption(
        "Designed for the current Lab Manager export. The first row is a title and the second row contains the column headers."
    )

    st.info(
        "Storage names from the Excel file will be linked to Storage Units. "
        "If a storage name does not exist yet, it will be created automatically. "
        "You can rename the Storage Unit later without editing every chemical."
    )

    uploaded_excel = st.file_uploader(
        "Upload Lab Manager Excel File",
        type=["xlsx"],
    )

    if uploaded_excel is not None:
        try:
            # The current export has title in row 1 and headers in row 2.
            df = pd.read_excel(
                uploaded_excel,
                sheet_name="item list",
                header=1,
                dtype=object,
            )
            df = df.dropna(how="all")
            df.columns = [str(column).strip() for column in df.columns]

            if "물품명" not in df.columns:
                st.error(
                    "The expected source Item Name column was not found. Please check that this is a Lab Manager item-list export."
                )
            else:
                only_reagents = st.checkbox(
                    "Import only rows categorized as reagents",
                    value=True,
                )

                working_df = df.copy()

                if only_reagents and "카테고리 이름" in working_df.columns:
                    working_df = working_df[
                        working_df["카테고리 이름"]
                        .fillna("")
                        .astype(str)
                        .str.strip()
                        .eq("시약")
                    ]

                working_df = working_df[
                    working_df["물품명"].fillna("").astype(str).str.strip().ne("")
                ].copy()

                st.success(
                    f"✅ File loaded. {len(working_df):,} chemical row(s) selected for review."
                )

                # ------------------------------------------------
                # Prepare existing legacy IDs for duplicate check
                # ------------------------------------------------
                conn = get_connection()
                existing_legacy_ids = {
                    row["legacy_id"]
                    for row in conn.execute(
                        """
                        SELECT legacy_id
                        FROM chemicals
                        WHERE legacy_id IS NOT NULL
                          AND TRIM(legacy_id) != ''
                        """
                    ).fetchall()
                }
                existing_storage_names = {
                    row["storage_name"].lower(): row["storage_name"]
                    for row in conn.execute(
                        "SELECT storage_name FROM storage_units"
                    ).fetchall()
                }
                conn.close()

                prepared_rows = []
                skipped_existing = 0
                skipped_duplicate_in_file = 0
                seen_legacy_ids = set()
                new_storage_names = set()

                for _, source in working_df.iterrows():
                    legacy_id = clean_text(source.get("관리번호"))

                    if legacy_id:
                        if legacy_id in existing_legacy_ids:
                            skipped_existing += 1
                            continue
                        if legacy_id in seen_legacy_ids:
                            skipped_duplicate_in_file += 1
                            continue
                        seen_legacy_ids.add(legacy_id)

                    storage_name = clean_text(source.get("보관함 경로"))
                    if storage_name and storage_name.lower() not in existing_storage_names:
                        new_storage_names.add(storage_name)

                    initial_amount = clean_number(source.get("물품사이즈(용량)"))
                    remaining_amount = clean_number(source.get("잔여 용량"))
                    opened_date = clean_date(source.get("개봉일"))

                    if remaining_amount is None:
                        remaining_amount = initial_amount

                    status = "In Use" if opened_date else "Unopened"

                    prepared_rows.append(
                        {
                            "category_classification": clean_text(
                                source.get("카테고리 분류")
                            ),
                            "category_name": clean_text(source.get("카테고리 이름")),
                            "legacy_id": legacy_id,
                            "company_management_id": clean_text(
                                source.get("사내관리번호")
                            ),
                            "chemical_name": clean_text(source.get("물품명")),
                            "description": clean_text(source.get("서브네임")),
                            "legacy_location_path": clean_text(source.get("위치 경로")),
                            "legacy_storage_name": storage_name,
                            "manufacturer": clean_text(source.get("브랜드")),
                            "registered_by": clean_text(source.get("등록자")),
                            "owner": clean_text(source.get("담당자")),
                            "purchaser": clean_text(source.get("구매자")),
                            "purchase_date": clean_date(source.get("구매일")),
                            "purchase_price": clean_text(source.get("구매 가격")),
                            "source_registered_date": clean_date(source.get("등록일")),
                            "source_last_modified": clean_date(source.get("최근 변경일시")),
                            "comments": clean_text(source.get("코멘트")),
                            "tags": clean_text(source.get("태그")),
                            "catalog_number": clean_text(source.get("제품번호")),
                            "lot_number": clean_text(source.get("Lot No.")),
                            "cas_number": clean_text(source.get("Cas No.")),
                            "initial_amount": initial_amount,
                            "remaining_amount": remaining_amount,
                            "unit": normalize_unit(source.get("물품사이즈(용량) 단위")),
                            "safety_stock": clean_text(source.get("안전재고량")),
                            "concentration_density": clean_text(source.get("농도(밀도)")),
                            "molecular_weight": clean_text(source.get("분자량")),
                            "mol": clean_text(source.get("MOL")),
                            "purity": clean_text(source.get("순도")),
                            "solubility": clean_text(source.get("용해도")),
                            "molecular_formula": clean_text(source.get("분자식")),
                            "storage_temperature": clean_text(source.get("보관 온도")),
                            "flash_point": clean_text(source.get("발화점")),
                            "boiling_point": clean_text(source.get("끓는점")),
                            "melting_point": clean_text(source.get("녹는점")),
                            "opened_date": opened_date,
                            "expiration_date": clean_date(source.get("유효기간")),
                            "status": status,
                        }
                    )

                # ------------------------------------------------
                # Summary
                # ------------------------------------------------
                summary_cols = st.columns(4)
                summary_cols[0].metric("Ready to Import", f"{len(prepared_rows):,}")
                summary_cols[1].metric("Already in DB", f"{skipped_existing:,}")
                summary_cols[2].metric(
                    "Duplicate IDs in File", f"{skipped_duplicate_in_file:,}"
                )
                summary_cols[3].metric("New Storage Units", len(new_storage_names))

                if new_storage_names:
                    with st.expander("Storage Units that will be created automatically"):
                        for name in sorted(new_storage_names):
                            st.write(f"• {name}")

                # ------------------------------------------------
                # Preview
                # ------------------------------------------------
                st.subheader("📋 Import Preview")

                preview = pd.DataFrame(prepared_rows)
                if not preview.empty:
                    preview = preview.rename(
                        columns={
                            "legacy_id": "Legacy ID",
                            "chemical_name": "Chemical Name",
                            "description": "Description / Grade",
                            "cas_number": "CAS No.",
                            "manufacturer": "Manufacturer",
                            "catalog_number": "Catalog No.",
                            "initial_amount": "Initial Amount",
                            "remaining_amount": "Remaining Amount",
                            "unit": "Unit",
                            "legacy_storage_name": "Storage Name",
                            "owner": "Owner",
                            "opened_date": "Opened Date",
                            "expiration_date": "Expiration Date",
                            "status": "Status",
                        }
                    )
                    preview_columns = [
                        "Legacy ID",
                        "Chemical Name",
                        "Description / Grade",
                        "CAS No.",
                        "Manufacturer",
                        "Catalog No.",
                        "Initial Amount",
                        "Remaining Amount",
                        "Unit",
                        "Storage Name",
                        "Owner",
                        "Opened Date",
                        "Expiration Date",
                        "Status",
                    ]
                    st.dataframe(
                        preview[preview_columns].head(200),
                        width="stretch",
                        hide_index=True,
                    )
                    if len(preview) > 200:
                        st.caption(
                            f"Preview shows the first 200 of {len(preview):,} rows. All reviewed rows will be imported."
                        )

                    st.divider()
                    confirm_import = st.checkbox(
                        "I reviewed the preview and want to import these chemicals."
                    )

                    if st.button(
                        "📥 Import Chemicals",
                        type="primary",
                        disabled=not confirm_import,
                    ):
                        conn = get_connection()
                        imported_count = 0

                        try:
                            for item in prepared_rows:
                                storage_unit_id = get_or_create_storage(
                                    conn, item["legacy_storage_name"]
                                )

                                cursor = conn.execute(
                                    """
                                    INSERT INTO chemicals (
                                        category_classification,
                                        category_name,
                                        legacy_id,
                                        company_management_id,
                                        chemical_name,
                                        description,
                                        legacy_location_path,
                                        legacy_storage_name,
                                        manufacturer,
                                        registered_by,
                                        owner,
                                        purchaser,
                                        purchase_date,
                                        purchase_price,
                                        source_registered_date,
                                        source_last_modified,
                                        comments,
                                        tags,
                                        catalog_number,
                                        lot_number,
                                        cas_number,
                                        amount,
                                        initial_amount,
                                        remaining_amount,
                                        unit,
                                        safety_stock,
                                        concentration_density,
                                        molecular_weight,
                                        mol,
                                        purity,
                                        solubility,
                                        molecular_formula,
                                        storage_temperature,
                                        flash_point,
                                        boiling_point,
                                        melting_point,
                                        opened_date,
                                        expiration_date,
                                        storage_unit_id,
                                        shelf_number,
                                        status,
                                        created_at,
                                        updated_at
                                    )
                                    VALUES (
                                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                                        ?, ?, ?, ?, ?, ?, ?
                                    )
                                    """,
                                    (
                                        item["category_classification"],
                                        item["category_name"],
                                        item["legacy_id"],
                                        item["company_management_id"],
                                        item["chemical_name"],
                                        item["description"],
                                        item["legacy_location_path"],
                                        item["legacy_storage_name"],
                                        item["manufacturer"],
                                        item["registered_by"],
                                        item["owner"],
                                        item["purchaser"],
                                        item["purchase_date"],
                                        item["purchase_price"],
                                        item["source_registered_date"],
                                        item["source_last_modified"],
                                        item["comments"],
                                        item["tags"],
                                        item["catalog_number"],
                                        item["lot_number"],
                                        item["cas_number"],
                                        item["initial_amount"],
                                        item["initial_amount"],
                                        item["remaining_amount"],
                                        item["unit"],
                                        item["safety_stock"],
                                        item["concentration_density"],
                                        item["molecular_weight"],
                                        item["mol"],
                                        item["purity"],
                                        item["solubility"],
                                        item["molecular_formula"],
                                        item["storage_temperature"],
                                        item["flash_point"],
                                        item["boiling_point"],
                                        item["melting_point"],
                                        item["opened_date"],
                                        item["expiration_date"],
                                        storage_unit_id,
                                        None,
                                        item["status"],
                                        datetime.now().strftime("%Y-%m-%d %H:%M"),
                                        datetime.now().strftime("%Y-%m-%d %H:%M"),
                                    ),
                                )

                                chemical_id = cursor.lastrowid
                                bottle_id = f"ICL-{chemical_id:06d}"
                                conn.execute(
                                    "UPDATE chemicals SET bottle_id=? WHERE id=?",
                                    (bottle_id, chemical_id),
                                )
                                imported_count += 1

                            write_audit_log(
                                "EXCEL_IMPORT",
                                target_type="bulk_import",
                                details=json.dumps(
                                    {
                                        "imported_count": imported_count,
                                        "source": uploaded_excel.name,
                                    },
                                    ensure_ascii=False,
                                ),
                                conn=conn,
                            )
                            conn.commit()
                            set_notice(
                                f"✅ {imported_count:,} chemical(s) imported successfully."
                            )
                            st.rerun()

                        except Exception as exc:
                            conn.rollback()
                            st.error(f"Import failed. No changes were committed.\n\n{exc}")
                        finally:
                            conn.close()
                else:
                    st.info("There are no new rows to import.")

        except ValueError as exc:
            st.error(
                "The 'item list' sheet could not be found. Please upload the original Lab Manager export."
            )
            st.code(str(exc))
        except Exception as exc:
            st.error(f"Unable to read the Excel file: {exc}")


# ============================================================
# STORAGE MANAGEMENT
# ============================================================

elif menu == "Storage Management":
    if st.button('Open Storage Scan', key='storage_open_scan'):
        st.session_state['_pending_navigation'] = 'Inventory Scan'
        st.rerun()


    st.header("🗄️ Storage Management")

    st.caption(
        "Manage storage units, shelves, and chemical locations."
    )

    # ========================================================
    # TABS
    # ========================================================

    tab1, tab2 = st.tabs(
        [
            "🗄️ Storage Overview",
            "➕ Add Storage Unit"
        ]
    )

    # ========================================================
    # TAB 1 - STORAGE OVERVIEW
    # ========================================================

    with tab1:

        # ----------------------------------------------------
        # Load storage units
        # ----------------------------------------------------

        conn = get_connection()

        storage_units = conn.execute(
            """
            SELECT
                su.id,
                su.storage_name,
                su.notes,
                COALESCE(su.hidden_from_lists, FALSE)
                    AS hidden_from_lists,

                COUNT(
                    DISTINCT ss.id
                ) AS shelf_count,

                COUNT(
                    DISTINCT CASE
                        WHEN c.status != 'Disposed'
                        THEN c.id
                    END
                ) AS chemical_count

            FROM storage_units su

            LEFT JOIN storage_shelves ss
                ON su.id = ss.storage_unit_id

            LEFT JOIN chemicals c
                ON su.id = c.storage_unit_id

            GROUP BY
                su.id,
                su.storage_name,
                su.notes,
                su.hidden_from_lists

            ORDER BY
                su.storage_name
            """
        ).fetchall()

        conn.close()

        # ----------------------------------------------------
        # Summary
        # ----------------------------------------------------

        if storage_units:

            total_chemicals = sum(
                storage["chemical_count"]
                for storage in storage_units
            )

            col1, col2 = st.columns(2)

            col1.metric(
                "Storage Units",
                len(storage_units)
            )

            col2.metric(
                "Chemicals in Storage",
                total_chemicals
            )

            st.write("")

            # =================================================
            # STORAGE CARDS
            # =================================================

            # Two cards per row
            for i in range(
                0,
                len(storage_units),
                2
            ):

                card_columns = st.columns(2)

                for j in range(2):

                    index = i + j

                    if index >= len(storage_units):
                        break

                    storage = storage_units[index]

                    with card_columns[j]:

                        with st.container(
                            border=True
                        ):

                            st.subheader(
                                f"🗄️ {storage['storage_name']}"
                            )

                            if storage["notes"]:

                                st.caption(
                                    storage["notes"]
                                )

                            if storage["hidden_from_lists"]:

                                st.caption(
                                    "Hidden from Chemical Search lists"
                                )

                            metric1, metric2 = (
                                st.columns(2)
                            )

                            metric1.metric(
                                "Chemicals",
                                storage[
                                    "chemical_count"
                                ]
                            )

                            metric2.metric(
                                "Shelves",
                                storage[
                                    "shelf_count"
                                ]
                            )

                            button1, button2 = (
                                st.columns(2)
                            )

                            # ------------------------------
                            # OPEN BUTTON
                            # ------------------------------

                            if button1.button(
                                "Open",
                                key=(
                                    f"open_storage_"
                                    f"{storage['id']}"
                                ),
                                width="stretch"
                            ):

                                st.session_state[
                                    "selected_storage"
                                ] = storage["id"]

                                st.session_state.pop(
                                    "edit_storage",
                                    None
                                )

                                st.rerun()

                            # ------------------------------
                            # EDIT BUTTON
                            # ------------------------------

                            if button2.button(
                                "Edit",
                                key=(
                                    f"edit_storage_"
                                    f"{storage['id']}"
                                ),
                                width="stretch"
                            ):

                                st.session_state[
                                    "edit_storage"
                                ] = storage["id"]

                                st.session_state.pop(
                                    "selected_storage",
                                    None
                                )

                                st.rerun()

        else:

            st.info(
                "No storage units have been created yet."
            )

        # ====================================================
        # STORAGE DETAIL — popup
        # ====================================================
        if "selected_storage" in st.session_state:
            _icl_storage_dialog(st.session_state["selected_storage"])

        # ====================================================
        # EDIT STORAGE
        # ====================================================

        if "edit_storage" in st.session_state:

            storage_id = st.session_state[
                "edit_storage"
            ]

            conn = get_connection()

            storage = conn.execute(
                """
                SELECT
                    id,
                    storage_name,
                    notes,
                    COALESCE(hidden_from_lists, FALSE)
                        AS hidden_from_lists
                FROM storage_units
                WHERE id = ?
                """,
                (storage_id,)
            ).fetchone()

            conn.close()

            if storage:

                st.divider()

                st.subheader(
                    f"✏️ Edit Storage Unit"
                )

                st.caption(
                    "Renaming a storage unit will "
                    "automatically update the displayed "
                    "location for all linked chemicals."
                )

                with st.form(
                    f"edit_storage_form_{storage_id}",
                    enter_to_submit=False
                ):

                    new_storage_name = (
                        st.text_input(
                            "Storage Name *",
                            value=storage[
                                "storage_name"
                            ]
                        )
                    )

                    new_notes = st.text_area(
                        "Notes",
                        value=(
                            storage["notes"]
                            or ""
                        )
                    )

                    new_hidden = st.checkbox(
                        "Hide from Chemical Search lists "
                        "(e.g. throw / disposal lists)",
                        value=bool(
                            storage["hidden_from_lists"]
                        ),
                        help=(
                            "Bottles here are left out of the "
                            "default Chemical Search list. Select "
                            "this location or tick 'Show hidden "
                            "locations' to see them."
                        )
                    )

                    save_col, cancel_col = (
                        st.columns(2)
                    )

                    with save_col:

                        save_storage = (
                            st.form_submit_button(
                                "Save Changes",
                                type="primary"
                            )
                        )

                    with cancel_col:

                        cancel_storage = (
                            st.form_submit_button(
                                "Cancel"
                            )
                        )

                    # -----------------------------------------
                    # SAVE
                    # -----------------------------------------

                    if save_storage:

                        if not new_storage_name.strip():

                            st.error(
                                "Storage name is required."
                            )

                        else:

                            conn = get_connection()

                            duplicate = conn.execute(
                                """
                                SELECT id
                                FROM storage_units

                                WHERE
                                    LOWER(storage_name)
                                    = LOWER(?)

                                    AND id != ?
                                """,
                                (
                                    new_storage_name.strip(),
                                    storage_id
                                )
                            ).fetchone()

                            if duplicate:

                                st.error(
                                    "Another storage unit "
                                    "already uses this name."
                                )

                            else:

                                conn.execute(
                                    """
                                    UPDATE storage_units

                                    SET
                                        storage_name = ?,
                                        notes = ?,
                                        hidden_from_lists = ?

                                    WHERE id = ?
                                    """,
                                    (
                                        new_storage_name.strip(),
                                        new_notes.strip(),
                                        bool(new_hidden),
                                        storage_id
                                    )
                                )
                                write_audit_log(
                                    "EDIT_STORAGE",
                                    target_type="storage_unit",
                                    target_id=storage_id,
                                    details=json.dumps(
                                        {
                                            "old_name": storage["storage_name"],
                                            "new_name": new_storage_name.strip(),
                                            "notes": new_notes.strip(),
                                            "hidden_from_lists": bool(new_hidden),
                                        },
                                        ensure_ascii=False,
                                    ),
                                    conn=conn,
                                )

                                conn.commit()

                                del st.session_state[
                                    "edit_storage"
                                ]

                                st.success(
                                    "✅ Storage unit updated."
                                )

                            conn.close()

                            if not duplicate:
                                st.rerun()

                    # -----------------------------------------
                    # CANCEL
                    # -----------------------------------------

                    if cancel_storage:

                        del st.session_state[
                            "edit_storage"
                        ]

                        st.rerun()


    # ========================================================
    # TAB 2 - ADD STORAGE UNIT
    # ========================================================

    with tab2:

        st.subheader(
            "➕ Add Storage Unit"
        )

        st.caption(
            "Create a new cabinet, refrigerator, "
            "desiccator, glovebox, or other storage location."
        )

        with st.form(
            "add_storage_unit_form",
            clear_on_submit=True,
            enter_to_submit=False
        ):

            storage_name = st.text_input(
                "Storage Name *",
                placeholder=(
                    "e.g. Desiccator, Refrigerator 1, "
                    "Chemical Cabinet 1"
                )
            )

            shelf_count = st.number_input(
                "Number of Shelves",
                min_value=0,
                max_value=30,
                value=0,
                step=1,
                help=(
                    "Enter 0 if the storage unit "
                    "does not use shelves."
                )
            )

            notes = st.text_area(
                "Notes",
                placeholder=(
                    "Optional description or "
                    "storage information"
                )
            )

            create_storage = (
                st.form_submit_button(
                    "Create Storage Unit",
                    type="primary"
                )
            )

            if create_storage:

                if not storage_name.strip():

                    st.error(
                        "Storage name is required."
                    )

                else:

                    conn = get_connection()

                    duplicate = conn.execute(
                        """
                        SELECT id
                        FROM storage_units

                        WHERE
                            LOWER(storage_name)
                            = LOWER(?)
                        """,
                        (
                            storage_name.strip(),
                        )
                    ).fetchone()

                    if duplicate:

                        st.error(
                            "A storage unit with "
                            "this name already exists."
                        )

                    else:

                        cursor = conn.execute(
                            """
                            INSERT INTO storage_units (
                                storage_name,
                                notes,
                                created_at
                            )

                            VALUES (?, ?, ?)
                            """,
                            (
                                storage_name.strip(),
                                notes.strip(),
                                datetime.now().strftime(
                                    "%Y-%m-%d %H:%M"
                                )
                            )
                        )

                        new_storage_id = (
                            cursor.lastrowid
                        )

                        # Create shelves
                        for shelf in range(
                            1,
                            int(shelf_count) + 1
                        ):

                            conn.execute(
                                """
                                INSERT INTO storage_shelves (
                                    storage_unit_id,
                                    shelf_number
                                )

                                VALUES (?, ?)
                                """,
                                (
                                    new_storage_id,
                                    shelf
                                )
                            )

                        write_audit_log(
                            "ADD_STORAGE",
                            target_type="storage_unit",
                            target_id=new_storage_id,
                            details=json.dumps(
                                {
                                    "storage_name": storage_name.strip(),
                                    "shelf_count": int(shelf_count),
                                    "notes": notes.strip(),
                                },
                                ensure_ascii=False,
                            ),
                            conn=conn,
                        )
                        conn.commit()
                        conn.close()

                        st.success(
                            f"✅ {storage_name} created."
                        )

                        st.rerun()


# ============================================================
# EXCEL EXPORT
# ============================================================

elif menu == "Excel Export":
    st.header("📤 Excel Export")
    st.caption("Export the current chemical inventory as an Excel workbook.")

    scope = st.selectbox(
        "Export Scope",
        ["Active Inventory", "All Chemicals", "Disposed Chemicals"],
    )

    include_details = st.checkbox(
        "Include detailed chemical information",
        value=True,
        help=(
            "Includes formula, molecular weight, purity, storage temperature, "
            "purchase/opened/expiration dates, comments, and disposal information."
        ),
    )

    if scope == "Active Inventory":
        status_clause = "WHERE c.status != 'Disposed'"
        sheet_name = "Active Inventory"
        file_label = "Active_Inventory"
    elif scope == "Disposed Chemicals":
        status_clause = "WHERE c.status = 'Disposed'"
        sheet_name = "Disposed Chemicals"
        file_label = "Disposed_Chemicals"
    else:
        status_clause = ""
        sheet_name = "All Chemicals"
        file_label = "All_Chemicals"

    basic_columns = f"""
        SELECT
            c.bottle_id AS "Bottle ID",
            c.legacy_id AS "Legacy ID",
            c.chemical_name AS "Chemical Name",
            c.description AS "Description / Grade",
            c.cas_number AS "CAS No.",
            c.manufacturer AS "Manufacturer",
            c.catalog_number AS "Catalog No.",
            c.lot_number AS "Lot No.",
            c.initial_amount AS "Initial Amount",
            COALESCE(c.remaining_amount, c.amount) AS "Remaining Amount",
            c.unit AS "Unit",
            {LOCATION_EXPR} AS "Storage Location",
            c.owner AS "Owner",
            c.status AS "Status"
    """

    detailed_columns = f"""
        SELECT
            c.bottle_id AS "Bottle ID",
            c.legacy_id AS "Legacy ID",
            c.company_management_id AS "Company Management ID",
            c.category_classification AS "Category Classification",
            c.category_name AS "Category Name",
            c.chemical_name AS "Chemical Name",
            c.description AS "Description / Grade",
            c.cas_number AS "CAS No.",
            c.manufacturer AS "Manufacturer",
            c.catalog_number AS "Catalog No.",
            c.lot_number AS "Lot No.",
            c.initial_amount AS "Initial Amount",
            COALESCE(c.remaining_amount, c.amount) AS "Remaining Amount",
            c.unit AS "Unit",
            {LOCATION_EXPR} AS "Storage Location",
            c.shelf_number AS "Shelf Number",
            c.owner AS "Owner",
            c.registered_by AS "Registered By",
            c.purchaser AS "Purchaser",
            c.purchase_date AS "Purchase Date",
            c.purchase_price AS "Purchase Price",
            c.opened_date AS "Opened Date",
            c.expiration_date AS "Expiration Date",
            c.concentration_density AS "Concentration / Density",
            c.molecular_weight AS "Molecular Weight",
            c.mol AS "MOL",
            c.purity AS "Purity",
            c.solubility AS "Solubility",
            c.molecular_formula AS "Molecular Formula",
            c.storage_temperature AS "Storage Temperature",
            c.flash_point AS "Flash Point",
            c.boiling_point AS "Boiling Point",
            c.melting_point AS "Melting Point",
            c.comments AS "Comments",
            c.tags AS "Tags",
            c.status AS "Status",
            c.created_at AS "Created At",
            c.updated_at AS "Updated At",
            c.disposed_at AS "Disposed At",
            c.disposal_reason AS "Disposal Reason",
            c.disposal_note AS "Disposal Note",
            c.disposal_location AS "Disposal Location"
    """

    select_sql = detailed_columns if include_details else basic_columns

    conn = get_connection()
    export_rows = conn.execute(
        f"""
        {select_sql}
        FROM chemicals c
        LEFT JOIN storage_units su ON c.storage_unit_id = su.id
        {status_clause}
        ORDER BY c.chemical_name, c.id
        """
    ).fetchall()
    conn.close()

    # Build the DataFrame from psycopg rows directly.
    # pandas.read_sql_query expects a DB-API connection exposing .cursor(),
    # while this app intentionally uses CompatConnection.
    export_df = pd.DataFrame([dict(row) for row in export_rows])

    summary_col1, summary_col2 = st.columns(2)
    summary_col1.metric("Chemicals to Export", f"{len(export_df):,}")
    summary_col2.metric("Columns", len(export_df.columns))

    st.subheader("Preview")

    if export_df.empty:
        st.info("No chemicals are available for the selected export scope.")
    else:
        st.dataframe(export_df.head(100), width="stretch", hide_index=True)

        if len(export_df) > 100:
            st.caption(
                f"Preview shows the first 100 of {len(export_df):,} rows. "
                "The Excel file will include all rows."
            )

        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            export_df.to_excel(writer, index=False, sheet_name=sheet_name[:31])
            worksheet = writer.book[sheet_name[:31]]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            header_fill = PatternFill(fill_type="solid", fgColor="25324B")
            header_font = Font(bold=True, color="FFFFFF")

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter

                for cell in column_cells[:300]:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))

                worksheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 10), 35
                )

        output.seek(0)
        export_date = datetime.now().strftime("%Y-%m-%d")
        filename = f"ICL_{file_label}_{export_date}.xlsx"

        st.download_button(
            "📥 Download Excel File",
            data=output.getvalue(),
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            width="stretch",
        )

        st.caption(
            "The export is generated from the current database at the moment you click Download."
        )


# ============================================================
# INVENTORY MANAGEMENT
# ============================================================

elif menu == "Inventory Management":
    st.header("Chemicals")

    tab1, tab2 = st.tabs(["📦 Active Inventory", "📋 Disposal History"])

    # --------------------------------------------------------
    # Active inventory
    # --------------------------------------------------------
    with tab1:
        inventory_search = st.text_input(
            "Search Active Inventory",
            placeholder="Bottle ID, chemical, CAS No., manufacturer, owner, location...",
        )

        search_clause = ""
        params = []
        if inventory_search:
            like = f"%{inventory_search}%"
            search_clause = """
            AND (
                   c.bottle_id ILIKE ?
                OR c.legacy_id ILIKE ?
                OR c.chemical_name ILIKE ?
                OR c.cas_number ILIKE ?
                OR c.manufacturer ILIKE ?
                OR c.owner ILIKE ?
                OR su.storage_name ILIKE ?
            )
            """
            params = [like] * 7

        conn = get_connection()
        total_rows = conn.execute(
            f"""
            SELECT COUNT(*)
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed'
            {search_clause}
            """,
            params,
        ).fetchone()[0]

        page_size = 30
        max_page = max(1, (total_rows + page_size - 1) // page_size)
        page = st.selectbox("Page", list(range(1, max_page + 1)))
        offset = (page - 1) * page_size

        rows = conn.execute(
            f"""
            SELECT
                c.id,
                c.bottle_id,
                c.legacy_id,
                c.chemical_name,
                c.description,
                c.cas_number,
                c.manufacturer,
                c.catalog_number,
                c.lot_number,
                c.initial_amount,
                COALESCE(c.remaining_amount, c.amount) AS remaining_amount,
                c.unit,
                c.owner,
                c.expiration_date,
                c.status,
                COALESCE(
                    NULLIF(c.source_registered_date, ''),
                    LEFT(c.created_at, 10),
                    '-'
                ) AS registered_date,
                c.storage_unit_id,
                c.shelf_number,
                {LOCATION_EXPR} AS location_text
            FROM chemicals c
            LEFT JOIN storage_units su ON c.storage_unit_id = su.id
            WHERE c.status != 'Disposed'
            {search_clause}
            ORDER BY c.chemical_name, c.id
            LIMIT ? OFFSET ?
            """,
            params + [page_size, offset],
        ).fetchall()
        conn.close()

        st.caption(f"{total_rows:,} active chemical(s)")

        if rows:
            display_rows = []

            for chemical in rows:
                amount_text = "-"
                if chemical["remaining_amount"] is not None:
                    amount_text = (
                        f"{chemical['remaining_amount']:g} "
                        f"{chemical['unit'] or ''}"
                    )

                display_rows.append(
                    {
                        "Bottle ID": chemical["bottle_id"] or "-",
                        "Chemical": chemical["chemical_name"] or "-",
                        "Manufacturer": chemical["manufacturer"] or "-",
                        "Remaining": amount_text,
                        "Location": chemical["location_text"] or "-",
                        "Status": chemical["status"] or "-",
                        "Registered": chemical["registered_date"] or "-",
                    }
                )

            inventory_selection = st.dataframe(
                display_rows,
                width="stretch",
                height=470,
                hide_index=True,
                on_select="rerun",
                selection_mode="single-row",
                key="inventory_management_table",
            )

            selected_chemical = None
            if inventory_selection.selection.rows:
                selected_index = inventory_selection.selection.rows[0]
                selected_chemical = rows[selected_index]

                selected_id = selected_chemical["id"]

                # Open the information popup only when the selected bottle
                # changes. This prevents the dialog from reopening on every
                # Edit/Dispose interaction while preserving row-click behavior.
                if st.session_state.get("_icl_last_detail_id") != selected_id:
                    conn = get_connection()
                    detail = conn.execute(
                        f"""
                        SELECT
                            c.*,
                            {LOCATION_EXPR} AS location_text
                        FROM chemicals c
                        LEFT JOIN storage_units su
                            ON c.storage_unit_id = su.id
                        WHERE c.id = ?
                        """,
                        (selected_id,),
                    ).fetchone()
                    conn.close()

                    if detail:
                        st.session_state["_icl_last_detail_id"] = selected_id
                        _icl_inventory_dialog(detail)
            else:
                st.session_state.pop("_icl_last_detail_id", None)

            # Search and management are merged, but management itself is not
            # removed: actions remain visible at all times and activate when a
            # bottle row is selected.
            with st.container(border=True):
                st.markdown("**Bottle actions**")

                if selected_chemical is None:
                    st.caption("Select a bottle above to view or manage it.")
                else:
                    st.caption(
                        f"{selected_chemical['bottle_id'] or '-'} · "
                        f"{selected_chemical['chemical_name'] or '-'} · "
                        f"{selected_chemical['manufacturer'] or '-'} · "
                        f"{selected_chemical['location_text'] or '-'}"
                    )

                action_col1, action_col2, action_col3 = st.columns(3)

                if action_col1.button(
                    "View details",
                    key="inventory_selected_view",
                    width="stretch",
                    disabled=(selected_chemical is None),
                ):
                    conn = get_connection()
                    detail = conn.execute(
                        f"""
                        SELECT
                            c.*,
                            {LOCATION_EXPR} AS location_text
                        FROM chemicals c
                        LEFT JOIN storage_units su
                            ON c.storage_unit_id = su.id
                        WHERE c.id = ?
                        """,
                        (selected_chemical["id"],),
                    ).fetchone()
                    conn.close()
                    if detail:
                        _icl_inventory_dialog(detail)

                if action_col2.button(
                    "Edit bottle",
                    key="inventory_selected_edit",
                    width="stretch",
                    disabled=(selected_chemical is None),
                ):
                    edit_chemical_dialog(selected_chemical["id"])

                if action_col3.button(
                    "Dispose",
                    key="inventory_selected_dispose",
                    width="stretch",
                    disabled=(selected_chemical is None),
                ):
                    dispose_chemical_dialog(selected_chemical["id"])

        else:
            st.info("No active chemicals found.")

    # --------------------------------------------------------
    # Disposal history
    # --------------------------------------------------------
    with tab2:
        st.subheader("Disposal History")

        col1, col2 = st.columns([3, 2])
        with col1:
            disposal_search = st.text_input(
                "Search Disposal History",
                placeholder="Bottle ID, chemical name, CAS No., manufacturer...",
            )
        with col2:
            reason_filter = st.selectbox(
                "Disposal Reason",
                [
                    "All",
                    "Empty / Fully Used",
                    "Expired",
                    "Damaged Container",
                    "Contaminated",
                    "Quality Concern",
                    "No Longer Needed",
                    "Duplicate / Excess Stock",
                    "Other",
                ],
            )

        query = f"""
        SELECT
            c.bottle_id AS "Bottle ID",
            c.legacy_id AS "Legacy ID",
            c.chemical_name AS "Chemical Name",
            c.cas_number AS "CAS No.",
            c.manufacturer AS "Manufacturer",
            COALESCE(
                NULLIF(c.disposal_location, ''),
                {LOCATION_EXPR}
            ) AS "Previous Location",
            COALESCE(NULLIF(c.disposal_reason, ''), 'Legacy / Unspecified') AS "Reason",
            c.disposal_note AS "Note",
            c.disposed_at AS "Disposal Date"
        FROM chemicals c
        LEFT JOIN storage_units su ON c.storage_unit_id = su.id
        WHERE c.status='Disposed'
        """
        params = []

        if disposal_search:
            like = f"%{disposal_search}%"
            query += """
            AND (
                   c.bottle_id ILIKE ?
                OR c.legacy_id ILIKE ?
                OR c.chemical_name ILIKE ?
                OR c.cas_number ILIKE ?
                OR c.manufacturer ILIKE ?
            )
            """
            params.extend([like] * 5)

        if reason_filter != "All":
            query += " AND c.disposal_reason=? "
            params.append(reason_filter)

        query += " ORDER BY c.disposed_at DESC, c.id DESC "

        conn = get_connection()
        disposed = conn.execute(query, params).fetchall()
        conn.close()

        if disposed:
            st.dataframe(
                [dict(row) for row in disposed],
                width="stretch",
                hide_index=True,
            )
            st.caption(f"{len(disposed):,} disposal record(s)")
        else:
            st.info("No disposal records found.")


# ============================================================
# SDS
# ============================================================

elif menu == "SDS":
    st.header("📄 SDS & Safety Management")
    st.caption(
        "Review automatic PubChem GHS screening and attach the official manufacturer SDS."
    )

    st.warning(
        "PubChem safety screening is not a substitute for the official product SDS. "
        "Signal words, H-codes, handling instructions, PPE, storage, and incompatibility "
        "must be verified against the manufacturer's current SDS."
    )

    include_disposed = st.checkbox("Include disposed chemicals", value=False)
    sds_search = st.text_input(
        "Search Chemical",
        placeholder="Chemical name, Bottle ID, CAS No., manufacturer, catalog number...",
    )

    conditions = []
    params = []

    if not include_disposed:
        conditions.append("c.status != 'Disposed'")

    if sds_search.strip():
        like = f"%{sds_search.strip()}%"
        conditions.append(
            """
            (
                   c.bottle_id ILIKE ?
                OR c.legacy_id ILIKE ?
                OR c.chemical_name ILIKE ?
                OR c.cas_number ILIKE ?
                OR c.manufacturer ILIKE ?
                OR c.catalog_number ILIKE ?
            )
            """
        )
        params.extend([like] * 6)

    where_sql = ""
    if conditions:
        where_sql = "WHERE " + " AND ".join(conditions)

    conn = get_connection()
    sds_rows = conn.execute(
        f"""
        SELECT
            c.id,
            c.bottle_id,
            c.chemical_name,
            c.cas_number,
            c.manufacturer,
            c.catalog_number,
            c.status,
            c.safety_signal_word,
            c.safety_hazard_statements,
            c.safety_checked_at,
            CASE WHEN sd.chemical_id IS NULL THEN 0 ELSE 1 END AS has_sds
        FROM chemicals c
        LEFT JOIN sds_documents sd ON sd.chemical_id = c.id
        {where_sql}
        ORDER BY c.chemical_name, c.id
        LIMIT 300
        """,
        params,
    ).fetchall()
    conn.close()

    if not sds_rows:
        st.info("No chemicals match the search.")
    else:
        st.caption(f"{len(sds_rows):,} result(s) shown")

        selection_options = [row["id"] for row in sds_rows]
        rows_by_id = {row["id"]: row for row in sds_rows}

        selected_id = st.selectbox(
            "Select Chemical",
            selection_options,
            format_func=lambda chemical_id: (
                f"{rows_by_id[chemical_id]['chemical_name']}  |  "
                f"{rows_by_id[chemical_id]['bottle_id'] or '-'}  |  "
                f"CAS {rows_by_id[chemical_id]['cas_number'] or '-'}  |  "
                f"{safety_badge(rows_by_id[chemical_id]['safety_signal_word'], rows_by_id[chemical_id]['safety_hazard_statements'], rows_by_id[chemical_id]['safety_checked_at'])}  |  "
                f"{'📄 SDS attached' if rows_by_id[chemical_id]['has_sds'] else 'No SDS'}"
            ),
        )

        conn = get_connection()
        chemical = conn.execute(
            "SELECT * FROM chemicals WHERE id=?",
            (selected_id,),
        ).fetchone()
        sds_doc = conn.execute(
            "SELECT * FROM sds_documents WHERE chemical_id=?",
            (selected_id,),
        ).fetchone()
        conn.close()

        if chemical:
            st.divider()
            st.subheader(f"🧪 {chemical['chemical_name']}")
            st.caption(
                f"Bottle ID: {chemical['bottle_id'] or '-'}   |   "
                f"CAS No.: {chemical['cas_number'] or '-'}   |   "
                f"Manufacturer: {chemical['manufacturer'] or '-'}   |   "
                f"Catalog No.: {chemical['catalog_number'] or '-'}"
            )

            st.markdown("### 1. Safety Screening")

            action1, action2 = st.columns([1, 2])
            with action1:
                check_label = (
                    "Refresh Safety Screening"
                    if chemical["safety_checked_at"]
                    else "Check Safety Screening"
                )
                check_safety = st.button(
                    check_label,
                    type="primary",
                    width="stretch",
                    key=f"check_safety_{selected_id}",
                )

            with action2:
                if chemical["safety_checked_at"]:
                    st.caption(
                        f"Last checked: {chemical['safety_checked_at']}  |  "
                        f"Matched using: {chemical['safety_match_term'] or '-'}"
                    )

            if check_safety:
                with st.spinner("Checking PubChem GHS information..."):
                    safety = fetch_pubchem_safety(
                        chemical["cas_number"],
                        chemical["chemical_name"],
                    )

                if safety:
                    save_pubchem_safety(selected_id, safety)
                    set_notice("✅ Safety screening has been updated.")
                    st.rerun()
                else:
                    st.error(
                        "No reliable PubChem match was found from the CAS No. or chemical name. "
                        "Do not assume the chemical is non-hazardous; verify the official SDS."
                    )

            render_safety_panel(
                signal_word=chemical["safety_signal_word"],
                h_codes=chemical["safety_h_codes"],
                hazard_statements=chemical["safety_hazard_statements"],
                pictograms=chemical["safety_pictograms"],
                source_url=chemical["safety_source_url"],
                checked_at=chemical["safety_checked_at"],
            )

            st.divider()
            st.markdown("### 2. Official SDS")

            if sds_doc:
                st.success("✅ An official SDS record is attached to this bottle.")

                meta1, meta2, meta3 = st.columns(3)
                meta1.markdown(
                    f"**Source**\n\n{sds_doc['source_name'] or chemical['manufacturer'] or '-'}"
                )
                meta2.markdown(
                    f"**Revision Date**\n\n{sds_doc['revision_date'] or '-'}"
                )
                meta3.markdown(
                    f"**Uploaded**\n\n{sds_doc['uploaded_at'] or '-'}"
                )

                if sds_doc["source_url"]:
                    st.link_button("Open Official SDS Source", sds_doc["source_url"])

                if sds_doc["file_data"]:
                    st.download_button(
                        "📄 Download Attached SDS PDF",
                        data=bytes(sds_doc["file_data"]),
                        file_name=sds_doc["file_name"] or "SDS.pdf",
                        mime="application/pdf",
                        key=f"download_sds_{selected_id}",
                    )

                st.caption(
                    "Uploading a new PDF below will replace the current attached SDS record."
                )
            else:
                st.info(
                    "No official manufacturer SDS is attached yet. "
                    "Upload the current product SDS or save its official source link below."
                )

            default_source_name = ""
            default_source_url = ""
            default_revision_date = ""

            if sds_doc:
                default_source_name = sds_doc["source_name"] or ""
                default_source_url = sds_doc["source_url"] or ""
                default_revision_date = sds_doc["revision_date"] or ""
            else:
                default_source_name = chemical["manufacturer"] or ""

            with st.form(
                f"official_sds_form_{selected_id}",
                clear_on_submit=False,
                enter_to_submit=False,
            ):
                source_name = st.text_input(
                    "SDS Source / Manufacturer",
                    value=default_source_name,
                    placeholder="e.g. Sigma-Aldrich",
                )
                source_url = st.text_input(
                    "Official SDS URL",
                    value=default_source_url,
                    placeholder="https://...",
                )
                revision_date = st.text_input(
                    "SDS Revision Date",
                    value=default_revision_date,
                    placeholder="YYYY-MM-DD",
                )
                sds_pdf = st.file_uploader(
                    "Upload Official SDS PDF",
                    type=["pdf"],
                    key=f"official_sds_pdf_{selected_id}",
                )

                save_sds = st.form_submit_button(
                    "Save Official SDS",
                    type="primary",
                )

            if save_sds:
                existing_file_name = sds_doc["file_name"] if sds_doc else None
                existing_file_data = sds_doc["file_data"] if sds_doc else None

                if sds_pdf is not None:
                    file_name = sds_pdf.name
                    file_data = sds_pdf.getvalue()
                else:
                    file_name = existing_file_name
                    file_data = existing_file_data

                if not file_data and not source_url.strip():
                    st.error(
                        "Upload an SDS PDF or provide the official SDS source URL."
                    )
                else:
                    conn = get_connection()
                    conn.execute(
                        """
                        INSERT INTO sds_documents (
                            chemical_id, file_name, file_data,
                            source_name, source_url, revision_date, uploaded_at
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                        ON CONFLICT(chemical_id) DO UPDATE SET
                            file_name=excluded.file_name,
                            file_data=excluded.file_data,
                            source_name=excluded.source_name,
                            source_url=excluded.source_url,
                            revision_date=excluded.revision_date,
                            uploaded_at=excluded.uploaded_at
                        """,
                        (
                            selected_id,
                            file_name,
                            file_data,
                            source_name.strip(),
                            source_url.strip(),
                            revision_date.strip(),
                            datetime.now().strftime("%Y-%m-%d %H:%M"),
                        ),
                    )
                    write_audit_log(
                        "SAVE_SDS",
                        target_type="chemical",
                        target_id=selected_id,
                        bottle_id=chemical["bottle_id"],
                        details=json.dumps(
                            {
                                "source_name": source_name.strip(),
                                "source_url": source_url.strip(),
                                "revision_date": revision_date.strip(),
                                "file_name": file_name,
                            },
                            ensure_ascii=False,
                        ),
                        conn=conn,
                    )
                    conn.commit()
                    conn.close()
                    set_notice("✅ Official SDS record has been saved.")
                    st.rerun()

            if sds_doc:
                if st.button(
                    "Remove Attached SDS Record",
                    key=f"remove_sds_{selected_id}",
                ):
                    conn = get_connection()
                    conn.execute(
                        "DELETE FROM sds_documents WHERE chemical_id=?",
                        (selected_id,),
                    )
                    write_audit_log(
                        "REMOVE_SDS",
                        target_type="chemical",
                        target_id=selected_id,
                        bottle_id=chemical["bottle_id"],
                        details="Removed attached official SDS record",
                        conn=conn,
                    )
                    conn.commit()
                    conn.close()
                    set_notice("SDS record removed.")
                    st.rerun()

    st.divider()
    with st.expander("How this safety system works"):
        st.write(
            """
            1. New manually registered chemicals are screened automatically using CAS No. first.
            2. PubChem GHS data are stored as a safety screening record.
            3. Danger/Warning information is displayed prominently when available.
            4. The manufacturer SDS is managed separately as the official safety document.
            5. A missing PubChem warning must never be interpreted as proof that a chemical is safe.
            """
        )



# ============================================================
# USER MANAGEMENT
# ============================================================

elif menu == "User Management":
    if CURRENT_USER["role"] != "admin":
        st.error("Administrator access is required.")
        st.stop()

    st.header("👥 User Management")
    st.caption(
        "Only active users listed here can access the laboratory inventory after Google sign-in."
    )

    conn = get_connection()
    users = conn.execute(
        """
        SELECT id, email, display_name, role, active, created_at, updated_at, created_by
        FROM app_users
        ORDER BY active DESC, role, LOWER(email)
        """
    ).fetchall()
    conn.close()

    active_users = sum(1 for user in users if bool(user["active"]))
    admin_users = sum(
        1
        for user in users
        if bool(user["active"]) and str(user["role"]).lower() == "admin"
    )

    metric1, metric2, metric3 = st.columns(3)
    metric1.metric("Registered Users", len(users))
    metric2.metric("Active Users", active_users)
    metric3.metric("Active Admins", admin_users)

    st.divider()
    st.subheader("Add or Update User")

    user_choices = ["➕ New user"] + [user["email"] for user in users]
    selected_user_choice = st.selectbox(
        "User",
        user_choices,
        key="user_management_selector",
    )

    selected_user = None
    if selected_user_choice != "➕ New user":
        selected_user = next(
            (user for user in users if user["email"] == selected_user_choice),
            None,
        )

    default_email = selected_user["email"] if selected_user else ""
    default_name = selected_user["display_name"] if selected_user else ""
    default_role = (
        str(selected_user["role"] or "member").lower()
        if selected_user
        else "member"
    )
    default_active = bool(selected_user["active"]) if selected_user else True

    role_options = ["member", "admin"]
    if default_role not in role_options:
        role_options.append(default_role)

    with st.form("manage_user_form", enter_to_submit=False):
        user_email = st.text_input(
            "Google Account Email *",
            value=default_email,
            disabled=selected_user is not None,
            placeholder="name@example.com",
        )
        display_name = st.text_input(
            "Display Name",
            value=default_name or "",
            placeholder="Name shown in audit records",
        )
        role = st.selectbox(
            "Role",
            role_options,
            index=role_options.index(default_role),
            help="Admins can manage users and view the full activity log.",
        )
        active = st.checkbox(
            "Active",
            value=default_active,
            help="Inactive users can sign in to Google but cannot access the inventory.",
        )
        save_user = st.form_submit_button(
            "Save User",
            type="primary",
        )

    if save_user:
        normalized_email = user_email.strip().lower()
        normalized_name = display_name.strip() or normalized_email

        if not normalized_email or "@" not in normalized_email:
            st.error("Enter a valid Google account email.")
        elif (
            normalized_email == CURRENT_USER["email"]
            and (not active or role != "admin")
        ):
            st.error(
                "You cannot deactivate or remove admin rights from your own current account."
            )
        else:
            conn = get_connection()
            existing = conn.execute(
                """
                SELECT id, email, role, active
                FROM app_users
                WHERE LOWER(email)=LOWER(?)
                """,
                (normalized_email,),
            ).fetchone()

            # Protect the last active administrator.
            if (
                existing
                and str(existing["role"]).lower() == "admin"
                and bool(existing["active"])
                and (role != "admin" or not active)
            ):
                active_admin_count = conn.execute(
                    """
                    SELECT COUNT(*)
                    FROM app_users
                    WHERE active=TRUE AND LOWER(role)='admin'
                    """
                ).fetchone()[0]

                if active_admin_count <= 1:
                    conn.close()
                    st.error("At least one active administrator must remain.")
                    st.stop()

            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            if existing:
                conn.execute(
                    """
                    UPDATE app_users
                    SET display_name=?,
                        role=?,
                        active=?,
                        updated_at=?
                    WHERE id=?
                    """,
                    (
                        normalized_name,
                        role,
                        active,
                        now,
                        existing["id"],
                    ),
                )
                user_id = existing["id"]
                audit_action = "EDIT_USER"
            else:
                cursor = conn.execute(
                    """
                    INSERT INTO app_users (
                        email, display_name, role, active,
                        created_at, updated_at, created_by
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        normalized_email,
                        normalized_name,
                        role,
                        active,
                        now,
                        now,
                        CURRENT_USER["email"],
                    ),
                )
                user_id = cursor.lastrowid
                audit_action = "ADD_USER"

            write_audit_log(
                audit_action,
                target_type="user",
                target_id=user_id,
                details=json.dumps(
                    {
                        "email": normalized_email,
                        "display_name": normalized_name,
                        "role": role,
                        "active": active,
                    },
                    ensure_ascii=False,
                ),
                conn=conn,
            )
            conn.commit()
            conn.close()
            set_notice(f"✅ User settings saved for {normalized_email}.")
            st.rerun()

    st.divider()
    st.subheader("Registered Users")

    if users:
        user_table = []
        for user in users:
            user_table.append(
                {
                    "Name": user["display_name"] or "-",
                    "Email": user["email"],
                    "Role": str(user["role"] or "member").title(),
                    "Status": "Active" if bool(user["active"]) else "Inactive",
                    "Created": user["created_at"] or "-",
                }
            )

        st.dataframe(
            user_table,
            hide_index=True,
            width="stretch",
        )
    else:
        st.info("No users are registered yet.")


# ============================================================
# ACTIVITY LOG
# ============================================================

elif menu == "Activity Log":
    if CURRENT_USER["role"] != "admin":
        st.error("Administrator access is required.")
        st.stop()

    st.header("📋 Activity Log")
    st.caption(
        "Shows who registered, edited, moved, disposed of, or otherwise changed laboratory records."
    )

    conn = get_connection()
    action_rows = conn.execute(
        """
        SELECT DISTINCT action
        FROM audit_logs
        WHERE action IS NOT NULL AND TRIM(action) != ''
        ORDER BY action
        """
    ).fetchall()
    conn.close()

    action_options = ["All Actions"] + [row["action"] for row in action_rows]

    filter1, filter2 = st.columns([2.5, 1.2])
    with filter1:
        activity_search = st.text_input(
            "Search Activity",
            placeholder="User, email, action, Bottle ID, details...",
        )
    with filter2:
        selected_action = st.selectbox(
            "Action",
            action_options,
        )

    conditions = []
    params = []

    if activity_search.strip():
        like = f"%{activity_search.strip()}%"
        conditions.append(
            """
            (
                   user_name ILIKE ?
                OR user_email ILIKE ?
                OR action ILIKE ?
                OR bottle_id ILIKE ?
                OR details ILIKE ?
            )
            """
        )
        params.extend([like] * 5)

    if selected_action != "All Actions":
        conditions.append("action=?")
        params.append(selected_action)

    where_sql = ""
    if conditions:
        where_sql = "WHERE " + " AND ".join(conditions)

    conn = get_connection()
    logs = conn.execute(
        f"""
        SELECT
            created_at,
            user_name,
            user_email,
            action,
            target_type,
            target_id,
            bottle_id,
            details
        FROM audit_logs
        {where_sql}
        ORDER BY id DESC
        LIMIT 1000
        """,
        params,
    ).fetchall()
    conn.close()

    st.caption(f"{len(logs):,} activity record(s) shown")

    if logs:
        display_logs = []
        for log in logs:
            display_logs.append(
                {
                    "Time": log["created_at"],
                    "User": log["user_name"] or "-",
                    "Email": log["user_email"],
                    "Action": log["action"],
                    "Bottle ID": log["bottle_id"] or "-",
                    "Target": (
                        f"{log['target_type'] or '-'}"
                        + (
                            f" #{log['target_id']}"
                            if log["target_id"] is not None
                            else ""
                        )
                    ),
                    "Details": log["details"] or "-",
                }
            )

        st.dataframe(
            display_logs,
            hide_index=True,
            width="stretch",
        )
    else:
        st.info("No activity records match the selected filters.")


# ============================================================
# SETTINGS
# ============================================================

elif menu == "Settings":
    st.header("⚙️ Settings")

    conn = get_connection()
    active_count = conn.execute(
        "SELECT COUNT(*) FROM chemicals WHERE status != 'Disposed'"
    ).fetchone()[0]
    disposed_count = conn.execute(
        "SELECT COUNT(*) FROM chemicals WHERE status='Disposed'"
    ).fetchone()[0]
    storage_count = conn.execute("SELECT COUNT(*) FROM storage_units").fetchone()[0]
    user_count = conn.execute(
        "SELECT COUNT(*) FROM app_users WHERE active=TRUE"
    ).fetchone()[0]
    audit_count = conn.execute("SELECT COUNT(*) FROM audit_logs").fetchone()[0]
    conn.close()

    st.subheader("System Information")
    st.write("Database: Supabase PostgreSQL")
    st.write(f"Signed in as: {CURRENT_USER['name']} ({CURRENT_USER['email']})")
    st.write(f"Role: {CURRENT_USER['role'].title()}")
    st.write(f"Active chemicals: {active_count:,}")
    st.write(f"Disposed chemicals: {disposed_count:,}")
    st.write(f"Storage units: {storage_count:,}")
    st.write(f"Active users: {user_count:,}")
    st.write(f"Audit records: {audit_count:,}")

    st.divider()
    st.subheader("Current Features")
    st.write(
        """
        - Google OIDC login
        - Approved-user whitelist with Admin / Member roles
        - User activity / audit logging
        - Manual chemical registration
        - Lab Manager Excel bulk import
        - Excel inventory export
        - Automatic Bottle ID generation
        - Legacy management ID retention
        - Storage-unit auto creation from Excel
        - Storage renaming and shelf management
        - Chemical search
        - Active inventory pagination
        - Chemical editing and relocation
        - Disposal with predefined reasons
        - Disposal history
        - Automatic PubChem GHS safety screening for new chemicals
        - Official SDS PDF / source-link management
        - Safety warnings in chemical detail view
        - Modal T4 serverless GPU OCR (ReagentOCRV7)
        - Multi-image / ZIP Inventory Scan with 3×2 review cards
        - Existing bottle reconciliation + new physical-bottle registration
        - Physical bottle image / scan history
        - OCR-assisted Add Chemical prefill
        - App-wide OCR backend shared by all logged-in users
        """
    )

    st.divider()
    st.subheader("Planned Features")
    st.write(
        """
        - QR code generation
        - Manufacturer-specific SDS auto-discovery
        - AI SDS summarization
        - Natural-language inventory search
        - Automated backup / reporting
        """
    )

# ICL_V8_GUARD_APPLIED


def _icl_live_response(response):
    status = response.status_code
    if status == 401:
        raise RuntimeError("인증 실패(401): Streamlit Secrets의 OCR_API_KEY와 Modal Secret icl-ocr-api의 OCR_API_KEY가 일치해야 합니다. 관리자라면 아래 임시 키 입력으로 확인할 수 있습니다.")
    try:
        payload = response.json()
    except (ValueError, TypeError) as exc:
        raise RuntimeError("OCR 서버 HTTP " + str(status) + ": JSON 응답이 아닙니다. 서버 실행 로그를 확인하세요.") from exc
    if not isinstance(payload, dict):
        raise RuntimeError("OCR 서버 HTTP " + str(status) + ": 응답 형식이 올바르지 않습니다.")
    failed = status != 200 or payload.get("ok") is False or bool(payload.get("error"))
    if failed:
        kind = str(payload.get("error_type") or "ServerError")
        detail = str(payload.get("error") or "서버가 실패 응답을 반환했지만 상세 메시지가 없습니다.").strip()
        if not detail:
            detail = "서버 오류 메시지가 비어 있습니다."
        raise RuntimeError("OCR 서버 HTTP " + str(status) + " · " + kind + ": " + detail[:500])
    return payload



def _icl_public_report(value):
    # Signed tokens are credentials; diagnostic downloads omit them and images.
    excluded = {"review_token", "evidence_token", "draft_token", "label_crop_base64", "preview_base64"}
    if isinstance(value, dict):
        return {k: _icl_public_report(v) for k, v in value.items() if k not in excluded}
    if isinstance(value, list):
        return [_icl_public_report(v) for v in value]
    return value


def render_icl_live_test():
    st.header("OCR Live Test")
    st.caption("화면 버전: 20260907-fix2 · 전체 시약관리 사이트 내 테스트")
    st.caption("사진을 올려 원문과 검증된 제품 정보를 비교하세요. 이 화면의 테스트는 재고 DB에 저장하지 않습니다.")
    st.caption("연결 주소: " + get_ocr_api_url())
    if _ocr_user_is_admin():
        with st.expander("인증 문제가 있을 때만: 임시 API 키"):
            st.text_input("OCR_API_KEY 값", type="password", key="icl_v8_temporary_api_key", help="현재 관리자 세션에서만 사용합니다. 정상 연결되면 Streamlit Secrets에도 같은 값을 설정하세요.")
            st.caption("비워 두면 기존 Streamlit Secrets의 키를 사용합니다. Modal 로그인 토큰과는 다른 값입니다.")
    key = get_ocr_api_key()
    if not key:
        st.warning("관리자가 Streamlit Secrets에 OCR_API_KEY를 설정해야 합니다.")
    if st.button("서버·인증·GPU 연결 확인", key="icl_live_health", disabled=not bool(key)):
        try:
            with st.spinner("서버와 GPU를 확인하는 중입니다. 첫 실행은 시간이 걸릴 수 있습니다."):
                state = _icl_live_response(requests.post(get_ocr_api_url() + "/health/deep", headers={"X-API-Key": key}, json={}, timeout=240))
            if not state.get("ok") or not state.get("cuda"):
                st.error("GPU 실행 확인에 실패했습니다.")
            else:
                st.success("인증 성공 · GPU OCR 실행 확인")
            st.json(state)
        except Exception as exc:
            st.error(_icl_exception_message(exc))
    uploaded = st.file_uploader("시약 사진", type=["jpg", "jpeg", "png", "webp"], key="icl_live_photo")
    if uploaded is None:
        return
    data = uploaded.getvalue()
    image_id = hashlib.sha256(data).hexdigest()
    st.image(data, width=450)
    if st.button("이 사진 OCR 실행", key="icl_live_run", disabled=not bool(key)):
        st.session_state.pop("icl_live_result", None)
        st.session_state.pop("icl_live_draft", None)
        try:
            with st.spinner("원본 판독 및 제품 정보를 확인하고 있습니다."):
                response = requests.post(get_ocr_api_url() + "/ocr", headers={"X-API-Key": key}, files={"file": (uploaded.name, data, uploaded.type or "image/jpeg")}, timeout=240)
                result = _icl_live_response(response)
                if result.get("schema_version") != "icl.ocr.v8":
                    raise RuntimeError("V8 서버 주소인지 확인하세요.")
            st.session_state["icl_live_result"] = {"image_id": image_id, "result": result}
        except Exception as exc:
            st.error(_icl_exception_message(exc))
    saved = st.session_state.get("icl_live_result", {})
    if saved.get("image_id") != image_id:
        return
    result = saved["result"]
    guard = __import__("icl_ocr_ui_guard")
    guard.render_review(globals(), result)
    label_fields = result.get('fields') or {}
    label_codes = {str(x.get('value', '')).strip().upper() for x in label_fields.get('catalog', [])} - {''}
    choices = [c for c in guard.candidates(result) if c.get('catalog_correction') or not c.get('catalog') or not label_codes or str(c['catalog']).strip().upper() in label_codes]
    if len(choices) < len(guard.candidates(result)):
        st.warning('Candidates with a different catalog number were excluded. Rescan to refresh the lookup.')
    if label_codes:
        st.caption('Catalog read from label: ' + ' / '.join(sorted(label_codes)))
    if not choices:
        st.warning("확인 가능한 제품 후보가 없습니다. 라벨을 다른 각도에서 촬영하거나 수동 등록하세요.")
    else:
        selected = st.radio("제품 후보 선택", options=range(len(choices)), index=None, format_func=lambda i: guard.candidate_label(choices[i]), key="live_" + guard.choice_key(result))
        if selected is not None:
            candidate = choices[selected]
            guard.render_candidate(globals(), result, candidate)
            reviewed = bool(st.session_state.get(guard.review_key(result, candidate), False))
            blocked = result.get("verification_status") == "CONFLICT" or (not candidate.get("product_verified") and not reviewed)
            if st.button("등록에 적용될 값 확인", disabled=blocked, key="icl_live_draft_check"):
                st.session_state.pop("icl_live_draft", None)
                try:
                    draft = _icl_live_response(requests.post(get_ocr_api_url() + "/draft", headers={"X-API-Key": key}, json={"review_token": result.get("review_token"), "candidate_id": candidate["candidate_id"], "reviewed": reviewed}, timeout=60))
                    if draft["fields"].get("name") != candidate.get("name"):
                        raise RuntimeError("선택한 이름과 적용할 이름이 다릅니다. 적용을 중단했습니다.")
                    st.session_state["icl_live_draft"] = {"image_id": image_id, "review_token": result.get("review_token"), "candidate_id": candidate["candidate_id"], "fields": draft["fields"]}
                except Exception as exc:
                    st.error(_icl_exception_message(exc))
            draft = st.session_state.get("icl_live_draft", {})
            if draft.get("image_id") == image_id and draft.get("review_token") == result.get("review_token") and draft.get("candidate_id") == candidate["candidate_id"]:
                st.success("선택한 제품명과 적용할 이름이 일치합니다. DB에는 저장하지 않았습니다.")
                st.json(draft["fields"])
    st.download_button("진단 결과 다운로드", data=json.dumps(_icl_public_report(result), ensure_ascii=False, indent=2), file_name="ocr_live_result.json", mime="application/json")


if menu == "OCR Live Test":
    render_icl_live_test()


st.sidebar.caption(
    f"This page: {time.perf_counter() - _ICL_PERF['start']:.1f}s · "
    f"database: {_ICL_PERF['db_calls']} calls, {_ICL_PERF['db_seconds']:.1f}s"
)
